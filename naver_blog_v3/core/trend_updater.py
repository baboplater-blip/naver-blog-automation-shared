# core/trend_updater.py — ④ 실시간 검색량 자동 업데이트
# 네이버 데이터랩 API로 매주 키워드 트렌드 갱신
# https://developers.naver.com/docs/serviceapi/datalab/search/search.md

import json, os, sys, time, urllib.request
from datetime import datetime, date, timedelta
from pathlib import Path

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config.logger import get_logger
log = get_logger(__name__)
BASE_DIR      = Path(__file__).parent.parent
from config import profile as _nbaprofile
_DATA_DIR = _nbaprofile.app_data_root()

TREND_DB_PATH = _DATA_DIR / "logs" / "trend_data.json"


def _load_trend_db() -> dict:
    from config import safe_load_json
    return safe_load_json(TREND_DB_PATH, default=dict)

def _save_trend_db(data: dict):
    from config import atomic_write_json
    atomic_write_json(TREND_DB_PATH, data)


try:
    from config.settings import API_REQUEST_TIMEOUT
except (ImportError, AttributeError):
    API_REQUEST_TIMEOUT = 10


def _fetch_trend_scores(keywords: list, client_id: str, client_secret: str,
                        provider=None, transport=None) -> dict:
    """
    네이버 데이터랩 API 호출 → 키워드별 최근 4주 평균 검색 지수 반환
    한 번에 최대 5개 키워드. 여러 배치로 나눠 호출.
    Returns: {keyword: score(0~100)}
    """
    from core.naver_api import datalab_request, resolve_provider, runtime_credentials
    selected_provider = resolve_provider(provider)
    if selected_provider == "api_hub":
        client_id, client_secret = runtime_credentials(selected_provider)
    elif not client_id or not client_secret:
        client_id, client_secret = runtime_credentials(selected_provider)
    if not client_id or not client_secret:
        log.warning("선택한 네이버 API 제공자 자격증명 없음 → 트렌드 업데이트 스킵")
        return {}

    end   = date.today()
    start = end - timedelta(days=90)

    results = {}
    for i in range(0, len(keywords), 5):
        batch = keywords[i:i+5]
        groups = [{"groupName": kw, "keywords": [kw]} for kw in batch]
        body  = json.dumps({
            "startDate":     start.strftime("%Y-%m-%d"),
            "endDate":       end.strftime("%Y-%m-%d"),
            "timeUnit":      "week",
            "keywordGroups": groups,
        }, ensure_ascii=False).encode("utf-8")

        request = datalab_request(json.loads(body.decode("utf-8")), client_id, client_secret,
                                  selected_provider)
        req = urllib.request.Request(request["url"], data=request["body"],
                                     headers=request["headers"], method="POST")
        for _retry in range(2):  # 최대 1회 재시도
            try:
                if transport is not None:
                    data = transport(request)
                else:
                    with urllib.request.urlopen(req, timeout=API_REQUEST_TIMEOUT) as resp:
                        data = json.loads(resp.read())
                for item in data.get("results", []):
                    kw     = item["title"]
                    ratios = [p["ratio"] for p in item.get("data", [])[-4:] if "ratio" in p]
                    results[kw] = round(sum(ratios) / len(ratios), 1) if ratios else 0.0
                break  # 성공 시 재시도 루프 탈출
            except urllib.error.HTTPError as e:
                log.warning(f"데이터랩 API HTTP 오류 ({batch}): {e.code}")
                break  # HTTP 오류는 재시도하지 않음
            except (urllib.error.URLError, TimeoutError, ValueError, TypeError) as e:
                log.warning(f"데이터랩 API 네트워크 오류 ({len(batch)}개 키워드): {type(e).__name__}")
                if _retry == 0:
                    log.info("2초 후 재시도...")
                    time.sleep(2)

        if i + 5 < len(keywords):
            time.sleep(0.5)   # API 과호출 방지

    return results


def update_excel_trends(excel_path: str, sheet_name: str) -> dict:
    """
    엑셀에서 키워드를 읽어 검색 트렌드를 업데이트합니다.
    반환: {keyword: score} — GUI 그래프용
    """
    try:
        from config.settings import USE_TREND_UPDATE
        from config import settings as _settings
        if not USE_TREND_UPDATE:
            return {}
        legacy_client_id = str(getattr(_settings, "NAVER_API_CLIENT_ID", "") or "")
        legacy_client_secret = str(getattr(_settings, "NAVER_API_CLIENT_SECRET", "") or "")
    except ImportError:
        return {}

    if not os.path.exists(excel_path):
        return {}

    try:
        from openpyxl import load_workbook
        from openpyxl.styles import PatternFill, Font
        import re

        wb = load_workbook(excel_path)
        ws = wb[sheet_name]

        # 셀 제목 전체(30자 롱테일)를 데이터랩에 넣으면 검색되는 말이 아니라 지수가
        # 전부 0.0 이 된다(2026-07-19 실측·수리). 제목에 앵커된 헤드 키워드로 변환해
        # 조회하고, 중복 제거로 API 콜도 줄인다. 헤드 매칭이 안 되는 짧은 키워드
        # 풀(구 방식)은 원문 그대로 조회 — 기존 동작 보존.
        try:
            from core.search_volume import best_keyword as _best_kw
        except ImportError:
            _best_kw = None

        keywords = []
        title_to_head = {}
        for row in ws.iter_rows():
            for cell in row:
                if not cell.value:
                    continue
                kw = re.sub(r'^\d{4}-\d{2}-\d{2} — ', '', str(cell.value)).strip()
                if not kw:
                    continue
                head = ""
                if _best_kw is not None:
                    try:
                        head = _best_kw(kw)[0]
                    except Exception:  # noqa: BLE001 — 앵커 실패는 원문 조회
                        head = ""
                head = head or kw[:50]
                title_to_head[kw] = head
                if head not in keywords:
                    keywords.append(head)

        if not keywords:
            return {}

        log.info(f"{len(keywords)}개 키워드 트렌드 조회 중... "
                 f"(제목 {len(title_to_head)}건 → 헤드 키워드 축약)")
        scores = _fetch_trend_scores(keywords, legacy_client_id, legacy_client_secret)

        # 검색량 급등(70점↑) 키워드에 빨간 배경 표시
        RED_FILL = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
        updated  = 0
        for row in ws.iter_rows():
            for cell in row:
                if not cell.value:
                    continue
                kw = re.sub(r'^\d{4}-\d{2}-\d{2} — ', '', str(cell.value)).strip()
                s  = scores.get(title_to_head.get(kw, kw), 0)
                if s >= 70:
                    cell.fill = RED_FILL
                    updated += 1

        wb.save(excel_path)

        # DB 저장
        db = _load_trend_db()
        db[date.today().strftime("%Y-%m-%d")] = scores
        _save_trend_db(db)

        log.info(f"트렌드 업데이트 완료 | 급등 키워드 {updated}개 하이라이트")
        return scores

    except FileNotFoundError as e:
        log.error(f"트렌드 업데이트 실패 (파일 없음): {e}")
        return {}
    except KeyError as e:
        log.error(f"트렌드 업데이트 실패 (시트 없음): {e}")
        return {}
    except (OSError, ValueError) as e:
        log.error(f"트렌드 업데이트 실패: {e}")
        return {}


def get_top_keywords(n: int = 10) -> list:
    """
    최신 트렌드 DB에서 검색 지수 상위 N개 반환
    Returns: [{"keyword": str, "score": float, "date": str}, ...]
    """
    db = _load_trend_db()
    if not db:
        return []
    latest_date = max(db.keys())
    scores = db[latest_date]
    sorted_kws = sorted(scores.items(), key=lambda x: x[1], reverse=True)
    return [
        {"keyword": kw, "score": score, "date": latest_date}
        for kw, score in sorted_kws[:n]
    ]


def get_trend_history(keyword: str) -> list:
    """특정 키워드의 날짜별 스코어 히스토리 반환"""
    db = _load_trend_db()
    history = []
    for d in sorted(db.keys()):
        s = db[d].get(keyword)
        if s is not None:
            history.append({"date": d, "score": s})
    return history
