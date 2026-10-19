# scheduler/auto_scheduler.py — v3.0
# 전체 파이프라인: 중복검사 + 품질검증 + 트렌드 + 통계 + 예약발행

import math
import os, sys, json, re, time, traceback, urllib.error, shutil, random
from datetime import datetime, date, timedelta
from typing import Optional
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from pathlib import Path

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config.logger import get_logger
log = get_logger(__name__)

from core.content_generator import detect_category, generate_blog_post, markdown_to_naver_html
from core.duplicate_checker import check_duplicate, record_post
from core import institution_dedup
from core import draft_image_fill, draft_store, image_plan
from core.excluded_titles import add_excluded, is_excluded
from core.preview_saver import save_preview
from core.gov24_screenshot import capture_gov24_page
from core.image_generator import generate_all_images
from config import atomic_write_json, safe_load_json

# 하네스 관측 어댑터 — import 실패 시 _he=None 으로 모든 hook 이 no-op
try:
    from core import harness_emitter as _he
except Exception as _he_err:
    _he = None
    log.debug(f"harness_emitter import 실패 ({_he_err}) — 하네스 기록 비활성")

BASE_DIR           = Path(__file__).parent.parent
from config import profile as _nbaprofile
_DATA_DIR = _nbaprofile.app_data_root()

LOG_DIR            = _DATA_DIR / "logs"
LOG_DIR.mkdir(exist_ok=True)
# 발행 실패 큐 — 프로필별 데이터 루트에 격리(계정 간 재시도 큐 누수 방지).
# default 프로필은 _DATA_DIR == BASE_DIR(naver_blog_v3/) 이라 경로·동작 무변화.
FAILED_POSTS_PATH  = _DATA_DIR / "failed_posts.json"
YELLOW             = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
PIPELINE_LOCK_PATH = LOG_DIR / "pipeline.lock"


class PipelineLock:
    """파이프라인 중복 실행 방지용 프로세스 파일 락 (표준 라이브러리).

    작업 스케줄러 다중 등록 + GUI 동시 실행 시 상태파일 경합·이중예약을
    막는다. ``os.open`` 의 ``O_CREAT | O_EXCL`` 로 잠금파일을 배타 생성한다.
    이미 존재하면 획득 실패로 간주한다.

    크래시로 남은 stale 잠금파일은 기록된 PID 의 생존 여부로 판정해
    죽은 프로세스 소유면 회수(steal)한다. 신규 pip 의존성 없음.
    """

    def __init__(self, lock_path: Path) -> None:
        self._lock_path = Path(lock_path)
        self._fd: Optional[int] = None

    def _pid_alive(self, pid: int) -> bool:
        """PID 가 현재 살아 있는지 확인한다 (Windows/POSIX 공통)."""
        if pid <= 0:
            return False
        if os.name == "nt":
            import ctypes
            from ctypes import wintypes
            kernel = ctypes.WinDLL("kernel32", use_last_error=True)
            kernel.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
            kernel.OpenProcess.restype = wintypes.HANDLE
            kernel.WaitForSingleObject.argtypes = [wintypes.HANDLE, wintypes.DWORD]
            kernel.WaitForSingleObject.restype = wintypes.DWORD
            kernel.CloseHandle.argtypes = [wintypes.HANDLE]
            kernel.CloseHandle.restype = wintypes.BOOL
            # SYNCHRONIZE: 열린 핸들이 남은 종료 프로세스도 기다림 결과로 구별한다.
            handle = kernel.OpenProcess(0x00100000, False, pid)
            if not handle:
                # ERROR_INVALID_PARAMETER = 존재하지 않는 PID. 접근 거부 등은 보수적으로 유지.
                return ctypes.get_last_error() != 87
            try:
                return kernel.WaitForSingleObject(handle, 0) != 0
            finally:
                kernel.CloseHandle(handle)
        try:
            os.kill(pid, 0)
        except ProcessLookupError:
            return False
        except PermissionError:
            return True   # 존재하지만 권한 없음 → 살아 있음
        return True

    def _read_owner_pid(self) -> Optional[int]:
        """잠금파일에 기록된 소유 PID 를 읽는다 (파싱 실패 시 None)."""
        try:
            raw = self._lock_path.read_text(encoding="utf-8").strip()
            return int(raw) if raw else None
        except (OSError, ValueError):
            return None

    def acquire(self) -> bool:
        """락 획득 성공 시 True, 이미 잡혀 있으면 False."""
        try:
            self._lock_path.parent.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            log.warning(f"  ⚠️  잠금파일 디렉토리 생성 실패 ({exc}) — 락 없이 진행")
            return True   # 락 인프라 자체 실패는 파이프라인을 막지 않는다

        try:
            self._fd = os.open(
                str(self._lock_path),
                os.O_CREAT | os.O_EXCL | os.O_WRONLY,
            )
        except FileExistsError:
            owner_pid = self._read_owner_pid()
            if owner_pid is not None and not self._pid_alive(owner_pid):
                # 죽은 프로세스가 남긴 stale 락 → 회수 후 1회 재시도
                log.warning(
                    f"  ⚠️  죽은 프로세스(PID {owner_pid})의 잠금파일 회수 — "
                    f"{self._lock_path.name}"
                )
                try:
                    self._lock_path.unlink()
                    self._fd = os.open(
                        str(self._lock_path),
                        os.O_CREAT | os.O_EXCL | os.O_WRONLY,
                    )
                except (OSError, FileExistsError):
                    return False   # 경합에서 밀림 → 다른 실행에 양보
            else:
                return False
        except OSError as exc:
            log.warning(f"  ⚠️  잠금파일 생성 실패 ({exc}) — 락 없이 진행")
            return True   # 락 인프라 실패는 파이프라인을 막지 않는다

        # 소유 PID 기록 (stale 판정용)
        try:
            os.write(self._fd, str(os.getpid()).encode("utf-8"))
        except OSError:
            pass   # PID 기록 실패해도 배타 획득 자체는 유효
        return True

    def release(self) -> None:
        """락 해제 + 잠금파일 정리 (finally 에서 항상 호출 안전)."""
        if self._fd is not None:
            try:
                os.close(self._fd)
            except OSError:
                pass
            self._fd = None
        try:
            if self._lock_path.exists():
                self._lock_path.unlink()
        except OSError as exc:
            log.warning(f"  ⚠️  잠금파일 정리 실패 ({exc})")


def _cfg():
    """설정 동적 로드 (저장 후 즉시 반영). 재로드 정책은 config.reload_settings 한 곳."""
    import config.settings as s
    from config import reload_settings
    reload_settings()
    return s


def _keyword_excel(s) -> str:
    """활성 프로필의 키워드 엑셀 경로(브랜드 팩 우선). default 는 settings 와 동일."""
    _fallback = getattr(s, "KEYWORD_EXCEL_PATH", "")
    try:
        from config import brand as _brand
        return _brand.excel_path() or _fallback
    except Exception:  # noqa: BLE001
        return _fallback




# ── ★ 카테고리별 최적 발행 시간 슬롯 ─────────────────
#
#  독자층별 검색 피크 시간대:
#    복지·노인·장애  →  오전 (주부·노인 오전 활동 시간)
#    교육·장학금     →  오전~점심 (학부모 오전, 학생 점심)
#    바우처·의료     →  오후 (병원 다녀온 뒤 검색)
#    청년·취업       →  저녁 (퇴근 후, 취준생 집중 시간)
#    주거            →  밤   (부동산 고민은 밤에)
#    default         →  전 시간대 고르게
#
CATEGORY_TIME_SLOTS = {
    "복지":    ["09:00", "10:00", "11:00", "13:00", "14:00", "09:30", "10:30", "11:30"],
    "교육":    ["09:30", "11:00", "12:00", "13:30", "10:00", "11:30", "14:00", "09:00"],
    "바우처":  ["14:00", "15:00", "13:00", "16:00", "14:30", "15:30", "13:30", "16:30"],
    "의료":    ["10:00", "14:00", "11:00", "15:00", "10:30", "14:30", "11:30", "15:30"],
    "청년":    ["19:00", "20:00", "21:00", "19:30", "20:30", "22:00", "18:30", "21:30"],
    "취업":    ["19:30", "20:30", "21:00", "18:30", "20:00", "21:30", "19:00", "22:00"],
    "주거":    ["21:00", "22:00", "20:30", "21:30", "20:00", "22:30", "19:30", "23:00"],
    "default": ["07:00", "09:00", "11:00", "13:00", "15:00", "17:00", "19:30", "21:00"],
}

_SCHEDULE_AT_PATTERNS = (
    "%Y-%m-%d %H:%M",
    "%Y-%m-%dT%H:%M",
    "%Y-%m-%dT%H:%M:%S",
)


def _parse_schedule_at(raw) -> datetime | None:
    """publish_brief.yaml 의 schedule_at 을 datetime 으로 파싱.

    하네스에서 두 가지 포맷이 섞여 있다:
      - "2026-05-09 07:00"            (_lib/publish_brief.py 식)
      - "2026-05-22T14:00:00+09:00"   (harness_emitter.emit_publish_brief 식)
    타임존은 무시하고 KST 기준 naive datetime 으로 통일 — 비교는 KST 안에서만 한다.
    파싱 실패는 None.
    """
    if not raw:
        return None
    s = str(raw).strip()
    # 타임존 오프셋 (+09:00, Z) 제거
    s = re.sub(r"(?:Z|[+\-]\d{2}:?\d{2})$", "", s)
    # 분(min) 단위까지만 사용
    if len(s) >= 19 and s[10] in (" ", "T"):
        s = s[:19]
    elif len(s) >= 16 and s[10] in (" ", "T"):
        s = s[:16]
    for fmt in _SCHEDULE_AT_PATTERNS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


# 하네스 외곽 runs/ — 프로필 라우팅(2026-07-19 격리 감사): harness_emitter 가
# publish_brief 를 쓰는 곳과 같은 프로필 runs/ 를 스캔해야 자기 예약이 보인다.
# 테스트가 monkeypatch 할 수 있도록 모듈 상수로 분리. (2026-05-26)
try:
    from config import profile as _resv_profile
    _RESERVATIONS_DIR = _resv_profile.runs_dir()
except Exception:  # noqa: BLE001
    _RESERVATIONS_DIR = BASE_DIR.parent / "runs"


def _format_queue_diagnostic(*, target_date: str | None = None,
                              cfg=None) -> str:
    """run_pipeline 시작 시 출력할 큐 상태 한 줄을 생성.

    "큐: 비어 있음 → 새 예약 2026-05-27부터" 또는
    "큐: 2026-05-28까지 7편 등록됨 → 새 예약 2026-05-29부터 (+2일 이월)"
    같은 형식. RESERVATION_RESPECT_EXISTING=False 면 그 사실을 함께 표기.

    Args:
        target_date: 기준 일자 (YYYY-MM-DD). None 이면 오늘.
        cfg: 이미 로드된 settings. None 이면 내부에서 _cfg() 호출.

    Returns:
        사람이 읽는 한 줄 문자열. 예외 발생 시 빈 문자열.
    """
    try:
        s = cfg or _cfg()
        respect = bool(getattr(s, "RESERVATION_RESPECT_EXISTING", True))
        _t = target_date or date.today().strftime("%Y-%m-%d")
        base_d_plus_1 = (
            datetime.strptime(_t, "%Y-%m-%d").date() + timedelta(days=1)
        )

        existing = _load_existing_reservations() if respect else []
        last = max((r["date"] for r in existing), default=None)
        next_start = (
            max(base_d_plus_1, last + timedelta(days=1))
            if last else base_d_plus_1
        )
        carry = max((next_start - base_d_plus_1).days, 0)

        if not existing:
            head = "큐: 비어 있음"
        else:
            head = f"큐: {last}까지 {len(existing)}편 등록됨"
        msg = f"📅 {head} → 새 예약 {next_start}부터"
        if carry > 0:
            msg += f" (+{carry}일 이월)"
        if not respect:
            msg += " · ⚠ RESERVATION_RESPECT_EXISTING=False (옛 D+1 동작)"
        return msg
    except Exception as exc:  # noqa: BLE001 — 진단 실패는 메인 흐름 영향 0
        return f"📅 큐 진단 실패: {type(exc).__name__}: {exc}"


def _load_existing_reservations(*, runs_dir: Path | None = None,
                                 now: datetime | None = None) -> list[dict]:
    """runs/*/publish_brief.yaml 을 스캔해 **미래** 예약 슬롯 목록을 반환.

    네이버 블로그 큐에 실제로 등록된 예약 = `success: true` 인 brief 만 카운트.
    실패·차단·예외(`success: false`) 는 네이버 큐에 없어 슬롯 차지 안 함.

    Args:
        runs_dir: 명시되면 그 경로를 스캔. None 이면 모듈 상수 `_RESERVATIONS_DIR`
            (테스트 환경에서 monkeypatch 가능).

    Returns:
        [{"datetime": naive_kst_datetime, "date": date, "time": "HH:MM",
          "category": str, "slug": str}, ...]
        — 과거 항목은 자동 제외, datetime 오름차순 정렬.
    """
    now_dt = now or datetime.now()
    if runs_dir is None:
        runs_dir = _RESERVATIONS_DIR

    # runs 폴더가 없어도 예약 대장(logs/reservations.json)은 참조해야 하므로
    # 조기 반환하지 않는다 — brief 루프만 건너뛴다. (2026-07-04)
    out: list[dict] = []
    brief_iter = runs_dir.glob("*/publish_brief.yaml") if runs_dir.exists() else []
    for brief_path in brief_iter:
        try:
            import yaml  # type: ignore
            with brief_path.open("r", encoding="utf-8") as f:
                data = yaml.safe_load(f) or {}
        except (OSError, ImportError, Exception):
            continue
        if not isinstance(data, dict):
            continue
        if not data.get("success"):
            continue  # 미발행(실패·차단·예외) 은 큐 점유 없음
        dt = _parse_schedule_at(data.get("schedule_at"))
        if dt is None or dt < now_dt:
            continue
        out.append({
            "datetime": dt,
            "date":     dt.date(),
            "time":     dt.strftime("%H:%M"),
            "category": data.get("category") or "default",
            "slug":     data.get("slug") or brief_path.parent.name,
        })

    # ── 예약 대장 병합 (2026-07-04) ─────────────────────
    # brief 파일이 정리되거나 다른 환경이어도 이월이 끊기지 않도록, 영속 대장
    # (logs/reservations.json)의 미래 예약을 합친다. (slug, datetime) 로 중복 제거.
    try:
        from core import reservation_ledger as _rl
        seen = {(r["slug"], r["datetime"]) for r in out}
        for r in _rl.future_reservations(now=now_dt, prune=False):
            if (r["slug"], r["datetime"]) not in seen:
                out.append(r)
                seen.add((r["slug"], r["datetime"]))
    except Exception as _exc:  # noqa: BLE001 — 대장 실패는 brief 결과만 사용
        log.debug(f"[_load_existing_reservations] 예약 대장 병합 실패: {_exc}")

    out.sort(key=lambda r: r["datetime"])
    return out


def assign_reservation_times(count: int, target_date: str = None,
                              posts: list = None, cfg=None,
                              existing_reservations: list[dict] | None = None,
                              ) -> list:
    """
    포스팅 목록의 카테고리를 감지해서 각 글에 최적 발행 시간을 배정합니다.

    D+1부터 시작하여 POSTS_PER_DAY 개수만큼 하루에 배분합니다.
    예시: 10개 글, 하루 3개 → D+1에 3개, D+2에 3개, D+3에 3개, D+4에 1개

    - posts가 없으면 기존 방식(settings.py 슬롯 순서대로) 유지
    - 같은 시간대 카테고리끼리 충돌하지 않도록 30분 간격 자동 조정
    - 날짜가 바뀌면 used_times를 초기화하여 시간 슬롯 재사용 가능
    - cfg: 이미 로드된 설정 객체 (None이면 내부에서 로드)
    - existing_reservations: 이미 잡힌 미래 예약 슬롯 목록.
        None 이면 `_load_existing_reservations()` 로 runs/*/publish_brief.yaml
        에서 자동 로드. 마지막 예약 일자 + 1일을 새 시작점으로 사용해
        같은 날짜 슬롯 충돌을 막는다. (2026-05-26 도입)
    """
    s = cfg or _cfg()
    posts_per_day = getattr(s, 'POSTS_PER_DAY', 3)

    # 시작 기준일: target_date가 주어지면 그 다음 날(D+1), 없으면 내일(D+1)
    if target_date:
        base_date = datetime.strptime(target_date, "%Y-%m-%d").date()
    else:
        base_date = date.today()
    start_date = base_date + timedelta(days=1)

    # ── 예약 이월 보정 (2026-05-26) ─────────────────────
    # 이미 등록된 미래 예약이 있으면 그 마지막 일자 + 1 부터 시작해
    # 같은 슬롯에 두 번 잡히는 사고를 방지. settings.RESERVATION_RESPECT_EXISTING
    # = False 면 옛 동작(D+1 고정) 유지.
    if getattr(s, "RESERVATION_RESPECT_EXISTING", True):
        if existing_reservations is None:
            try:
                existing_reservations = _load_existing_reservations()
            except Exception as _exc:  # noqa: BLE001 — 헬퍼 실패는 무음 폴백
                log.debug(f"[assign_reservation_times] 예약 큐 로드 실패: {_exc}")
                existing_reservations = []
        if existing_reservations:
            last_reserved_date = max(r["date"] for r in existing_reservations)
            shifted = last_reserved_date + timedelta(days=1)
            if shifted > start_date:
                log.info(
                    f"  📅 예약 이월: 기존 예약 마지막 {last_reserved_date} → "
                    f"새 시작 {shifted} (기본 {start_date} 보다 +{(shifted - start_date).days}일)"
                )
                start_date = shifted

    if not posts:
        # 폴백: 기존 고정 슬롯 (다중 날짜 배분 적용)
        assigned = []
        for i in range(count):
            day_offset = i // posts_per_day
            current_date = (start_date + timedelta(days=day_offset)).strftime("%Y-%m-%d")
            slot_idx = i % min(posts_per_day, len(s.RESERVATION_TIME_SLOTS))
            slot = s.RESERVATION_TIME_SLOTS[slot_idx] if slot_idx < len(s.RESERVATION_TIME_SLOTS) else s.RESERVATION_TIME_SLOTS[0]
            assigned.append(f"{current_date} {slot}")
        return assigned

    assigned = []
    used_times = set()          # 현재 날짜에서 이미 배정된 시간 (충돌 방지)
    current_day_count = 0       # 현재 날짜에 배정된 글 수
    day_offset = 0              # 시작일로부터의 날짜 오프셋

    for info in posts[:count]:
        # 하루 할당량 초과 시 다음 날로 이동
        if current_day_count >= posts_per_day:
            day_offset += 1
            current_day_count = 0
            used_times = set()  # 새 날짜이므로 시간 슬롯 초기화

        current_date = (start_date + timedelta(days=day_offset)).strftime("%Y-%m-%d")

        raw_title = info.get("title", "")
        clean     = re.sub(r'^\d{4}-\d{2}-\d{2} — ', '', raw_title)
        category  = detect_category(clean)

        slots = CATEGORY_TIME_SLOTS.get(category, CATEGORY_TIME_SLOTS["default"])

        # 충돌 없는 첫 번째 슬롯 선택
        chosen = None
        for slot in slots:
            if slot not in used_times:
                chosen = slot
                break

        # 모든 슬롯이 사용된 경우 → 마지막 슬롯에서 30분씩 밀기
        if chosen is None:
            last = slots[-1]
            h, m  = map(int, last.split(":"))
            total_min = h * 60 + m
            for offset in range(30, 300, 30):
                candidate_min = (total_min + offset) % (24 * 60)
                candidate = f"{candidate_min // 60:02d}:{candidate_min % 60:02d}"
                if candidate not in used_times:
                    chosen = candidate
                    break
            if chosen is None:
                chosen = slots[0]   # 최후 폴백

        used_times.add(chosen)
        current_day_count += 1
        assigned.append(f"{current_date} {chosen}")

    return assigned


# ── 엑셀에서 미완료 포스팅 추출 ──────────────────────
def _load_latest_trend_scores() -> dict:
    """trend_data.json에서 최신 날짜의 키워드별 스코어를 반환합니다.
    Returns: {keyword: score} 또는 빈 dict
    """
    trend_path = _DATA_DIR / "logs" / "trend_data.json"
    trend_db = safe_load_json(trend_path, default=dict)
    if not trend_db:
        return {}
    latest_date = max(trend_db.keys())
    return trend_db.get(latest_date, {})


def _load_deadline_days() -> dict:
    """deadline_data.json에서 키워드별 잔여일수를 반환합니다.
    Returns: {keyword: days_left}
    """
    deadline_path = _DATA_DIR / "logs" / "deadline_data.json"
    items = safe_load_json(deadline_path, default=list)
    result = {}
    for item in items:
        kw = item.get("keyword", "").strip()
        if kw:
            result[kw] = item.get("days_left", 999)
    return result


def _calc_deadline_urgency(days_left: int) -> float:
    """마감 잔여일수를 긴급도 점수(0~100)로 변환합니다."""
    if days_left <= 0:
        return 100.0
    elif days_left <= 7:
        return 80.0
    elif days_left <= 14:
        return 50.0
    else:
        return 0.0


def _extract_keyword_from_title(title: str) -> str:
    """엑셀 제목에서 순수 키워드를 추출합니다.
    예: '2026-04-06 — 청년 월세 지원' → '청년 월세 지원'
    예: '2026-04-06 — 청년 월세 지원 (D-3 마감임박)' → '청년 월세 지원'
    """
    kw = re.sub(r'^\d{4}-\d{2}-\d{2}\s*—\s*', '', title).strip()
    kw = re.sub(r'\s*\(D-\d+\s*마감임박\)\s*$', '', kw).strip()
    return kw


def _volume_score(volume: int) -> float:
    """실측 월간검색량 → 0~100 점수(로그 스케일).

    100회→25 · 1,000회→50 · 10,000회→75 · 100,000회 이상→100.
    데이터랩 상대지수(0~100)와 같은 눈금이라 복합식 가중치를 그대로 쓴다.
    """
    if volume <= 0:
        return 0.0
    return max(0.0, min(100.0, 25.0 * (math.log10(volume) - 1.0)))


def _sort_by_trend_priority(pending: list) -> list:
    """검색 수요와 마감 긴급도를 결합하여 발행 우선순위를 정렬합니다.
    복합 점수 = (수요 스코어 * 0.6) + (마감 긴급도 * 0.4)

    수요 스코어 소스 우선순위(2026-07-19 무작동 수리):
      ① 제목에 앵커된 헤드 키워드의 실측 월간검색량(검색량 엑셀, 검색광고 API 병합분)
      ② 폴백: 데이터랩 상대지수(trend_data.json) — 단 0.0 은 '데이터 없음'으로 무시
      ③ 둘 다 없으면 중립 40점(엑셀 순서 보존 · 실측 500회≈42점이 미지보다 우선)
    제목 전체(30자 롱테일)를 데이터랩에 넣어 전부 0.0 이 되던 병을 rank_tracker #7 과
    같은 수술(헤드 키워드 추출)로 고쳤다. 정렬 근거는 상위 10건 로그 +
    logs/queue_sort_rationale.json 으로 남겨 홈 탭이 보여준다.

    이전 사이클에서 100% 중복으로 Skip된 제목은 excluded_titles.json 에 누적되며
    여기서 후보 단계 사전 차감한다. trend_data.json 자체는 건드리지 않는다.
    """
    # ── 사전 제외: 이미 중복으로 확인된 제목을 후보에서 차감 ──
    if pending:
        try:
            kept = [item for item in pending if not is_excluded(item.get("title", ""))]
            removed = len(pending) - len(kept)
            if removed > 0:
                log.info(f"  🚫 제외 리스트 적용: {removed}건 사전 차감")
                pending = kept
        except Exception as _ex_err:
            log.warning(f"  excluded_titles 사전 차감 실패(무시): {_ex_err}")

    # ── 정렬 결과 캐시 조회 (2026-09-02) — 홈 탭 '발행 준비' 신호등 지연(≈19초) 단축.
    # 이 함수는 실발행 경로(get_pending_posts)도 타므로 히트 결과는 캐시 없이
    # 전체 재계산한 결과와 반드시 동일해야 한다. 예외 시 캐시가 없었던 것처럼
    # 아래 기존 전체 재계산 경로로 그대로 진행(fail-open).
    try:
        from core import queue_sort_cache as _qsc
        _cache_hit = _qsc.try_get(pending)
    except Exception:  # noqa: BLE001
        _qsc, _cache_hit = None, None
    if _cache_hit is not None:
        log.info(f"  ⚡ 정렬 캐시 사용 ({len(pending)}건, {_cache_hit['age_s']:.0f}초 전 계산)")
        try:
            atomic_write_json(str(_DATA_DIR / "logs" / "queue_sort_rationale.json"),
                               _cache_hit["rationale"])
        except Exception as _re:  # noqa: BLE001
            log.debug(f"  정렬 근거 저장 실패(무시): {_re}")
        return _qsc.reorder(pending, _cache_hit["order"])

    trend_scores = _load_latest_trend_scores()
    deadline_days = _load_deadline_days()

    # 뉴스 선행지표(news_radar) — 스테일/없음이면 빈 dict = 가산 0 (fail-open).
    # 뉴스는 '가산점'만 — 실측 검색량 1순위 원칙 불변.
    try:
        from core import news_radar as _nr
        _news_map = _nr.signal_map()
    except Exception:  # noqa: BLE001
        _nr, _news_map = None, {}

    # 실측 검색량 인덱스 — 데이터랩 상대지수보다 절대량이라 1순위.
    try:
        from core import search_volume as _sv
        _has_volume = bool(_sv.volume_index())
    except Exception:  # noqa: BLE001 — 인덱스 실패는 데이터랩 폴백
        _sv, _has_volume = None, False

    # 근거 데이터가 하나도 없으면 정렬 무의미 — 정직하게 알리고 엑셀 순서 유지
    if not _has_volume and not trend_scores and not deadline_days:
        log.info("  📊 트렌드 정렬 스킵: 검색량·트렌드·마감 데이터 없음 → 엑셀 순서 유지")
        return pending

    scored = []
    for item in pending:
        kw = _extract_keyword_from_title(item["title"])

        # ① 실측 검색량: 제목에 앵커된 헤드 키워드의 월간검색수
        head_kw, volume = "", 0
        if _has_volume:
            try:
                head_kw, volume = _sv.best_keyword(kw)
            except Exception:  # noqa: BLE001
                head_kw, volume = "", 0

        if volume > 0:
            t_score = _volume_score(volume)
            src = "volume"
        else:
            # ② 데이터랩 폴백 (부분 매칭). 0.0 은 '조회했지만 검색량 없음'이므로
            #    매칭에서 제외 — 제목 전체를 넣던 시절의 0.0 이 점수처럼 붙는 것 방지.
            t_score = trend_scores.get(kw) or None
            if t_score is None:
                for tk, tv in trend_scores.items():
                    if tv > 0 and (kw in tk or tk in kw):
                        t_score = tv
                        break
            src = "datalab" if t_score else "none"
            if t_score is None:
                # ③ 중립 기본값 — 실측 500회(≈42점)가 '수요 미지'보다 앞서도록 42 미만
                t_score = 40.0

        # 마감 긴급도 계산
        d_left = deadline_days.get(kw, None)
        if d_left is None:
            for dk, dv in deadline_days.items():
                if kw in dk or dk in kw:
                    d_left = dv
                    break
        if d_left is None:
            d_left = 999
        d_urgency = _calc_deadline_urgency(d_left)

        # 복합 점수 + 뉴스 선행지표 가산(스파이크 +8 · 변경 보도 +6, 캡 10)
        composite = (t_score * 0.6) + (d_urgency * 0.4)
        n_boost = 0.0
        if _news_map and _nr is not None:
            try:
                n_boost = _nr.boost_for(head_kw or kw, _news_map)
            except Exception:  # noqa: BLE001
                n_boost = 0.0
        composite += n_boost

        scored.append({
            **item,
            "_keyword":      head_kw or kw,
            "_volume":       volume,
            "_source":       src,
            "_trend_score":  t_score,
            "_deadline_days": d_left,
            "_urgency":      d_urgency,
            "_news":         n_boost,
            "_composite":    composite,
        })

    # 복합 점수 높은 순 정렬. 동점(10만회+ 는 모두 100점 캡)은 실측 검색량 큰 순,
    # 그것도 같으면 stable sort 라 엑셀 순서 보존.
    scored.sort(key=lambda x: (x["_composite"], x["_volume"]), reverse=True)

    # 같은 헤드 키워드 반복 감쇠 — 고수요 제도 하나(예: 문화누리카드 141k)가 하루
    # 큐 상단을 독식하면 하류 duplicate_checker 가 어차피 차단해 생성만 낭비된다.
    # 반복 등장마다 6점씩 감쇠해 다른 수요 키워드와 인터리브한다(결정적).
    _seen: dict = {}
    for s in scored:
        k = s["_keyword"]
        if s["_source"] == "volume":     # 실측 매칭 항목만 — 중립(제목 키)은 반복 아님
            s["_composite"] -= 6.0 * _seen.get(k, 0)
            _seen[k] = _seen.get(k, 0) + 1
    scored.sort(key=lambda x: (x["_composite"], x["_volume"]), reverse=True)

    # 로그: 전 항목 수백 줄 대신 근거 요약 1줄 + 상위 10건
    n_vol = sum(1 for s in scored if s["_source"] == "volume")
    n_trd = sum(1 for s in scored if s["_source"] == "datalab")
    n_dl = sum(1 for s in scored if s["_deadline_days"] < 999)
    n_news = sum(1 for s in scored if s.get("_news", 0) > 0)
    if n_vol == 0 and n_trd == 0 and n_dl == 0 and n_news == 0:
        log.info(f"  📊 트렌드 정렬: {len(scored)}건 모두 근거 매칭 실패(전부 중립 40점) "
                 "→ 사실상 엑셀 순서 유지")
    else:
        log.info(f"  📊 트렌드 정렬: {len(scored)}건 — 실측 검색량 {n_vol} · "
                 f"데이터랩 {n_trd} · 마감 {n_dl} · 뉴스 {n_news}건 매칭")
        for s in scored[:10]:
            d_str = (f"D-{s['_deadline_days']}" if s["_deadline_days"] < 999
                     else "마감없음")
            if s["_source"] == "volume":
                v_str = f"월 {s['_volume']:,}회"
            elif s["_source"] == "datalab":
                v_str = f"데이터랩 {s['_trend_score']:.0f}"
            else:
                v_str = "근거없음·중립"
            news_s = f" · 📰+{s['_news']:.0f}" if s.get("_news", 0) > 0 else ""
            log.info(f"    {s['_composite']:5.1f}점  {s['_keyword'][:24]}"
                     f"  ({v_str} · {d_str}{news_s})")
        if len(scored) > 10:
            log.info(f"    … 외 {len(scored) - 10}건")

    # 정렬 근거를 홈 탭이 읽도록 저장 — 실패해도 정렬 결과는 유효(fail-open)
    _rationale_payload = None
    try:
        _rationale_payload = {
            "date": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "total": len(scored),
            "matched_volume": n_vol,
            "matched_datalab": n_trd,
            "matched_deadline": n_dl,
            "matched_news": n_news,
            "top": [{
                "title": s.get("title", ""),
                "keyword": s["_keyword"],
                "volume": s["_volume"],
                "source": s["_source"],
                "score": round(s["_trend_score"], 1),
                "deadline_days": (s["_deadline_days"]
                                  if s["_deadline_days"] < 999 else None),
                "news_boost": round(s.get("_news", 0.0), 1),
                "composite": round(s["_composite"], 1),
            } for s in scored[:10]],
        }
        atomic_write_json(str(_DATA_DIR / "logs" / "queue_sort_rationale.json"),
                           _rationale_payload)
    except Exception as _re:  # noqa: BLE001
        log.debug(f"  정렬 근거 저장 실패(무시): {_re}")

    # 내부 점수 필드 제거 후 반환
    for s in scored:
        for f in ("_keyword", "_volume", "_source", "_trend_score",
                  "_deadline_days", "_urgency", "_news", "_composite"):
            s.pop(f, None)

    # 정렬 결과 캐시 저장 (2026-09-02) — 실패해도 무시, 다음 호출은 재계산(fail-open).
    if _rationale_payload is not None:
        try:
            from core import queue_sort_cache as _qsc2
            _qsc2.store(pending, [s.get("title", "") for s in scored], _rationale_payload)
        except Exception:  # noqa: BLE001
            pass

    return scored


def get_pending_posts(target_date: str = None, cfg=None, limit: int = None) -> list:
    """미완료 키워드를 트렌드·마감 우선순위 정렬해 반환.

    limit: None 이면 settings.DAILY_POST_LIMIT 적용 (examplebrand 운영 기본).
           양의 정수면 그 값으로 cap. save_only 보관함 채우기처럼 큰 배치가 필요할 때
           명시적으로 풀어줄 수 있다.
    """
    s = cfg or _cfg()
    if not target_date:
        target_date = date.today().strftime("%Y-%m-%d")

    excel = _keyword_excel(s)
    if not os.path.exists(excel):
        log.error(f"❌ 엑셀 파일 없음: {excel}")
        return []

    # 지난 실행에서 엑셀 잠금으로 보류된 완료 표시를 먼저 반영하고,
    # 여전히 잠겨 남은 항목은 이번 큐에서 완료로 간주(중복 발행 차단).
    apply_pending_marks(cfg=s)
    _held = {(_norm_xl(m.get("excel_path")), m.get("row"), m.get("col"))
             for m in _load_pending_marks()}
    _excel_key = _norm_xl(excel)

    wb = load_workbook(excel)
    ws = wb[s.KEYWORD_PUBLISH_POOL_SHEET]
    pending = []
    for row in ws.iter_rows():
        for cell in row:
            val = str(cell.value or "").strip()
            # 날짜 형식(YYYY-MM-DD — 제목)이 아니면 건너뜀
            if not val or len(val) < 12 or val[4:5] != "-":
                continue
            try:
                rgb  = cell.fill.fgColor.rgb if cell.fill and cell.fill.fgColor else ""
                done = "FFFF00" in rgb or rgb == "FFFFFF00"
            except (AttributeError, TypeError):
                done = False
            if not done and _held and (_excel_key, cell.row, cell.column) in _held:
                done = True   # 발행은 이미 성공 — 표시만 잠금으로 보류 중
            if not done:
                pending.append({
                    "title":      val,
                    "row":        cell.row,
                    "col":        cell.column,
                    "excel_path": excel,
                })

    # 트렌드 스코어 + 마감 긴급도 기반 우선순위 정렬
    pending = _sort_by_trend_priority(pending)

    cap = limit if (limit is not None and limit > 0) else s.DAILY_POST_LIMIT
    count = min(len(pending), cap)
    log.info(f"📋 [{target_date}] 미완료: {len(pending)}개 → 처리: {count}개 (limit={cap})")
    return pending[:count]


def _automatic_topic_editorial(cfg=None) -> dict:
    """Expose only supported explicit topic fields from the current settings object."""
    return {key: getattr(cfg, key, None) for key in
            ("TOPIC_MODE", "BLOG_GENRE", "BRAND_GENRE", "topic_mode", "brand_genre")}


def _automatic_topic_configured(cfg=None, profile: Optional[str] = None) -> bool:
    """Use one side-effect-free predicate for execution and preflight queue routing."""
    from core import auto_topics
    profile_name = profile or _nbaprofile.active_profile()
    return auto_topics.is_auto_configured(profile_name, _automatic_topic_editorial(cfg))


def _automatic_topic_posts(cfg=None, *, refresh: bool, limit: int = 1,
                           profile: Optional[str] = None):
    """Return auto candidates when configured, or None to preserve the Excel queue."""
    profile_name = profile or _nbaprofile.active_profile()
    editorial = _automatic_topic_editorial(cfg)
    if not _automatic_topic_configured(cfg, profile_name):
        return None
    from core import auto_topics
    candidates = auto_topics.select_candidates(profile_name, editorial=editorial,
        refresh=refresh, limit=max(1, int(limit)))
    if _nbaprofile.active_profile() != profile_name:
        log.warning("자동 주제 선정 중 프로필 변경 감지 — 후보를 폐기합니다")
        return []
    return candidates


# ── 보관함 우선 소진 (archive-first drain, 2026-07-13) ──────────
# 발행 큐(트렌드 정렬 상위 N편)를 신규 생성하기 전에, 보관함(draft_ready elite)을
# 먼저 소진하도록 드레인 리스트를 큐 앞에 붙인다. 프리펜드된 드레인 제목은 메인 루프의
# find_draft_for_title 경로에 걸려 신규 생성 없이 보관본 그대로 발행 준비된다.
def _drain_image_ready_rank(entry: dict, post_map: dict) -> int:
    """드레인 후보 하나의 이미지 완성 순위 키(완성=0, 미완/판정불가=1)를 낸다.

    core.draft_image_fill.images_complete 단일 진실 원본을 재사용한다(이미지 경로
    문자열 유무가 아니라 실파일·계획 완결 여부까지 본다). 슬러그를 못 찾거나 조회·
    판정 중 예외가 나면 **완성으로 오승격하지 않도록** 미완(1)으로 안전하게 취급한다.
    """
    slug = entry.get("slug") or ""
    hit = post_map.get(slug) if slug else None
    if not hit:
        return 1
    try:
        complete, _detail = draft_image_fill.images_complete(hit.get("post_data") or {})
    except Exception:    # noqa: BLE001 — 개별 판정 실패는 그 항목만 뒤로(미완 취급)
        return 1
    return 0 if complete else 1


def _prioritize_image_ready(candidates: list, post_map: dict = None) -> list:
    """드레인 후보를 '이미지 완성 글 우선'으로 안정 정렬한다(cap 적용은 호출자 몫).

    동일 완성도 안에서는 candidates 가 들어온 순서(기존 수요·품질 정렬)를 그대로
    유지한다(Python sorted() 는 stable). draft_store 조회는 기본적으로 여기서 딱
    1회(load_drafts)만 하고 slug→entry map 으로 모든 후보에 재사용한다 — 후보별
    디스크 재탐색 없음. map 구축 자체가 실패하면 원래 순서를 그대로 반환한다
    (fail-open — 전체 조회 실패가 드레인을 막지 않음).

    Args:
        post_map: 호출자가 이미 조회해 둔 slug→entry map(2026-09-21, `_apply_archive_drain`
            이 드레인·fresh 판정에 같은 map 을 재사용할 때 넘긴다). None 이면 여기서
            직접 조회한다(기존 동작 그대로, 단독 호출·테스트 호환).
    """
    if not candidates:
        return candidates
    if post_map is None:
        try:
            post_map = {
                d["slug"]: d for d in draft_store.load_drafts(status_filter="draft_ready")
                if isinstance(d, dict) and d.get("slug")
            }
        except Exception as exc:  # noqa: BLE001 — 조회 실패는 원 순서 유지
            log.warning(f"  ⚠️  드레인 이미지 우선순위 조회 실패(원 순서 유지): {exc}")
            return candidates
    return sorted(candidates, key=lambda c: _drain_image_ready_rank(c, post_map))


def _archive_drain_candidates(cap, cfg=None, *, allowed_cats=None,
                              exclude_keys=None, plan=None, post_map=None) -> list:
    """보관함 elite 를 발행 큐 앞에 채울 드레인 후보 dict 리스트를 만든다(최대 cap 개).

    build_draft_backlog_plan(순수·엑셀 미변경)이 낸 발행 버킷 제목 중,
      · 이미 fresh 큐에 있는 키(exclude_keys)와 겹치지 않고
      · allowed_cats 필터가 있으면 그 카테고리에 속하는 제목만
    취한 뒤(원래 수요·품질 순서 보존), **이미지 완성 글을 앞으로** 안정 정렬하고
    (2026-09-21) 그 다음에 cap 을 적용한다 — 완성 글이 캡에 밀려 사라지지 않게.
    각 항목은 fresh(get_pending_posts) dict 와 같은 인터페이스(title)를 갖되,
    엑셀 셀이 없으므로 _drain=True 로 표시한다(발행 성공 시 mark_done 스킵).

    plan: 주입 가능(테스트용). None 이면 build_draft_backlog_plan 라이브 로드.
    실패·빈 계획은 [] 반환(fail-open — 발행 흐름 무중단).
    """
    if not cap or cap <= 0:
        return []
    try:
        from core.performance_db import normalize_key
    except ImportError:
        return []
    if plan is None:
        try:
            from core.queue_injector import build_draft_backlog_plan
            plan = build_draft_backlog_plan().get("plan", [])
        except Exception as exc:  # noqa: BLE001 — 드레인 계획 실패가 발행을 막지 않음
            log.warning(f"  ⚠️  보관함 드레인 계획 실패(무시): {exc}")
            return []
    exclude = set(exclude_keys or ())
    seen: set = set()
    filtered: list = []
    for item in plan:
        title = (item.get("title") or "").strip()
        if not title:
            continue
        key = normalize_key(title)
        if not key or key in exclude or key in seen:
            continue
        if allowed_cats:
            _clean = re.sub(r'^\d{4}-\d{2}-\d{2}\s*[—\-]\s*', '', title)
            if detect_category(_clean) not in allowed_cats:
                continue
        seen.add(key)
        filtered.append({
            "title": title,
            "row": None, "col": None, "excel_path": None,
            "_drain": True,
            "slug": item.get("slug", ""),
            "category": item.get("category", ""),
            "score": item.get("score", 0),
        })
    filtered = _prioritize_image_ready(filtered, post_map=post_map)
    return filtered[:cap]


def select_only_titles(only_titles, cfg=None, *, target_date: str = None) -> list:
    """명시 지정된 제목만 후보로 만드는 **제한 경로** (2026-09-18).

    "하루 편수 1편"은 *몇 편인가*만 정하지 *어느 글인가*는 정하지 않는다. 파일럿처럼
    지정한 한 편만 준비·검증하려면 후보 자체를 그 제목으로 못박아야 한다.

    ① 발행 풀(엑셀)에 있으면 그 셀 정보를 그대로 쓴다(발행 후 완료 표시 정상 동작).
    ② 엑셀에 없고 보관본에 있으면 드레인과 같은 모양(`_drain=True`)으로 만든다.
    ③ 어느 쪽에도 없는 제목은 **조용히 넘기지 않고** 제외 목록으로 로그에 남긴다.

    반환 순서는 `only_titles` 가 준 순서다(호출자가 최대 편수를 자른다).
    """
    wanted = [str(t).strip() for t in (only_titles or []) if str(t or "").strip()]
    if not wanted:
        return []
    s = cfg or _cfg()
    try:
        from core.performance_db import normalize_key
    except ImportError:
        def normalize_key(text):  # type: ignore[misc]
            return str(text or "").strip()

    pool: dict = {}
    try:
        for item in get_pending_posts(target_date, cfg=s, limit=10 ** 6):
            key = normalize_key(item.get("title", ""))
            if key and key not in pool:
                pool[key] = item
    except Exception as exc:  # noqa: BLE001 — 엑셀을 못 읽어도 보관본 경로는 살린다
        log.warning(f"  ⚠️  지정 후보 조회 중 발행 풀 읽기 실패(보관본만 확인): {exc}")

    out, missing = [], []
    for title in wanted:
        key = normalize_key(title)
        hit = pool.get(key)
        if hit is not None:
            out.append(dict(hit))
            continue
        entry = None
        try:
            from core import draft_store as _ds
            entry = _ds.find_draft_for_title(title)
            if entry is None:
                # 슬러그를 그대로 준 경우도 받는다 — 보관본의 발행 큐 제목은
                # 슬러그 키(대시→공백)라서 사람이 제목으로 지목하기 어렵다.
                for cand in _ds.load_drafts(status_filter="draft_ready") or []:
                    if isinstance(cand, dict) and cand.get("slug") == title:
                        entry = cand
                        break
        except Exception as exc:  # noqa: BLE001
            log.warning(f"  ⚠️  보관본 조회 실패({type(exc).__name__}) — 지정 제목 제외")
        if entry is None:
            missing.append(title)
            continue
        post = entry.get("post_data") or {}
        # 루프의 `find_draft_for_title` 이 다시 찾을 수 있는 형태(= 드레인과 동일)로
        # 제목을 맞춘다: 슬러그 키의 대시를 공백으로.
        queue_title = title
        try:
            from core import draft_store as _ds2
            key = _ds2._key_from_slug(entry.get("slug") or "") or ""
            queue_title = key.replace("-", " ").strip() or title
        except Exception:  # noqa: BLE001
            pass
        out.append({"title": queue_title, "row": None, "col": None,
                    "excel_path": None, "_drain": True,
                    "slug": entry.get("slug", ""),
                    "category": post.get("category", ""), "score": 0})
    if missing:
        log.warning(f"  ⚠️  지정 제목을 발행 풀·보관함 어디에서도 찾지 못했습니다: {missing}")
    log.info(f"  🎯 지정 후보 제한 경로: 요청 {len(wanted)}건 → 후보 {len(out)}건")
    return out


def _apply_archive_drain(posts: list, cap, cfg=None, *,
                         allowed_cats=None, plan=None) -> tuple:
    """보관함 elite 를 posts 앞에 prepend 하고 총합이 cap 을 넘지 않게 fresh 를 트림한다.

    Returns:
        (new_posts, n_drain, n_fresh_kept). settings.DRAFT_DRAIN_ENABLED(기본 True)가
        꺼져 있거나 cap<=0 이거나 드레인 후보가 없으면 posts 를 그대로 돌려준다.
        어떤 예외도 posts 원본을 반환(fail-open) — 발행 흐름 무중단.
    """
    s = cfg or _cfg()
    try:
        if not getattr(s, "DRAFT_DRAIN_ENABLED", True):
            return posts, 0, len(posts)
        if not cap or cap <= 0:
            return posts, 0, len(posts)
        from core.performance_db import normalize_key
        fresh_keys = set()
        for p in posts:
            k = normalize_key(p.get("title", ""))
            if k:
                fresh_keys.add(k)
        try:
            post_map = {
                d["slug"]: d for d in draft_store.load_drafts(status_filter="draft_ready")
                if isinstance(d, dict) and d.get("slug")
            }
        except Exception:
            post_map = {}
        drain = _archive_drain_candidates(
            cap, cfg=s, allowed_cats=allowed_cats,
            exclude_keys=fresh_keys, plan=plan, post_map=post_map)
        ready_ranks = {}
        matches = {}

        def priority(item):
            slug = item.get("slug") or ""
            hit = post_map.get(slug)
            if hit is None and not item.get("_drain"):
                title = item.get("title", "")
                if title not in matches:
                    # 기존 조회의 나이 제한·동일 키워드 최신본 선택을 그대로 적용.
                    matches[title] = draft_store.find_draft_for_title(title)
                hit = matches[title]
                slug = (hit or {}).get("slug", "")
            if not hit:
                return 1 if item.get("_drain") else 2
            if slug not in ready_ranks:
                ready_ranks[slug] = _drain_image_ready_rank(
                    {"slug": slug}, {slug: hit})
            return ready_ranks[slug]

        selected = sorted(list(drain) + list(posts), key=priority)[:cap]
        n_drain = sum(bool(p.get("_drain")) for p in selected)
        return selected, n_drain, len(selected) - n_drain
    except Exception as exc:  # noqa: BLE001 — 드레인은 부가기능, 실패해도 기존 흐름 유지
        log.warning(f"  ⚠️  보관함 우선 소진 스킵(예외): {exc}")
        return posts, 0, len(posts)


# ── fact_gate/QA 실패 본문 보관 (버리지 말고 needs_fix, 2026-07-13) ──
# 발행 차단(_failed) 글 중 품질/사실 문제로 막힌 것은 본문(≈2,000자)이 살아 있다.
# 그냥 continue 로 버리면 그 본문이 사라지므로, needs_fix 상태로 보관해 사람/후속
# 수정 대기용으로 남긴다. needs_fix 는 find_draft_for_title(draft_ready)과 드레인·
# triage 에서 자동 제외되므로 자동 발행되지 않는다.
_QUALITY_SAVE_REASONS = frozenset({
    "retries_exceeded", "qa_failed",
    # 2026-09-18: 새 글 발행 경로에서 본문/QA 는 통과했는데 계획 기반 이미지 채움이
    # 미완성으로 끝난 경우 — 품질 실패가 아니라 "다음 시도가 이어 만들면 되는" 보류다.
    "new_post_images_not_ready",
})
_FACT_GATE_MARKERS = ("허구수치", "수치모순", "무근거_YMYL", "[fact_gate]")
# needs_fix 로 보관할 최소 본문 길이 — 이보다 짧으면 placeholder/too_short 같은
# 깡통 본문으로 보고 저장하지 않는다(수정 가치가 없음).
_NEEDS_FIX_MIN_BODY = 500


def _is_fixable_quality_failure(reason, qa_failures) -> bool:
    """발행 차단 사유가 '수정하면 살릴 수 있는' 품질/사실 실패인지 판정.

    True: retries_exceeded·qa_failed, 또는 fact_gate 유래 실패(허구수치·수치모순·
          무근거_YMYL)가 실패 목록에 있는 경우.
    False: placeholder_text·too_short 등 깡통 사유(본문 자체가 못 쓸 것) — 저장 대상 아님.
    """
    if reason in _QUALITY_SAVE_REASONS:
        return True
    for f in (qa_failures or []):
        if isinstance(f, str) and any(m in f for m in _FACT_GATE_MARKERS):
            return True
    return False


def _save_failed_draft_for_fix(title, post_data, image_paths, category, reason,
                               *, slug=None, target_date=None,
                               schedule_at=None) -> bool:
    """발행 차단된 품질/사실 실패 본문을 draft_store 에 needs_fix 로 보관한다.

    깡통 사유(placeholder/too_short)나 너무 짧은 본문은 저장하지 않는다.
    실패는 fail-open(로그만) — 호출자의 기존 스킵 흐름을 절대 바꾸지 않는다.

    Returns: 저장하면 True, 저장 안 함/실패면 False.
    """
    try:
        qa_failures = post_data.get("_qa_failures") or []
        if not _is_fixable_quality_failure(reason, qa_failures):
            return False
        body_html = post_data.get("content_html") or post_data.get("content") or ""
        if len(body_html) < _NEEDS_FIX_MIN_BODY:
            return False   # 깡통 본문 — 저장 가치 없음

        # 실패 사유 요약 — fact_gate 유래 라인(어떤 수치가 문제였는지)을 우선한다.
        fact_lines = [f for f in qa_failures
                      if isinstance(f, str) and any(m in f for m in _FACT_GATE_MARKERS)]
        detail = " | ".join((fact_lines or [f for f in qa_failures if isinstance(f, str)])[:6])
        error_summary = f"{reason}: {detail}" if detail else str(reason)

        draft_slug = slug or _make_fallback_slug(title, target_date)
        from core.draft_store import save_draft
        draft_post_data = {
            "title":    post_data.get("title", title),
            "content":  body_html,
            "tags":     list(post_data.get("tags", []) or []),
            "category": category,
        }
        # 계획·슬롯을 함께 넘겨야 다음 시도(이미지 채움/재발행)가 기획을 다시 만들지
        # 않고 빠진 슬롯만 채운다(2026-09-18 — new_post_images_not_ready 보류 경로).
        if isinstance(post_data.get("image_plan"), dict):
            draft_post_data["image_plan"] = post_data["image_plan"]
        if isinstance(post_data.get("image_slots"), dict):
            draft_post_data["image_slots"] = post_data["image_slots"]
        save_draft(
            slug=draft_slug,
            post_data=draft_post_data,
            image_paths=list(image_paths or []),
            schedule_at=schedule_at,
            status="needs_fix",
            error=error_summary[:500],
        )
        log.info(
            f"  🗄  발행 차단 본문 보관(needs_fix): {draft_slug} "
            f"({len(body_html)}자 / 사유 {reason})"
        )
        return True
    except Exception as exc:  # noqa: BLE001 — 보관 실패가 기존 스킵 흐름을 깨면 안 됨
        log.warning(f"  ⚠️  needs_fix 보관 실패(무시): {type(exc).__name__}: {exc}")
        return False


# ── 엑셀 완료 표시 ────────────────────────────────────
# 발행 중 같은 엑셀을 읽는 내부 경합(GUI 홈 preflight·대시보드 스캔) 또는
# 백신 스캔·엑셀 열림이 겹치면 os.replace 가 WinError 5 로 실패한다.
# 발행(예약)은 이미 성공한 뒤라 여기서 예외를 올리면 성공 글이 실패로 기록되고
# 다음 실행에서 같은 글이 중복 발행된다(2026-07-19 exampleprofile 실측, 사용자는
# 엑셀을 열지 않았음 — 자기 경합) — 점증 재시도로 순간 경합을 흡수하고,
# 그래도 잠겨 있으면 저널(pending_excel_marks.json)에 보류 기록,
# 다음 실행의 get_pending_posts 가 보류분을 완료로 간주 + 재적용한다.
PENDING_MARKS_PATH = LOG_DIR / "pending_excel_marks.json"


def _norm_xl(p) -> str:
    return os.path.normcase(os.path.normpath(str(p)))


def _write_done_cell(excel_path: str, row: int, col: int, cfg=None):
    s  = cfg or _cfg()
    wb = load_workbook(excel_path)
    ws = wb[s.KEYWORD_PUBLISH_POOL_SHEET]
    ws.cell(row=row, column=col).fill = YELLOW
    tmp_path = excel_path + ".tmp"
    try:
        wb.save(tmp_path)
        os.replace(tmp_path, excel_path)
    except OSError:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except OSError:
                pass
        raise


def _load_pending_marks() -> list:
    marks = safe_load_json(PENDING_MARKS_PATH, default=list)
    return marks if isinstance(marks, list) else []


def apply_pending_marks(cfg=None) -> int:
    """이전 실행에서 엑셀 잠금으로 보류된 완료 표시를 재적용.

    성공 항목은 저널에서 제거. 여전히 잠겨 있으면 보류 유지 —
    get_pending_posts 가 보류 항목을 완료로 간주하므로 중복 발행은 없다.
    """
    marks = _load_pending_marks()
    if not marks:
        return 0
    remaining, applied = [], 0
    for m in marks:
        try:
            _write_done_cell(m["excel_path"], int(m["row"]), int(m["col"]), cfg=cfg)
            applied += 1
        except OSError:
            remaining.append(m)          # 잠금 지속 — 다음 기회에 재시도
        except (KeyError, ValueError, TypeError):
            pass                         # 손상 항목은 폐기
    atomic_write_json(str(PENDING_MARKS_PATH), remaining)
    if applied or remaining:
        log.info(f"  📌 보류된 엑셀 완료 표시 {applied}건 반영"
                 + (f" (잠금 지속 {len(remaining)}건 보류 유지)" if remaining else ""))
    return applied


def mark_done(excel_path: str, row: int, col: int, cfg=None) -> bool:
    """엑셀 셀 완료(노랑) 표시. 잠금이면 점증 재시도 → 저널 보류(예외 미전파)."""
    for attempt in range(4):
        try:
            _write_done_cell(excel_path, row, col, cfg=cfg)
            return True
        except OSError:
            if attempt < 3:
                time.sleep(2 * (attempt + 1))   # 2·4·6초 — 읽기 경합·AV 스캔 흡수
    marks = _load_pending_marks()
    entry = {"excel_path": str(excel_path), "row": row, "col": col}
    if entry not in marks:
        marks.append(entry)
    atomic_write_json(str(PENDING_MARKS_PATH), marks)
    log.warning(
        f"  ⚠️  엑셀이 잠겨 완료 표시 보류(row={row}, col={col}) — 저널 기록, "
        f"다음 실행 시 자동 반영 (원인: 내부 조회 경합·백신 스캔·파일 열림 등)"
    )
    return False


# ── JSON 로그 ─────────────────────────────────────────
def write_log(title: str, success: bool,
              reserve_time: str = "", detail: str = ""):
    log_file = LOG_DIR / f"{date.today().strftime('%Y-%m-%d')}.json"
    logs = safe_load_json(log_file, default=list)
    logs.append({
        "time":         datetime.now().strftime("%H:%M:%S"),
        "title":        title,
        "success":      success,
        "reserve_time": reserve_time,
        "detail":       detail,
    })
    atomic_write_json(log_file, logs)


# ── 실패 포스트 재시도 ────────────────────────────────
def retry_failed_post(post_title: str, stop_flag=None,
                      progress_cb=None) -> dict:
    """failed_posts.json에서 해당 제목의 포스트를 찾아 재시도합니다.

    Args:
        post_title: 재시도할 포스트 제목
        stop_flag: threading.Event — 외부 중지 신호
        progress_cb: GUI 콜백 (title, status)

    Returns:
        {"success": bool, "title": str, "error": str}
    """
    failed = safe_load_json(str(FAILED_POSTS_PATH), default=list)

    # 제목이 일치하는 항목 찾기
    target = None
    target_idx = -1
    for idx, item in enumerate(failed):
        if item.get("title", "") == post_title:
            target = item
            target_idx = idx
            break

    if target is None:
        return {"success": False, "title": post_title,
                "error": "failed_posts.json에서 해당 제목을 찾을 수 없습니다"}

    if progress_cb:
        progress_cb(post_title, "retry_start")

    try:
        # 글 재생성
        post_data = generate_blog_post(post_title, stop_flag=stop_flag)

        from core.quality_validator import validate_and_fix
        post_data = validate_and_fix(post_data, post_title, stop_flag=stop_flag)

        # ── 품질 게이트 (2026-07-03) ──────────────────────────
        # validate_and_fix 직후 품질/중지 플래그를 확인한다. 통과하지
        # 못하면 pop/record 하지 않고 실패로 반환해 failed 목록에 남겨
        # 다음 재시도 대상으로 유지한다. (_qa_failed 는 별도 작업에서
        # 추가 예정이나 .get() 이라 없어도 안전.)
        if (post_data.get("_failed") or post_data.get("_qa_failed")
                or post_data.get("_stopped")):
            log.warning(
                f"  ⚠️  재시도 품질 게이트 차단: {post_title[:46]} — "
                f"failed={bool(post_data.get('_failed'))} "
                f"qa_failed={bool(post_data.get('_qa_failed'))} "
                f"stopped={bool(post_data.get('_stopped'))}"
            )
            write_log(post_title, False, "", "retry_quality_gate_blocked")
            if progress_cb:
                progress_cb(post_title, "retry_fail")
            return {"success": False, "title": post_title,
                    "error": "quality_gate 또는 stopped"}

        post_data["content_html"] = markdown_to_naver_html(
            post_data["content"], post_title, post_data.get("category", ""))

        # 성공 시 failed_posts.json에서 해당 항목 제거
        failed.pop(target_idx)
        atomic_write_json(str(FAILED_POSTS_PATH), failed)

        # 이력 기록 (본문 키워드 포함 → 본문 중복 검사에 활용)
        record_post(post_title, "", content=post_data.get("content", ""))
        write_log(post_title, True, "", "retry_success")

        if progress_cb:
            progress_cb(post_title, "retry_ok")

        log.info(f"  ✅ 재시도 성공: {post_title[:46]}")
        return {"success": True, "title": post_title, "error": ""}

    except (ValueError, OSError, RuntimeError) as e:
        log.error(f"  ❌ 재시도 실패: {post_title[:46]} — {e}", exc_info=True)
        write_log(post_title, False, "", f"retry_failed: {e}")

        if progress_cb:
            progress_cb(post_title, "retry_fail")

        return {"success": False, "title": post_title, "error": str(e)}


# ── 엑셀 백업 ────────────────────────────────────────
def _backup_excel(excel_path: str, max_backups: int = 7):
    """엑셀 파일을 logs/backup/에 날짜별로 백업합니다.

    - 이미 오늘 날짜 백업이 있으면 건너뜁니다.
    - 백업이 max_backups개 이상이면 가장 오래된 것을 삭제합니다.
    """
    src = Path(excel_path)
    if not src.is_file():
        return

    backup_dir = LOG_DIR / "backup"
    backup_dir.mkdir(parents=True, exist_ok=True)

    today_str = date.today().strftime("%Y%m%d")
    ext = src.suffix  # .xlsx
    dest = backup_dir / f"keywords_{today_str}{ext}"

    # 오늘 이미 백업했으면 건너뜀
    if dest.exists():
        log.debug(f"오늘 백업 이미 존재: {dest.name}")
        return

    shutil.copy2(str(src), str(dest))
    log.info(f"  📦 엑셀 백업 완료: {dest.name}")

    # 회전: 오래된 백업 삭제
    backups = sorted(backup_dir.glob(f"keywords_*{ext}"),
                     key=lambda f: f.stat().st_mtime)
    while len(backups) > max_backups:
        oldest = backups.pop(0)
        oldest.unlink()
        log.info(f"  🗑️  오래된 백업 삭제: {oldest.name}")


# ── 실패 글 자동 재시도 ──────────────────────────────
def _auto_retry_failed(max_retry: int = 3, stop_flag=None, progress_cb=None):
    """failed_posts.json에서 retry_count가 max_retry 미만인 항목을 자동으로 재시도합니다.

    - 재시도 성공 시 해당 항목을 failed_posts.json에서 제거합니다.
    - 재시도 실패 시 retry_count를 증가시킵니다.
    """
    failed = safe_load_json(str(FAILED_POSTS_PATH), default=list)
    if not failed:
        return

    retryable = [
        (idx, item) for idx, item in enumerate(failed)
        if item.get("retry_count", 0) < max_retry
    ]
    if not retryable:
        return

    log.info(f"\n♻️  실패 글 자동 재시도: {len(retryable)}건")
    removed_indices = []

    for idx, item in retryable:
        if stop_flag and stop_flag.is_set():
            log.info("  ⏹ 재시도 중지 (사용자 요청)")
            break

        title = item.get("title", "")
        retry_count = item.get("retry_count", 0)
        log.info(f"  ♻️  재시도 ({retry_count+1}/{max_retry}): {title[:46]}")

        result = retry_failed_post(title, stop_flag=stop_flag,
                                   progress_cb=progress_cb)
        if result.get("success"):
            removed_indices.append(idx)
            log.info(f"  ✅ 재시도 성공: {title[:46]}")
        else:
            # retry_count 증가
            item["retry_count"] = retry_count + 1
            log.info(f"  ❌ 재시도 실패 ({item['retry_count']}/{max_retry}): "
                     f"{title[:46]}")

    # 루프 종료 후 처리
    # 1) 최신 파일 로드 (retry_failed_post가 성공 항목을 이미 제거했을 수 있음)
    latest = safe_load_json(str(FAILED_POSTS_PATH), default=list)
    # 2) 실패한 항목들의 retry_count를 latest에 반영
    for orig_idx, item in retryable:
        title = item.get("title", "")
        new_count = item.get("retry_count", 0)
        for rec in latest:
            if rec.get("title") == title:
                rec["retry_count"] = new_count
                break
    # 3) 저장
    atomic_write_json(str(FAILED_POSTS_PATH), latest)


# ── 이미지 dispatcher 통합 (2026-05-16) ───────────────────
# Day 1~3 에 만들어둔 codex_cli_bot + image_inbox + image_dispatcher 를
# 실제 발행 파이프라인에 끼우는 진입점. 한도(5h 30회)·Tier 2 폴백·인박스
# 매칭은 부품들이 자체적으로 처리하므로 여기서는 1) DispatcherFactory 초기화,
# 2) 본문에서 BodyRecipe/InfoRecipe 추출, 3) dispatch(...) 호출, 4) 성공한 PNG
# 경로를 image_paths 로 모아 반환만 수행한다.
#
# 호출 규약:
#   image_paths = _run_image_dispatcher(slug, title, category, post_data, ...)
#   - 실패 시 빈 리스트를 반환. 에이전트 흐름은 멈추지 않는다.
#   - settings.IMAGE_INJECT_ENABLED == False 면 즉시 빈 리스트.
#   - dispatcher_factory 가 주입되면 (테스트) 내부 import 를 우회.


def _archive_images_complete(post: dict, meta: dict) -> bool:
    """Return whether every planned image exists after a fill attempt."""
    paths = post.get("image_paths") or []
    if not paths or meta.get("skipped") or meta.get("slots_failed"):
        return False
    if draft_image_fill.plan_mode_enabled():
        # 계획 경로의 판정은 `draft_image_fill.images_complete` 가 단일 진실 원본이다
        # (재고 집계도 같은 함수를 쓴다 — 2026-09-18). 규칙은 이전과 동일:
        # "계획된 슬롯 전부에 실재 파일이 붙어 있는가".
        complete, _detail = draft_image_fill.images_complete(post)
        return complete
    expected = 1 + int(meta.get("cards_grounded") or 0)
    return len(paths) >= expected


_ARCHIVE_IMAGE_REASON_CODES = {
    "already_running", "bad_post_data", "budget_exhausted", "codex_not_found",
    "config_broken", "error", "genspark_disabled", "image_fill_disabled",
    "incomplete_images", "login_required", "no_article", "no_plan", "no_title",
    "no_verified_images", "ok", "outdir_failed", "rate_limited",
    "slots_unfilled", "stale_no_longer_ready", "unverified_assets",
    "unverifiable_assets", "reference_assets", "placeholder_assets",
    "tampered_assets", "undecodable_assets",
}


def _archive_image_reason(value: object, fallback: str = "incomplete_images") -> str:
    """Keep archive diagnostics to known, non-sensitive reason codes."""
    return value if isinstance(value, str) and value in _ARCHIVE_IMAGE_REASON_CODES else fallback


def _set_archive_image_result(entry: dict, reason: str, *, slots: int = 0,
                              unverified: int = 0, produced: int = 0,
                              reused: int = 0) -> None:
    """Attach count-only diagnostics to the transient archive entry."""
    entry["_image_prepare_result"] = {
        "reason": _archive_image_reason(reason),
        "slots": max(0, slots) if isinstance(slots, int) else 0,
        "unverified": max(0, unverified) if isinstance(unverified, int) else 0,
        "produced": max(0, produced) if isinstance(produced, int) else 0,
        "reused": max(0, reused) if isinstance(reused, int) else 0,
    }


def _fill_archive_images_locked(entry: dict, *, repair_legacy_images: bool = False) -> bool:
    """Fill and persist a ready archive entry under the profile image lock."""
    slug = entry["slug"]
    if not draft_image_fill._draft_still_fillable(draft_store, slug):
        _set_archive_image_result(entry, "stale_no_longer_ready")
        return False
    post = dict(entry.get("post_data") or {})
    replace_unverified = False
    if draft_image_fill.plan_mode_enabled():
        initially_complete, inventory = draft_image_fill.images_complete(post)
        unverified_count = inventory.get("unverified", 0)
        reasons = inventory.get("unverified_reasons")
        if not initially_complete and unverified_count:
            legacy_only = (
                isinstance(reasons, dict)
                and reasons.get("unverified_legacy") == unverified_count
                and sum(value for value in reasons.values()
                        if isinstance(value, int) and value >= 0) == unverified_count
            )
            if legacy_only and repair_legacy_images:
                current_paths = []
                findings = []
                raw_paths = list(post.get("image_paths") or [])
                slot_paths = post.get("image_slots")
                if isinstance(slot_paths, dict):
                    raw_paths.extend(slot_paths.values())
                for raw_path in raw_paths:
                    path = str(raw_path or "").strip()
                    if not path or path in current_paths:
                        continue
                    try:
                        if Path(path).is_file():
                            current_paths.append(path)
                    except (OSError, TypeError, ValueError):
                        reason = "unverifiable_assets"
                        break
                else:
                    reason = ""
                if not reason:
                    try:
                        from core import image_asset_guard
                        _publishable, findings, _paths = image_asset_guard.verify_publishable(
                            current_paths, title=post.get("title", ""),
                            category=post.get("category", ""))
                    except (ImportError, OSError, TypeError, ValueError, RuntimeError):
                        findings = []
                        reason = "unverifiable_assets"
                blocked = [finding for finding in findings
                           if finding.get("blocking")]
                nonlegacy = [finding for finding in blocked
                             if finding.get("verdict") != "unverified_legacy"]
                reason = ""
                if nonlegacy:
                    reason_map = {
                        "unverifiable": "unverifiable_assets",
                        "reference_capture": "reference_assets",
                        "placeholder": "placeholder_assets",
                        "provenance_mismatch": "tampered_assets",
                        "undecodable": "undecodable_assets",
                    }
                    reason = reason_map.get(nonlegacy[0].get("verdict"),
                                            "unverifiable_assets")
                elif not blocked:
                    # The fast inventory and content-bound check disagreed. Do not
                    # replace an asset whose current bytes are not proven legacy.
                    reason = "unverifiable_assets"
                else:
                    from core import image_plan
                    mapped, _missing, _extras = image_plan.slot_assets(
                        post.get("image_plan") or {}, post)
                    mapped_paths = set(mapped.values())
                    if any(finding.get("path") not in mapped_paths
                           for finding in blocked):
                        # fill_images_for_draft only replaces plan-slot assets.
                        # Keep unassigned legacy files for explicit user review.
                        reason = "unverified_assets"
                if reason:
                    _set_archive_image_result(
                        entry, reason, slots=inventory.get("slots", 0),
                        unverified=unverified_count)
                    draft_image_fill._persist_fill_state(
                        slug, ok=False, reason=reason, produced=0, reused=0)
                    return False
                replace_unverified = True
            else:
                if not legacy_only:
                    reason_map = {
                        "unverifiable": "unverifiable_assets",
                        "reference_capture": "reference_assets",
                        "placeholder": "placeholder_assets",
                        "provenance_mismatch": "tampered_assets",
                        "undecodable": "undecodable_assets",
                    }
                    blocked_reasons = [key for key in (reasons or {})
                                       if key != "unverified_legacy"]
                    reason = (reason_map.get(blocked_reasons[0], "unverified_assets")
                              if blocked_reasons else "unverifiable_assets")
                else:
                    reason = "unverified_assets"
            if not (legacy_only and repair_legacy_images):
                _set_archive_image_result(
                    entry, reason, slots=inventory.get("slots", 0),
                    unverified=unverified_count)
                draft_image_fill._persist_fill_state(
                    slug, ok=False, reason=reason, produced=0, reused=0)
                return False
    result = draft_image_fill.fill_images_for_draft(
        post, slug=slug, category=post.get("category"), reuse=True,
        replace_unverified=replace_unverified)
    meta = result.get("image_meta") or {}
    candidates = list(result.get("image_paths") or []) + list(post.get("image_paths") or [])
    post["image_paths"] = list(dict.fromkeys(str(p) for p in candidates if Path(p).is_file()))
    complete = _archive_images_complete(post, meta)
    if complete:
        reason = "ok"
    else:
        _complete, inventory = draft_image_fill.images_complete(post)
        reason = _archive_image_reason(
            meta.get("skipped") or inventory.get("reason"), "incomplete_images")
    if not draft_image_fill._draft_still_fillable(draft_store, slug):
        _set_archive_image_result(entry, "stale_no_longer_ready",
                                  slots=inventory.get("slots", 0) if not complete else 0,
                                  unverified=inventory.get("unverified", 0) if not complete else 0,
                                  produced=meta.get("generated", 0), reused=meta.get("reused", 0))
        return False
    draft_store.save_draft(slug, post, post["image_paths"],
                           schedule_at=entry.get("schedule_at"),
                           error=None if complete else reason)
    entry["post_data"] = post
    _set_archive_image_result(
        entry, reason, slots=inventory.get("slots", 0) if not complete else 0,
        unverified=inventory.get("unverified", 0) if not complete else 0,
        produced=meta.get("generated", 0), reused=meta.get("reused", 0))
    draft_image_fill._persist_fill_state(
        slug, ok=complete, reason=reason, produced=meta.get("generated", 0),
        reused=meta.get("reused", 0))
    return complete


def _prepare_archive_images(entry: dict, *, repair_legacy_images: bool = False) -> bool:
    """Return publish readiness, preserving text-only behavior when fill is off."""
    if not draft_image_fill.image_fill_gate().get("enabled"):
        return True
    lock = draft_image_fill._ImageFillLock()
    # Batch locks allow same-PID reentry; publishing must not race a GUI fill thread.
    if not lock.acquire() or not lock._acquired:
        _set_archive_image_result(entry, "already_running")
        log.warning("이미지 채움이 이미 실행 중입니다 — 보관본 발행을 보류합니다")
        return False
    try:
        return _fill_archive_images_locked(entry,
                                           repair_legacy_images=repair_legacy_images)
    except (OSError, ValueError, TypeError, KeyError, RuntimeError):
        _set_archive_image_result(entry, "incomplete_images")
        log.error("보관본 이미지 준비/저장 실패 — 발행을 보류합니다", exc_info=True)
        return False
    finally:
        lock.release()


def _image_fill_plan_mode_ready() -> bool:
    """이미지 채움 게이트 ON + 계획 경로 ON 인가(save_only 와 무관한 공통 판정).

    ``_plan_fill_is_sole_image_producer``(보관함 채우기 전용)와
    ``_plan_fill_covers_new_post_publish``(새 글 발행 전용) 이 같은 조건을
    복제하지 않도록 여기 한 곳에 둔다. 판정 실패(설정 읽기 예외 등)는 False —
    즉 호출부가 예전 경로를 유지한다(fail-open).
    """
    try:
        from core import draft_image_fill
        if not draft_image_fill.image_fill_gate().get("enabled"):
            return False
        return bool(draft_image_fill.plan_mode_enabled())
    except Exception:  # noqa: BLE001 — 판정 실패가 파이프라인을 막지 않는다
        return False


def _plan_fill_is_sole_image_producer(save_only: bool) -> bool:
    """보관함 채우기에서 [HOOK G] 새 이미지 채움이 **유일한** 생성 주체인가.

    2026-09-10 마감 검토 실측 — save_only(보관함 채우기)에서 위쪽
    ``IMAGE_INJECT_ENABLED`` 분기가 구 ``_run_image_dispatcher`` 를 돌리고, 곧이어
    [HOOK G] 가 ``fill_images_for_draft`` 를 또 돌려 **같은 글에 이미지를 두 번**
    만들 수 있었다(라이브·예시 설정 모두 ``IMAGE_INJECT_ENABLED=True``). 구 경로는
    계획·슬롯을 모르므로 그 산출물은 슬롯 map 에 들어가지도 못한 채 시간·호출만 쓴다.

    그래서 **보관함 채우기 + 이미지 채움 ON + 계획 경로(기본)** 이 셋이 모두 참일
    때에만 구 dispatcher 를 건너뛰고 새 채움 하나로 생성한다. 발행 경로
    (``save_only=False``)·계획 경로 OFF(``IMAGE_FILL_PLAN_MODE=False``)·채움 OFF 는
    판정이 False 라 예전 동작 그대로다.

    Returns:
        True 면 호출부가 구 dispatcher 를 건너뛴다. 판정 자체가 실패하면(설정 읽기
        예외 등) False — 즉 예전 경로를 유지한다(fail-open).
    """
    if not save_only:
        return False
    return _image_fill_plan_mode_ready()


def _plan_fill_covers_new_post_publish() -> bool:
    """**새 글을 곧장 발행하는 경로**에서 계획 기반 채움이 이미지 완전성을 책임지는가.

    2026-09-18 마감 검토 실측 — 보관본 재사용(draft_hit) 발행은 ``_prepare_archive_images``
    가 계획·완전성을 확인하지만, 그 자리에서 막 집필해 곧장 업로드로 넘어가는 새 글은
    ``_run_image_dispatcher``/``generate_all_images`` 만 거치고 계획·슬롯을 몰랐다.
    그 결과 이미지가 0장이어도 "의도한 텍스트 전용"인지 "생성 실패"인지 구분 없이
    그대로 발행됐다(업로드 직전 어떤 완전성 검사도 없었다 — auto_scheduler 의
    "── 업로드 단계 ──" 진입 직전이 바로 그 지점).

    ``save_only`` 와 무관하게 **이미지 채움 게이트 ON + 계획 경로 ON** 일 때만 True.
    둘 중 하나라도 꺼져 있으면 False — 사용자가 의도적으로 끈 텍스트 전용 발행을
    그대로 존중한다(예전 경로 불변).
    """
    return _image_fill_plan_mode_ready()


def _fill_new_post_images(post_data: dict, *, slug: str, category: str,
                          existing_images: list) -> tuple:
    """새 글 발행 경로 — 계획대로 이미지를 채우고 완전성을 판정한다.

    ``_fill_archive_images_locked``(보관본 경로)와 **같은 API·같은 완전성 판정**
    (``fill_images_for_draft`` + ``_archive_images_complete``)을 쓴다. 이미 만들어진
    이미지(``existing_images`` — generate_all_images/_run_image_dispatcher 산출물)는
    ``reuse=True`` 로 넘겨 성공 슬롯을 재생성하지 않는다(사라진 경로만 다시 대상).

    ``post_data`` 는 **부분 성공이라도** image_plan/image_slots/_force_image_insert 를
    그 자리에서 갱신한다 — 완전성 미달로 발행이 보류돼도 호출자가 그대로
    ``_save_failed_draft_for_fix`` 에 넘기면 다음 시도가 빠진 슬롯만 만든다.

    Returns:
        (complete: bool, image_paths: list, image_meta: dict).
    """
    fill_src: dict = {
        "title": post_data.get("title", ""),
        "content": post_data.get("content", ""),
        "category": category,
    }
    if isinstance(post_data.get("image_plan"), dict):
        fill_src["image_plan"] = post_data["image_plan"]
    # 편집 브리프(2026-09-18)도 함께 넘긴다 — 계획이 없어 새로 만드는 경우에도 썸네일에
    # 얹을 문구가 **집필이 정한 메시지**가 되게(본문·썸네일이 같은 질문을 겨냥).
    if isinstance(post_data.get("editorial_brief"), dict):
        fill_src["editorial_brief"] = post_data["editorial_brief"]
    existing = [str(p) for p in (existing_images or []) if str(p or "").strip()]
    if existing:
        fill_src["image_paths"] = existing
    try:
        result = draft_image_fill.fill_images_for_draft(
            fill_src, slug=slug, category=category, reuse=True)
    except Exception:  # noqa: BLE001 — 채움 실패가 발행 파이프라인을 깨면 안 됨(fail-open)
        log.warning("새 글 이미지 채움 예외 — 부분 결과 없이 미완성으로 처리", exc_info=True)
        return False, existing, {"skipped": "fill_exception"}

    meta = result.get("image_meta") or {}
    images = list(result.get("image_paths") or existing)

    # 계획·슬롯은 성공/실패와 무관하게 승계 — 다음 시도가 빠진 슬롯만 만들게 한다.
    if isinstance(fill_src.get("image_plan"), dict):
        post_data["image_plan"] = fill_src["image_plan"]
    if isinstance(fill_src.get("image_slots"), dict):
        post_data["image_slots"] = fill_src["image_slots"]
    if fill_src.get("_force_image_insert"):
        post_data["_force_image_insert"] = True

    complete = _archive_images_complete(fill_src, meta)
    return complete, images, meta


def _run_image_dispatcher(
        slug: str,
        title: str,
        category: str,
        post_data: dict,
        *,
        runs_root=None,
        inbox_dir=None,
        dispatcher_factory=None,
        notifier=None) -> list:
    """본문에서 레시피를 뽑아 ImageDispatcher 를 호출하고 PNG 경로 목록을 반환.

    Args:
        slug: 발행 슬러그. runs/{slug}/draft.html 의 src 를 채우기 위한 키.
        title: 본문 제목 (한국어 OK).
        category: examplebrand 카테고리 — 디지털만 8 슬롯, 그 외 6 슬롯.
        post_data: content_generator 의 결과 dict. content_html / content / title.
        runs_root: runs/ 디렉토리 경로(테스트 주입용). None 이면 기본값.
        inbox_dir: 인박스 폴더 경로(테스트 주입용). None 이면 기본값.
        dispatcher_factory: ImageDispatcher 인스턴스 팩토리.
            None 이면 실제 ImageDispatcher + CodexCLIBot + ImageInbox 생성.
            테스트는 mock dispatcher 를 반환하는 callable 을 주입한다.
        notifier: GUI 알림 콜백 — Tier 2 폴백 시 메시지 전달.

    Returns:
        성공한 PNG 의 절대 경로 리스트. 한 장도 못 만들면 빈 리스트.
        이 결과는 호출자가 post_data["image_paths"] 로 저장해 draft_store /
        naver_uploader 에 전달한다.
    """
    s = _cfg()
    if not getattr(s, "IMAGE_INJECT_ENABLED", False):
        return []
    # 카테고리별 단계적 활성화 (2026-05-25) — Codex Plus 5h 30회 한도 보호.
    by_cat = getattr(s, "IMAGE_INJECT_BY_CATEGORY", None)
    if isinstance(by_cat, dict) and by_cat:
        if category and category in by_cat:
            if not bool(by_cat[category]):
                log.info(f"[image_dispatch] {category} 카테고리는 비활성 (IMAGE_INJECT_BY_CATEGORY)")
                return []
        elif not bool(by_cat.get("default", True)):
            log.info(f"[image_dispatch] {category!r} 카테고리는 기본 비활성")
            return []
    if not isinstance(slug, str) or not slug.strip():
        log.warning("[image_dispatch] slug 가 비어있음 — 이미지 생성 스킵")
        return []

    # 1) 레시피 추출 — recipe_extractor 가 본문 H2/리스트/표를 인식
    try:
        from core.recipe_extractor import extract_recipes_from_post
        body_recipes, info_recipes = extract_recipes_from_post(
            post_data, category,
            body_slot_count=2,
            info_slot_count=3,
        )
    except (ImportError, KeyError, AttributeError) as exc:
        log.warning(
            f"[image_dispatch] 레시피 추출 실패 ({type(exc).__name__}: {exc}) — "
            "이미지 생성 스킵"
        )
        return []

    # 2) dispatcher 인스턴스 — 테스트 주입이 있으면 그대로 사용
    try:
        if dispatcher_factory is not None:
            dispatcher = dispatcher_factory()
        else:
            # 실제 환경 — CodexCLIBot + ImageInbox + ImageDispatcher 조립
            # 이미지 체인은 반드시 naver_blog_v3.core.* (절대 패키지) 로 임포트한다.
            # image_dispatcher 내부가 naver_blog_v3.core.codex_cli_bot 을 참조하므로,
            # 여기서 core.* 로 만든 bot/inbox 를 넘기면 별개 모듈 정체성이 되어
            # ImageDispatcher.__init__ 의 isinstance(bot, CodexCLIBot) 검사가 TypeError 로
            # 실패한다. gui/image_panel·e2e·tests 와 동일하게 절대 임포트로 통일. (2026-07-03 버그픽스)
            # (위 recipe_extractor 임포트가 프로젝트 루트를 sys.path 에 이미 선삽입한 상태)
            from pathlib import Path as _Path
            from naver_blog_v3.core.codex_cli_bot import CodexCLIBot
            from naver_blog_v3.core.image_inbox import ImageInbox
            from naver_blog_v3.core.image_dispatcher import ImageDispatcher

            _runs_root = _Path(runs_root) if runs_root else _RESERVATIONS_DIR
            _inbox_dir = (_Path(inbox_dir) if inbox_dir
                          else (BASE_DIR / "image_inbox"))
            _runs_root.mkdir(parents=True, exist_ok=True)
            _inbox_dir.mkdir(parents=True, exist_ok=True)

            bot = CodexCLIBot(inbox_dir=_inbox_dir)
            # 로컬 ComfyUI 백엔드가 켜져 있으면 codex 대신 사용(동일 인터페이스).
            # ⚠️ 절대 임포트(naver_blog_v3.core.*) — isinstance(bot, CodexCLIBot) 통과 위해
            # 반드시 dispatcher 와 같은 모듈 정체성이어야 한다(이중 정체성 TypeError 방지).
            try:
                from naver_blog_v3.core import nvidia_image as _nvimg
                if _nvimg.available():
                    from naver_blog_v3.core.nvidia_image_bot import NvidiaImageBot
                    bot = NvidiaImageBot(inbox_dir=_inbox_dir)
                    log.info("[image_dispatch] 로컬 ComfyUI 이미지 백엔드 사용")
            except Exception as _nv_exc:  # noqa: BLE001 — 실패 시 codex 로 폴백
                log.debug(f"[image_dispatch] 로컬 ComfyUI 이미지 백엔드 미사용: {_nv_exc}")
            inbox = ImageInbox(inbox_dir=_inbox_dir, runs_root=_runs_root)
            dispatcher = ImageDispatcher(
                bot=bot, inbox=inbox, notifier=notifier or (lambda _m: None)
            )
    except (ImportError, OSError, TypeError) as exc:
        log.warning(
            f"[image_dispatch] dispatcher 초기화 실패 — "
            f"({type(exc).__name__}: {exc}) 이미지 생성 스킵"
        )
        return []

    # 3) dispatch 호출 — 한도 초과·codex 미설치는 dispatcher 가 Tier 2 폴백.
    subtitle = ""
    try:
        result = dispatcher.dispatch(
            slug=slug,
            title=title,
            subtitle=subtitle,
            category=category,
            body_recipes=body_recipes,
            info_recipes=info_recipes,
        )
    except (TypeError, ValueError) as exc:
        log.warning(
            f"[image_dispatch] dispatch 호출 시그너처 오류 — "
            f"({type(exc).__name__}: {exc})"
        )
        return []
    except Exception as exc:  # noqa: BLE001 — dispatcher 자체 예외는 메인 흐름 안 막음
        log.warning(
            f"[image_dispatch] dispatch 실행 실패 — "
            f"({type(exc).__name__}: {exc})"
        )
        return []

    # 4) 결과 로깅 + image_paths 수집
    tier = getattr(result, "tier_used", "unknown")
    gen = getattr(result, "generation", None)
    paths: list = []
    if gen is not None and getattr(gen, "success_paths", None):
        paths = [str(p) for p in gen.success_paths if p]

    matched_count = 0
    expected_slots = 6
    reconcile = getattr(result, "reconcile", None)
    if reconcile is not None:
        matched_count = getattr(reconcile, "matched_count", 0) or 0
        expected_slots = getattr(reconcile, "expected_slots", 6) or 6

    log.info(
        f"  🖼  dispatcher: {tier} {matched_count}/{expected_slots} PASS "
        f"(success_paths={len(paths)})"
    )
    fallback_reason = getattr(result, "fallback_reason", None)
    if fallback_reason:
        log.info(f"     ↳ 폴백 사유: {fallback_reason}")

    return paths


# ── 메인 파이프라인 ───────────────────────────────────
def _make_fallback_slug(title: str, target_date: str = None) -> str:
    """harness_emitter import 가 실패한 환경에서 쓰는 안전한 폴백 슬러그.
    형식: '{date}_{kebab-title}' — title 을 한글/영숫자/하이픈만 남기고 60자 cap."""
    raw = re.sub(r'^\d{4}-\d{2}-\d{2}\s*[—\-]\s*', '', title or '').strip()
    raw = re.sub(r'\s+', '-', raw)
    raw = re.sub(r'[^\w가-힣\-]', '', raw, flags=re.UNICODE).strip('-')[:60]
    if not raw:
        raw = "untitled"
    d = target_date or date.today().strftime("%Y-%m-%d")
    return f"{d}_{raw}"


# ── 사전 중복 필터 (2026-05-16) ───────────────────────
# 배경:
#   "보관함 채우기 5편" 요청 시 카테고리 필터 후 5건 트림 → 그 5건이 모두
#   post_history 의 발행본과 100% 매칭이면 [1/5]~[5/5] 모두 Skip 되어 1편만
#   생성되는 사고가 있었다 (2026-05-16). 트림 *전* 에 발행본 중복 후보를
#   먼저 차감해야 사용자가 요청한 N편이 실제로 채워진다.
#
# 동작:
#   - 각 post 의 title 을 duplicate_checker.check_duplicate(title) 로 1회 비교
#   - is_dup True (settings.DUPLICATE_THRESHOLD = 0.80 이상) 면 차감
#   - 차감 항목은 excluded_titles.add_excluded(title, reason=...) 로 누적
#   - 영속 파일(post_history.json·trend_data.json)은 절대 수정하지 않음
#
# 비용:
#   후보 50건 × 발행본 500건 Jaccard = ~25k 비교. 500ms 미만, 캐시 불필요.
def _prefilter_known_duplicates(
        posts: list,
        *,
        reason: str = "prefilter_history",
        mutate_excluded: bool = True) -> tuple[list, list]:
    """post_history 와 중복(score ≥ DUPLICATE_THRESHOLD)인 후보를 사전 차감.

    Args:
        posts: pending posts 리스트. 각 항목은 dict 로 'title' 키를 가진다.
        reason: excluded_titles 에 기록될 사유 태그.
        mutate_excluded: True 면 차감 시 add_excluded() 로 누적 저장.
            preflight_check 처럼 부수 효과 0 을 보장해야 하는 호출자는 False.

    Returns:
        (kept, dropped):
            kept   — 중복 없는 후보 (원본 순서 보존)
            dropped — score ≥ THRESHOLD 로 차감된 후보
    """
    if not posts:
        return [], []

    kept: list = []
    dropped: list = []
    for info in posts:
        title = info.get("title", "") if isinstance(info, dict) else ""
        if not title:
            kept.append(info)
            continue
        try:
            is_dup, _matched, _score = check_duplicate(title)
        except (OSError, ValueError, TypeError) as exc:
            # 후보를 1개의 비교 실패로 잃지 않도록 — keep + 경고
            log.warning(
                f"  ⚠️  사전 중복 비교 실패 ({type(exc).__name__}: {exc}) — "
                f"keep: {title[:46]}"
            )
            kept.append(info)
            continue

        if is_dup:
            dropped.append(info)
            # 다음 사이클에서도 후보 진입 전 차감되도록 excluded_titles 누적.
            # 저장 실패는 메인 흐름을 막지 않는다.
            if mutate_excluded:
                try:
                    add_excluded(title, reason=reason)
                except OSError as exc:
                    log.warning(
                        f"  ⚠️  excluded_titles 추가 실패 ({exc}) — 메모리 차감만 적용"
                    )
        else:
            kept.append(info)

    return kept, dropped


def _release_driver(driver, mode: str, *, quiet: bool = False) -> None:
    """발행 드라이버 정리 — spawn 은 quit(기존 동작), attach 는 세션 유지(quit 생략).

    attach 드라이버는 사용자가 직접 로그인한 **세션 브라우저**라 발행이 끝나도 닫으면
    안 된다(다음 발행에 재사용). spawn 은 우리가 띄운 일회용 크롬이라 반드시 정리한다.
    어떤 예외도 삼켜 정리/종료 흐름을 방해하지 않는다(회귀 0: spawn quit 동작 불변).

    Args:
        quiet: True 면 attach 유지 안내 로그만 억제한다(정리 동작 자체는 동일).
            GUI 유휴 세션 프로브처럼 반복 호출되는 읽기 전용 경로 전용이며,
            기본 False 라 발행 경로 로그는 그대로다.
    """
    if driver is None:
        return
    if str(mode).lower() == "attach":
        if not quiet:
            log.info("  🖥 attach 세션 브라우저는 닫지 않습니다(다음 발행 재사용)")
        return
    try:
        driver.quit()
    except Exception:  # noqa: BLE001
        log.debug("메인 드라이버 종료 실패")


# ── S2: 보관본 우선 발행 — 지연 준비 작업 헬퍼 (2026-09-10) ──────────
# 첫 발행 대상이 보관본(draft_ready)이면 그 발행엔 필요 없는 준비 작업들을
# 본 루프(생성+업로드) 종료 뒤로 미룬다. 아래 헬퍼들은 "지연 실행" 시에만
# 쓰이고, fast-path 가 아닐 때는 기존 인라인 코드가 그대로 즉시 실행된다
# (동작 100% 보존 — 헬퍼는 새 경로에서만 호출).

def _deadline_tracker_active(s) -> bool:
    """마감 추적 기능이 이 프로필에서 활성인지 판정(브랜드 게이트 포함).

    settings.DEADLINE_TRACKER_ENABLED 와 config.brand.deadline_enabled() 가 모두
    True 여야 활성. 브랜드 조회 실패는 fail-open(활성 유지) — 기존 동작 보존.
    """
    active = getattr(s, "DEADLINE_TRACKER_ENABLED", True)
    if active:
        try:
            from config import brand as _brand
            active = _brand.deadline_enabled()
        except Exception as _be:  # noqa: BLE001 — 브랜드 조회 실패 시 기존 동작(활성) 유지
            log.debug(f"  브랜드 마감 설정 조회 실패 — 활성 유지: {_be}")
    return active


def _run_deadline_detection_deferred(s) -> dict:
    """마감 임박 공고 감지 + 엑셀 삽입만 수행(포스팅 목록 재로드 없음).

    fast-path 로 이 스텝이 첫 발행 뒤로 미뤄질 때 호출한다. 호출 시점엔
    이미 본 루프가 끝나 있어 이번 실행의 posts/total 을 다시 조립해도
    아무도 소비하지 않는다 — fetch_deadlines()·엑셀 삽입이라는 실제로
    느린 작업 자체는 그대로 수행되므로(조용히 삭제 아님), 다음 실행의
    get_pending_posts 가 이번에 삽입된 항목을 집어간다.

    Returns: pipeline_status 엔트리 하나.
    """
    if not _deadline_tracker_active(s):
        log.info("  ⏭  마감 추적 비활성 프로필 — skip(지연)")
        return {"success": True, "message": "비활성 프로필 — skip"}
    try:
        from core.deadline_tracker import fetch_deadlines, inject_urgent_to_excel
        log.info("\n🚨 마감 임박 공고 감지 중... (지연 실행 — 다음 실행 큐에 반영)")
        deadlines = fetch_deadlines()
        urgent_cnt = sum(1 for d in deadlines if d.get("urgency") in ("urgent", "warning"))
        if urgent_cnt > 0:
            inserted = inject_urgent_to_excel(
                _keyword_excel(s), s.KEYWORD_PUBLISH_POOL_SHEET, deadlines)
            if inserted:
                log.info(f"  ✅ 마감 임박 {inserted}개 삽입(다음 실행에 반영)")
        return {"success": True, "message": f"{len(deadlines)}건 수집(지연)"}
    except (ImportError, urllib.error.URLError, FileNotFoundError) as e:
        log.warning(f"  ⚠️  마감 감지 실패(지연): {e}")
        return {"success": False, "message": str(e)}


def _run_retry_failed_deferred(stop_flag, progress_cb) -> dict:
    """실패 글 자동 재시도(LLM 재생성)를 지연 실행(fast-path 전용)."""
    try:
        _auto_retry_failed(stop_flag=stop_flag, progress_cb=progress_cb)
        return {"success": True, "message": "지연 실행 완료"}
    except Exception as e:    # noqa: BLE001 — 재시도 실패가 발행 결과에 영향 없음
        log.warning(f"  ⚠️  실패 글 자동 재시도 오류(지연): {e}")
        return {"success": False, "message": str(e)}


def _run_trend_update_deferred(s) -> dict:
    """월요일 트렌드 업데이트를 지연 실행(fast-path 전용). 그 외 요일은 no-op."""
    if date.today().weekday() != 0:
        return {"success": True, "message": "월요일 아님 — skip"}
    try:
        from core.trend_updater import update_excel_trends
        log.info("\n📈 주간 트렌드 업데이트 중... (지연 실행)")
        update_excel_trends(_keyword_excel(s), s.KEYWORD_PUBLISH_POOL_SHEET)
        return {"success": True, "message": "트렌드 업데이트 완료(지연)"}
    except (ImportError, FileNotFoundError, KeyError) as e:
        log.warning(f"  ⚠️  트렌드 업데이트 실패(지연): {e}")
        return {"success": False, "message": str(e)}


def _run_stats_and_rank_deferred(dry_run: bool) -> dict:
    """방문자·유입 통계 수집 + (월요일) 순위 추적을 지연 실행(fast-path 전용)."""
    try:
        from core.stats_fetcher import collect_stats
        from core import performance_db as pdb
        pdb.sync_from_history()
        collect_stats()
        msg = "통계 수집 완료(지연)"
        if date.today().weekday() == 0 and not dry_run:
            try:
                from core.rank_tracker import run_rank_tracking
                log.info("\n🔎 주간 순위 추적 중... (지연 실행)")
                rr = run_rank_tracking(max_keywords=120)
                msg += f" · 순위 {rr.get('ranked', 0)}/{rr.get('checked', 0)} (TOP10 {rr.get('top10', 0)})"
            except (ImportError, urllib.error.URLError, TimeoutError) as e:
                log.warning(f"  ⚠️  순위 추적 실패(지연): {e}")
        return {"success": True, "message": msg}
    except (ImportError, urllib.error.URLError, TimeoutError) as e:
        log.warning(f"  ⚠️  통계 수집 실패(지연): {e}")
        return {"success": False, "message": str(e)}


def _run_refresh_batch_deferred(dry_run: bool, target_date, reserve_times, driver) -> dict:
    """화·목 30일 리프레시 배치를 지연 실행(fast-path 전용).

    공유 드라이버(driver)를 그대로 재사용한다 — 호출자는 이 함수를
    ``_release_driver`` 이전(드라이버가 아직 살아 있는 지점)에서만 호출해야 한다.
    """
    if date.today().weekday() not in (1, 3):
        return {"success": True, "message": "화/목 아님 — skip"}
    try:
        from core.post_refresher import run_refresh_batch
        log.info("\n♻️  30일 글 리프레시 확인 중... (지연 실행, 공유 세션)")
        run_refresh_batch(max_count=3, dry_run=dry_run,
                          target_date=target_date, after_slots=reserve_times,
                          driver=driver)
        return {"success": True, "message": "리프레시 완료(지연)"}
    except (ImportError, FileNotFoundError) as e:
        log.warning(f"  ⚠️  리프레시 실패(지연): {e}")
        return {"success": False, "message": str(e)}


def run_pipeline(target_date: str = None, dry_run: bool = False,
                 progress_cb=None, stop_flag=None,
                 auth_event=None, auth_cb=None,
                 save_only: bool = False,
                 max_count=None,
                 allowed_categories=None,
                 fill_category=None,
                 only_titles=None,
                 repair_legacy_images: bool = False) -> dict:
    """파이프라인 진입점 — 중복 실행 방지 프로세스 락으로 본체를 감싼다.

    작업 스케줄러 다중 등록 + GUI 동시 실행 시 상태파일 경합·이중예약을
    막기 위해 실제 발행/save 모드에서는 배타 파일 락(pipeline.lock)을 획득한다.
    이미 실행 중이면 즉시 조기 반환하고, 종료 시 finally 에서 락을 해제한다.

    dry_run(테스트 모드)은 부수 효과가 없어 락을 스킵한다 — 진단 목적의
    병렬 실행을 막을 이유가 없다. (2026-07-03)
    """
    # dry_run 은 발행/save 부수 효과가 없어 락 대상에서 제외.
    if dry_run:
        return _run_pipeline_impl(
            target_date=target_date, dry_run=dry_run,
            progress_cb=progress_cb, stop_flag=stop_flag,
            auth_event=auth_event, auth_cb=auth_cb,
            save_only=save_only, max_count=max_count,
            allowed_categories=allowed_categories,
            fill_category=fill_category, only_titles=only_titles,
            repair_legacy_images=False)

    lock = PipelineLock(PIPELINE_LOCK_PATH)
    if not lock.acquire():
        log.warning("이미 실행 중인 파이프라인이 있어 중복 실행을 막습니다")
        return {"total": 0, "success": 0, "pipeline_status": {},
                "error": "already_running"}
    try:
        return _run_pipeline_impl(
            target_date=target_date, dry_run=dry_run,
            progress_cb=progress_cb, stop_flag=stop_flag,
            auth_event=auth_event, auth_cb=auth_cb,
            save_only=save_only, max_count=max_count,
            allowed_categories=allowed_categories,
            fill_category=fill_category, only_titles=only_titles,
            repair_legacy_images=bool(repair_legacy_images and not save_only))
    finally:
        lock.release()


def _report_draft_save_failure(progress_cb, title: str, index: int, total: int) -> None:
    """Report a draft persistence failure independently from publication uncertainty."""
    if progress_cb:
        progress_cb(title, "fail", index + 1, total)


def _publish_failure_progress_state(unconfirmed: bool) -> str:
    """Expose uncertain publish outcomes for manual verification, not as failures."""
    return "verify" if unconfirmed else "fail"


def _run_pipeline_impl(target_date: str = None, dry_run: bool = False,
                       progress_cb=None, stop_flag=None,
                       auth_event=None, auth_cb=None,
                       save_only: bool = False,
                       max_count=None,
                       allowed_categories=None,
                       fill_category=None,
                       only_titles=None,
                       repair_legacy_images: bool = False) -> dict:
    """
    progress_cb(title, status) — GUI 콜백
    stop_flag: threading.Event — 외부에서 중지 신호
    auth_event: threading.Event — GUI 보안 인증 완료 대기용
    auth_cb: callable — 보안 인증 상태 알림 콜백
    save_only: True 면 글 생성·이미지·품질검증까지만 수행하고 업로드는 스킵.
        본문/이미지/예약시각을 runs/{slug}/post_data.json + _status.json 으로
        보관해 차후 별도 업로드 단계에서 발행. 기본 False(기존 흐름 100% 보존).
        (2026-05-09: 생성/발행 분리 1단계 — examplebrand 일 10편 자동 발행 무영향.)
    max_count: 양수면 pending 중 상위 N건만 처리. None/0 이면 전체.
    allowed_categories: 발행 실행 시 후보를 이 카테고리들로 필터링한다.
        2026-05-16 pre-flight 팝업에서 "매칭되는 카테고리만 실행" 선택 시 GUI 가
        전달. set / list / tuple 허용. 빈 collection 또는 None 이면 필터링 없음.
        trend_data.json / 엑셀 / excluded_titles.json 같은 영속 파일은 절대 수정
        하지 않는다 — 메모리 상 posts 리스트에서만 차감.
    fill_category: save_only 모드에서 보관할 카테고리를 1개로 한정한다.
        대시보드 "부족 카테고리 채우기" 다이얼로그가 전달. 값이 12 카테고리에
        포함되지 않으면 무시 (전체 채우기로 폴백).
    only_titles: **제한 경로**(2026-09-18). 지정하면 후보를 이 제목들로 못박는다 —
        트렌드 정렬·보관함 드레인·카테고리 필터·사전 중복 필터를 모두 건너뛰고
        사용자가 고른 글만 처리한다. 한 편 파일럿처럼 "어느 글인가"까지 정해야 할
        때 쓴다. None/빈 값이면 기존 동작 100% 그대로.
    """
    s     = _cfg()

    # ── 큐 진단 (2026-05-26) ─────────────────────────────
    # 새 예약 이월 로직이 메모리에 적재됐는지 즉시 확인할 수 있도록
    # run_pipeline 진입 직후 한 줄 로그. 사용자가 GUI/CLI 를 재시작했는데도
    # 같은 D+1 슬롯에 잡히면 이 라인이 안 나오거나 "비어 있음" 으로 찍힌다.
    log.info(_format_queue_diagnostic(target_date=target_date, cfg=s))

    # ── 카테고리 필터 정규화 (메모리 차감만 — 영속 파일 무수정) ───────
    # 빈 set/list 도 명시적으로 None 처리해 "필터 없음" 으로 본다.
    _allowed_cats_set = None
    if allowed_categories:
        try:
            _allowed_cats_set = {
                c.strip() for c in allowed_categories
                if isinstance(c, str) and c.strip()
            } or None
        except TypeError:
            log.warning(
                f"  ⚠️  allowed_categories 가 iterable 이 아님 — 무시: "
                f"{type(allowed_categories).__name__}"
            )
            _allowed_cats_set = None
    _fill_cat = None
    if save_only and isinstance(fill_category, str) and fill_category.strip():
        _fill_cat = fill_category.strip()

    _only = [str(t).strip() for t in (only_titles or []) if str(t or "").strip()]
    _topic_profile = _nbaprofile.active_profile() if not _only else None
    _auto_topic_mode = False
    if not _only:
        try:
            _auto_topic_mode = _automatic_topic_configured(s, _topic_profile)
        except (ImportError, OSError, TypeError, ValueError) as _auto_config_exc:
            log.warning(f"  ⚠️  자동 주제 설정 확인 실패: {_auto_config_exc} — Excel 흐름 유지")

    # Auto mode has no Excel source; preserve legacy backup behavior for every other path.
    if not _auto_topic_mode:
        try:
            _backup_excel(_keyword_excel(s))
        except Exception as e:
            log.warning(f"  ⚠️  엑셀 백업 실패: {e}")

    # save_only 일 때는 보관함 차감 + max_count 트림이 뒤따르므로
    # get_pending_posts 의 DAILY_POST_LIMIT 한도를 풀고 후보를 넉넉히 받는다.
    _pending_limit = None
    if save_only and max_count is not None and max_count > 0:
        _pending_limit = max(max_count * 5, 200)

    # ── 제한 경로: 지정된 제목만 (2026-09-18) ───────────────────────────
    if _only:
        posts = select_only_titles(_only, cfg=s, target_date=target_date)
        if not posts:
            log.warning("✨ 지정한 제목을 찾지 못해 처리할 항목이 없습니다")
            return {"total": 0, "success": 0, "pipeline_status": {},
                    "saved_drafts": 0, "drafts_used": 0,
                    "error": "only_titles_not_found"}
    elif _auto_topic_mode:
        try:
            topic_limit = _pending_limit or max_count or getattr(s, "DAILY_POST_LIMIT", 1)
            not_stopped = not (stop_flag and stop_flag.is_set())
            auto_posts = _automatic_topic_posts(
                s, refresh=not dry_run and not_stopped, limit=topic_limit,
                profile=_topic_profile)
            posts = auto_posts if auto_posts is not None else []
        except (ImportError, OSError, TypeError, ValueError) as _auto_topic_exc:
            log.warning(f"  ⚠️  자동 주제 후보 선택 실패: {_auto_topic_exc}")
            posts = []
    else:
        posts = get_pending_posts(target_date, cfg=s, limit=_pending_limit)
    if not posts:
        log.info("✨ 포스팅할 항목 없음")
        return {"total": 0, "success": 0, "pipeline_status": {}}

    # save_only=True 모드에서는 mark_done 이 호출되지 않아 같은 키워드가 다시 픽업된다.
    # 보관함에 이미 'draft_ready' 상태로 있는 키워드를 사전에 제외해야 50편을 만들 때
    # 정확히 50개의 unique 글이 쌓인다 (중복 _2/_3 접미사 회피).
    if save_only and posts and not _only:
        try:
            from core.draft_store import find_draft_for_title
            filtered = []
            skipped_drafted = 0
            for info in posts:
                if find_draft_for_title(info.get("title", "")) is not None:
                    skipped_drafted += 1
                    continue
                filtered.append(info)
            if skipped_drafted > 0:
                log.info(
                    f"  ⏭  보관함에 이미 있는 키워드 {skipped_drafted}건 제외 — "
                    f"{len(posts)}건 → {len(filtered)}건"
                )
            posts = filtered
            if not posts:
                log.info("✨ 보관함 차감 후 처리할 신규 키워드 없음")
                return {"total": 0, "success": 0, "pipeline_status": {},
                        "saved_drafts": 0, "drafts_used": 0}
        except Exception as _fe:
            log.warning(f"  ⚠️  보관함 차감 중 예외 — 차감 없이 진행: {_fe}")

    # ── 카테고리 필터: pre-flight 팝업 "매칭만 실행" / "특정 카테고리만 채우기" ──
    # max_count 적용 전에 차감해야 사용자가 요청한 N편이 필터 적용 후에도 의미를 갖는다.
    # 영속 파일은 절대 수정하지 않으며, posts 메모리 리스트만 줄인다.
    if posts and (_allowed_cats_set or _fill_cat) and not _only:
        try:
            filter_set = _allowed_cats_set or {_fill_cat}
            before = len(posts)
            posts = [
                p for p in posts
                if detect_category(
                    re.sub(r'^\d{4}-\d{2}-\d{2}\s*[—\-]\s*', '', p.get("title", ""))
                ) in filter_set
            ]
            removed = before - len(posts)
            if removed > 0:
                filter_label = (
                    f"fill_category={_fill_cat}" if _fill_cat
                    else f"allowed_categories={sorted(filter_set)}"
                )
                log.info(
                    f"  🎯 카테고리 필터 적용 ({filter_label}): "
                    f"{before}건 → {len(posts)}건 (차감 {removed})"
                )
            if not posts:
                log.info("✨ 카테고리 필터 적용 후 처리할 후보 없음")
                return {
                    "total": 0, "success": 0, "pipeline_status": {},
                    "saved_drafts": 0, "drafts_used": 0,
                    "skipped_quality": 0,
                }
        except (re.error, TypeError, ValueError) as _ce:
            log.warning(f"  ⚠️  카테고리 필터 적용 중 예외 — 필터 없이 진행: {_ce}")

    # ── 사전 중복 필터 (max_count 트림 *전* 에 적용) ─────────────────
    # dry_run 일 때는 진단 가치를 보존하기 위해 차감하지 않는다.
    if posts and not dry_run and not _only:
        try:
            before_dup = len(posts)
            kept, dropped = _prefilter_known_duplicates(posts)
            posts = kept
            if dropped:
                log.info(
                    f"  🧹 사전 중복 필터: {len(dropped)}건 차감 (이미 발행됨) — "
                    f"{before_dup}건 → {len(posts)}건"
                )
            if not posts:
                log.info("✨ 사전 중복 필터 후 처리할 신규 후보 없음")
                return {
                    "total": 0, "success": 0, "pipeline_status": {},
                    "saved_drafts": 0, "drafts_used": 0,
                    "skipped_quality": 0,
                }
        except Exception as _pe:    # noqa: BLE001 — 사전 필터 실패가 본 파이프라인을 막지 않게
            log.warning(f"  ⚠️  사전 중복 필터 예외 — 필터 없이 진행: {_pe}")

    _archive_priority = (not save_only and not dry_run and not _only and not _auto_topic_mode
                         and getattr(s, "DRAFT_DRAIN_ENABLED", True))
    if (not _archive_priority and max_count is not None
            and max_count > 0 and len(posts) > max_count):
        log.info(f"  ✂️  max_count={max_count} 적용: {len(posts)}건 → {max_count}건으로 트림")
        posts = posts[:max_count]
    elif (max_count is not None and max_count > 0
          and len(posts) < max_count and not dry_run):
        # 사전 필터 적용 후 풀이 요청보다 작으면 사용자에게 알린다.
        log.warning(
            f"  ⚠️  사전 필터 후 후보 부족: {len(posts)}건 < 요청 {max_count}건 — "
            f"가능한 만큼만 처리"
        )

    # ── 보관함 우선 소진 (archive-first drain, 2026-07-13) ──
    # 신규 생성 전에 보관함 elite 를 먼저 소진 — 큐 앞에 드레인 제목을 붙이고 fresh 를
    # 부족분만 남기게 트림한다. 드레인 제목은 루프의 find_draft_for_title 경로에서
    # 신규 생성 없이 보관본 그대로 발행 준비된다. save_only/dry_run 은 기존 흐름 100%
    # 보존(드레인 스킵), 실패는 fail-open. assign_reservation_times 는 재조립된 posts 기준.
    if not save_only and not dry_run and not _only and not _auto_topic_mode:
        _drain_cap = (max_count if (max_count is not None and max_count > 0)
                      else s.DAILY_POST_LIMIT)
        posts, _n_drain, _n_fresh = _apply_archive_drain(
            posts, _drain_cap, cfg=s, allowed_cats=_allowed_cats_set)
        if _archive_priority and _drain_cap > 0:
            posts = posts[:_drain_cap]
        if _n_drain:
            log.info(f"  📦 보관함 우선 소진: 드레인 {_n_drain}편 + 신규 {_n_fresh}편")

    total         = len(posts)
    reserve_times = assign_reservation_times(total, target_date, posts, cfg=s)   # ★ 카테고리 인식
    mode          = "예약발행" if s.USE_RESERVATION else "즉시발행"
    pipeline_status = {}  # step_name → {"success": bool, "message": str}
    skipped_dup_count = 0
    skipped_quality_count = 0   # 2026-05-08 추가: 품질 미달(타임아웃·placeholder)로 업로드 스킵 카운트
    saved_draft_count = 0       # 2026-05-09 추가: save_only=True 시 디스크 보관 카운트
    skipped_low_volume_count = 0  # 2026-07-03 추가: 검색량 하한 미달(무수요 초-롱테일)로 발행 스킵
    skipped_institution_count = 0  # 2026-07-08 추가: 제도-단위 중복 게이트로 발행 차단(사촌 자기잠식 방지)
    published_records: list = []   # 하네스 일일 리포트용 (성공·실패 모두)
    post_slugs: dict = {}           # i → 하네스 slug (hook B/C/D 간 공유)

    # ── S2: 보관본 우선 발행 조기 판별 (2026-09-10) ──────────────
    # 첫 발행 대상이 이미 QA 통과한 보관본이면, 그 발행에 필요 없는 준비 작업
    # (실패 글 재생성·마감 감지·트렌드·통계/순위·30일 리프레시·gov24 캡처
    # 드라이버 기동)을 첫 발행 뒤로 미룬다. 사용자 실측(2026-09-10 앱 로그):
    # 이 준비 작업들이 순차로 돌아 보관본이 있는데도 첫 업로드까지 7분+ 걸렸다.
    # find_draft_for_title 은 draft_store 의 순수 JSON 조회(부작용 0·LLM/네트워크
    # 호출 없음)이며, 본 루프가 i=0 에서 어차피 수행할 조회를 한 번 앞당길 뿐이다.
    # posts[0] 이 사전 필터를 통과했어도 루프 안 중복/제도포화/검색량 게이트에서
    # 스킵될 수 있어 "진짜 첫 발행"이 posts[1] 이 될 가능성은 있다(예측 한계) —
    # 그 경우도 fail-open: 지연된 준비 작업은 루프 종료 후 정상 수행되므로
    # 발행 자체가 막히지는 않는다.
    _first_draft_hit = None
    if not save_only and not dry_run and posts:
        try:
            from core.draft_store import find_draft_for_title
            _first_draft_hit = find_draft_for_title(posts[0].get("title", ""))
        except Exception as _fdh_err:    # noqa: BLE001 — 조기 판별 실패는 일반 경로로
            log.debug(f"  보관본 조기 판별 예외 — 일반 경로로 진행: {_fdh_err}")
            _first_draft_hit = None
    # 제한 경로(only_titles) — 사용자가 "이 글만" 지정했다. 대규모 유지보수
    # (실패 글 재생성·마감 감지·트렌드·통계/순위·30일 리프레시)는 지정하지 않은
    # 다른 글을 건드리는 작업이므로 **아예 돌리지 않는다**(지연도 아님).
    # _fast_path 와 같은 게이트를 재사용해 앞쪽 즉시 실행을 막고, 루프 뒤
    # '지연 준비 작업' 블록도 `_limited_run` 으로 건너뛴다.
    _limited_run = bool(_only)
    _fast_path = _limited_run or (_first_draft_hit is not None)
    if _limited_run:
        log.info(f"  🎯 지정 발행({len(_only)}편) — 준비 작업(재시도·마감·트렌드·통계·"
                 f"리프레시)은 이 실행에서 돌리지 않습니다")
    elif _fast_path:
        log.info(
            f"  📦 저장해 둔 글이 있어 바로 발행합니다 — 집필 없이 첫 편 발행 후 "
            f"나머지 준비 작업(재시도·마감·트렌드·통계·리프레시)을 이어서 진행합니다 "
            f"(첫 대상: {_first_draft_hit.get('slug', '')})"
        )
        if progress_cb:
            try:
                progress_cb("저장된 글 재사용 — 준비 작업은 첫 발행 뒤로 미룸",
                             "prep", 0, total)
            except Exception:    # noqa: BLE001 — progress_cb 계약 밖 상황은 흐름을 막지 않음
                log.debug("  progress_cb(prep) 호출 실패 — 로그만 유지")

    log.info(f"\n{'━'*54}")
    log.info(f"🚀 자동 포스팅 시작 [{mode}]  총 {total}개")
    log.info(f"{'━'*54}")
    for i, (p, t) in enumerate(zip(posts, reserve_times), 1):
        clean = re.sub(r'^\d{4}-\d{2}-\d{2} — ', '', p['title'])[:36]
        cat   = detect_category(clean)
        log.info(f"  {i}. {t}  [{cat:4s}]  {clean}")

    # [HOOK A] 하네스 — 배치 brief 기록 (관측 전용)
    if _he is not None:
        try:
            _he.emit_batch_brief(
                target_date or date.today().strftime("%Y-%m-%d"),
                posts, reserve_times)
        except Exception:
            pass

    # ── 실패 글 자동 재시도 (retry_count < 3인 항목) ────────
    # save_only=True 면 driver 가 필요한 retry_failed_post 가 동작 못 하므로 스킵.
    # fast-path(첫 대상이 보관본)일 땐 LLM 재생성(느림)이 첫 발행을 막지 않도록
    # 루프 종료 후로 미룬다(아래 "지연 준비 작업" 블록).
    if not save_only and not _fast_path:
        try:
            _auto_retry_failed(stop_flag=stop_flag, progress_cb=progress_cb)
            pipeline_status["auto_retry_failed"] = {"success": True, "message": "즉시 실행"}
        except Exception as e:
            log.warning(f"  ⚠️  실패 글 자동 재시도 오류: {e}")
            pipeline_status["auto_retry_failed"] = {"success": False, "message": str(e)}
    elif _fast_path and not save_only:
        log.info("  ⏭  실패 글 자동 재시도 — 첫 발행(보관본) 뒤로 지연")
        pipeline_status["auto_retry_failed"] = {
            "success": True, "message": "보관본 우선 발행 — 첫 발행 뒤로 지연"}

    # ★ 마감 임박 키워드 자동 감지 → 엑셀 최상단 삽입 (매일 실행)
    #   정부24/정책포털 마감 추적은 예시브랜드(정부지원) 전용 기능이다. 영어교육 등
    #   비정부 분야 프로필에선 정부 공고가 발행 큐에 주입되면 안 되므로 게이트로 차단한다.
    #   default(예시브랜드)는 settings.DEADLINE_TRACKER_ENABLED=True + brand.deadline_enabled()=True
    #   라 기존대로 실행되고, named 비정부 프로필은 둘 중 하나가 False 면 통째로 skip 한다.
    # fast-path(첫 대상이 보관본)일 땐 네트워크 3소스 조회가 첫 발행을 막지 않도록
    # 루프 종료 후로 미룬다(`_run_deadline_detection_deferred`).
    if _fast_path:
        log.info("  ⏭  마감 임박 감지 — 첫 발행(보관본) 뒤로 지연")
        pipeline_status["deadline_tracker"] = {
            "success": True, "message": "보관본 우선 발행 — 첫 발행 뒤로 지연"}
        _deadline_on = False   # 아래 즉시 실행 블록 스킵용(지연 경로로 이미 처리됨)
    elif _auto_topic_mode:
        _deadline_on = False
    else:
        _deadline_on = _deadline_tracker_active(s)
    if _deadline_on:
        try:
            from core.deadline_tracker import fetch_deadlines, inject_urgent_to_excel
            log.info("\n🚨 마감 임박 공고 감지 중...")
            deadlines = fetch_deadlines()
            urgent_cnt = sum(1 for d in deadlines if d.get("urgency") in ("urgent","warning"))
            if urgent_cnt > 0:
                inserted = inject_urgent_to_excel(
                    _keyword_excel(s), s.KEYWORD_PUBLISH_POOL_SHEET, deadlines)
                if inserted:
                    # 엑셀 변경됐으므로 포스팅 목록 재로드 (save_only 시 한도 풀고)
                    posts = get_pending_posts(target_date, cfg=s, limit=_pending_limit)
                    # save_only 시 보관함 차감 (재로드 후에도 동일 가드)
                    if save_only and posts:
                        try:
                            from core.draft_store import find_draft_for_title
                            before = len(posts)
                            posts = [
                                p for p in posts
                                if find_draft_for_title(p.get("title", "")) is None
                            ]
                            if before > len(posts):
                                log.info(
                                    f"  ⏭  보관함 차감(재로드 후): {before}건 → {len(posts)}건"
                                )
                        except Exception as _fe:
                            log.warning(f"  ⚠️  보관함 차감 중 예외(재로드 후) — 무시: {_fe}")
                    # 마감 재로드 시점에도 카테고리 필터를 다시 적용해야 일관성 유지.
                    if posts and (_allowed_cats_set or _fill_cat):
                        try:
                            filter_set = _allowed_cats_set or {_fill_cat}
                            before = len(posts)
                            posts = [
                                p for p in posts
                                if detect_category(
                                    re.sub(r'^\d{4}-\d{2}-\d{2}\s*[—\-]\s*', '', p.get("title", ""))
                                ) in filter_set
                            ]
                            if before > len(posts):
                                log.info(
                                    f"  🎯 카테고리 필터(재로드 후): {before}건 → {len(posts)}건"
                                )
                        except (re.error, TypeError, ValueError) as _ce:
                            log.warning(f"  ⚠️  카테고리 필터(재로드 후) 예외 — 무시: {_ce}")
                    # 마감 재로드 후에도 사전 중복 필터 동일 적용 (트림 *전*)
                    if posts and not dry_run:
                        try:
                            before_dup = len(posts)
                            kept, dropped = _prefilter_known_duplicates(posts)
                            posts = kept
                            if dropped:
                                log.info(
                                    f"  🧹 사전 중복 필터(재로드 후): {len(dropped)}건 차감 — "
                                    f"{before_dup}건 → {len(posts)}건"
                                )
                        except Exception as _pe:    # noqa: BLE001
                            log.warning(
                                f"  ⚠️  사전 중복 필터(재로드 후) 예외 — 무시: {_pe}"
                            )
                    if max_count is not None and max_count > 0 and len(posts) > max_count:
                        log.info(f"  ✂️  max_count={max_count} 적용: {len(posts)}건 → {max_count}건으로 트림")
                        posts = posts[:max_count]
                    # 재로드로 초기 드레인이 유실됐으므로 보관함 우선 소진을 다시 적용한다
                    # (초기 경로와 동일 로직 — 계획은 결정적이라 멱등).
                    if not save_only and not dry_run:
                        _drain_cap = (max_count if (max_count is not None and max_count > 0)
                                      else s.DAILY_POST_LIMIT)
                        posts, _n_drain, _n_fresh = _apply_archive_drain(
                            posts, _drain_cap, cfg=s, allowed_cats=_allowed_cats_set)
                        if _n_drain:
                            log.info(
                                f"  📦 보관함 우선 소진(재로드 후): 드레인 {_n_drain}편 + "
                                f"신규 {_n_fresh}편")
                    total = len(posts)
                    reserve_times = assign_reservation_times(total, target_date, posts, cfg=s)
                    log.info(f"  ✅ 마감 임박 {inserted}개 삽입 → 포스팅 목록 재로드 ({total}개)")
            pipeline_status["deadline_tracker"] = {"success": True, "message": f"{len(deadlines)}건 수집"}
        except (ImportError, urllib.error.URLError, FileNotFoundError) as e:
            log.warning(f"  ⚠️  마감 감지 실패: {e}")
            pipeline_status["deadline_tracker"] = {"success": False, "message": str(e)}
    elif not _fast_path:
        log.info("  ⏭  마감 추적 비활성 프로필 — skip")
        pipeline_status["deadline_tracker"] = {"success": True, "message": "비활성 프로필 — skip"}

    # ★ 30일 자동 리프레시(화·목)는 아래 '메인 드라이버 로그인 성공 직후'로 이동했다.
    #   메인 예약 발행과 같은 드라이버·로그인 세션을 공유해 '오늘 실행'당 GUI 인증을
    #   1회로 만들기 위함(예전엔 여기서 리프레시가 별도 브라우저·별도 로그인을 띄워
    #   재인증이 2회 발생했다). 실제 블록은 인터리브 루프 직전을 참조.

    # ④ 트렌드 업데이트 (월요일 자동 실행) — fast-path 는 첫 발행 뒤로 지연.
    if _fast_path:
        if date.today().weekday() == 0:
            log.info("  ⏭  주간 트렌드 업데이트 — 첫 발행(보관본) 뒤로 지연")
        pipeline_status["trend_updater"] = {
            "success": True, "message": "보관본 우선 발행 — 첫 발행 뒤로 지연"}
    elif date.today().weekday() == 0:
        try:
            from core.trend_updater import update_excel_trends
            log.info("\n📈 주간 트렌드 업데이트 중...")
            update_excel_trends(_keyword_excel(s), s.KEYWORD_PUBLISH_POOL_SHEET)
            pipeline_status["trend_updater"] = {"success": True, "message": "트렌드 업데이트 완료"}
        except (ImportError, FileNotFoundError, KeyError) as e:
            log.warning(f"  ⚠️  트렌드 업데이트 실패: {e}")
            pipeline_status["trend_updater"] = {"success": False, "message": str(e)}

    # ⑥ 방문자·유입 성과 수집 (Phase 1 측정 루프) — fast-path 는 첫 발행 뒤로 지연.
    if _fast_path:
        log.info("  ⏭  방문자·유입 통계 수집 — 첫 발행(보관본) 뒤로 지연")
        pipeline_status["stats_fetcher"] = {
            "success": True, "message": "보관본 우선 발행 — 첫 발행 뒤로 지연"}
    else:
        try:
            from core.stats_fetcher import collect_stats
            from core import performance_db as pdb
            pdb.sync_from_history()          # 발행 카탈로그를 성과 DB에 최신화(항상, 저비용)
            collect_stats()                  # 방문자 총계 + 세션 기반 글별 수집
            msg = "통계 수집 완료"
            # 순위 추적은 API 부하를 고려해 월요일에만(자격증명 있으면).
            if date.today().weekday() == 0 and not dry_run:
                try:
                    from core.rank_tracker import run_rank_tracking
                    log.info("\n🔎 주간 순위 추적 중...")
                    rr = run_rank_tracking(max_keywords=120)
                    msg += f" · 순위 {rr.get('ranked', 0)}/{rr.get('checked', 0)} (TOP10 {rr.get('top10', 0)})"
                except (ImportError, urllib.error.URLError, TimeoutError) as e:
                    log.warning(f"  ⚠️  순위 추적 실패: {e}")
            pipeline_status["stats_fetcher"] = {"success": True, "message": msg}
        except (ImportError, urllib.error.URLError, TimeoutError) as e:
            log.warning(f"  ⚠️  통계 수집 실패: {e}")
            pipeline_status["stats_fetcher"] = {"success": False, "message": str(e)}

    # ════════════════════════════════════════════════════
    # 드라이버 초기화 (루프 전)
    # save_only=True 면 글 생성·이미지·QA·draft_store 만 동작하므로
    # Selenium driver / 네이버 로그인 / gov24 캡처 드라이버 일체 스킵.
    # 사용자가 GUI 를 띄워둔 상태에서도 충돌 없이 보관함을 채울 수 있다.
    # ════════════════════════════════════════════════════
    # ── ⏹ 운영 정지 스위치 — **브라우저를 띄우기 전에** 본다 (2026-09-23) ──────
    # 예전에는 정지 확인이 글 루프 안(새 글 경계)에만 있어서, 일시중지·비상정지
    # 상태에서도 ①크롬이 뜨고 ②네이버 로그인이 시도되고 ③화·목이면 루프 진입
    # 전에 도는 리프레시 배치가 실제로 예약까지 했다. 로그인도 외부 행동이다.
    # save_only(보관함 채우기)·dry_run(테스트로 보기)은 외부 행동이 아니라 통과.
    if not save_only and not dry_run:
        try:
            from core import publish_state_bridge as _psb0
            _pipeline_hold = _psb0.guard_new_item() or ""
        except Exception:  # noqa: BLE001 — 모듈 부재 = 기존 동작(막지 않음)
            _pipeline_hold = ""
        if _pipeline_hold:
            _hold_msg = ("전체 비상정지" if _pipeline_hold == "estop"
                         else "이 계정 일시중지")
            log.warning(f"⏹ 운영 정지 상태({_hold_msg}) — 브라우저·로그인·발행을 "
                        "하지 않고 종료합니다. 정지를 푼 뒤 다시 실행하세요.")
            pipeline_status["stopped"] = {"success": False, "message": _hold_msg}
            return {"total": 0, "success": 0, "pipeline_status": pipeline_status,
                    "saved_drafts": 0, "drafts_used": 0,
                    "held": _pipeline_hold, "error": f"operation_{_pipeline_hold}"}

    gov24_driver = None
    _gov24_attempted = False   # 지연 생성 시 한 번만 시도(재시도 스팸 방지)

    def _ensure_gov24_driver():
        """gov24 캡처 드라이버를 아직 만든 적 없으면 지금 만든다(1회만 시도).

        fast-path(첫 대상이 보관본)일 땐 루프 진입 전 즉시 기동을 건너뛰고,
        실제로 캡처가 필요한 항목(보관본이 아닌 신규 집필)을 만났을 때
        이 함수가 처음 호출되며 그때 기동한다 — 첫 발행(보관본)엔 어차피
        캡처가 필요 없으므로 불필요한 크롬 한 대를 안 띄운다.
        """
        nonlocal gov24_driver, _gov24_attempted
        if _gov24_attempted:
            return gov24_driver
        _gov24_attempted = True
        try:
            from core.gov24_screenshot import make_screenshot_driver
            gov24_driver = make_screenshot_driver()
            pipeline_status["gov24_driver"] = {"success": True, "message": "기동 완료"}
        except Exception as e:  # noqa: BLE001 — 캡처 드라이버 실패가 발행을 막지 않음
            log.warning(f"  ⚠️  gov24 드라이버 생성 실패: {e}", exc_info=True)
            pipeline_status["gov24_driver"] = {"success": False, "message": str(e)}
        return gov24_driver

    if save_only:
        log.info("💾 save_only 모드 — Selenium driver / 네이버 로그인 / gov24 캡처 모두 스킵")
    _gov24_needed = not dry_run and not s.SKIP_IMAGE_INSERT and not save_only
    if _gov24_needed and not _fast_path:
        _ensure_gov24_driver()
    elif _gov24_needed and _fast_path:
        log.info("  ⏭  gov24 캡처 드라이버 기동 지연 — 실제로 필요해질 때(신규 집필 항목) 생성")

    driver = None
    # 발행 경로 이원화(P0-① Phase 1): 기본 spawn(현행) vs attach(로그인된 세션 크롬에 얹기).
    # 기본값 spawn 이면 아래 attach 블록은 통째로 건너뛰고 현행 spawn 경로가 그대로 돈다.
    _publish_mode = str(getattr(s, "PUBLISH_MODE", "spawn")).lower()
    if (_publish_mode == "attach"
            and not dry_run and not s.USE_NAVER_API and not save_only):
        # 🖥 attach 경로 — naver_login 을 호출하지 않는다(세션이 이미 살아있음).
        # 신원 게이트(_assert_publish_identity)는 발행 직전 _do_upload 내부에서 자동 수행.
        _attach_port = int(getattr(s, "CHROME_ATTACH_PORT", 9222))
        if bool(getattr(s, "CHROME_ATTACH_AUTOSTART", True)):
            try:
                from core import chrome_session
                _sess_status = chrome_session.ensure_session(_attach_port)
                # 방금 띄운 크롬('started')은 디버깅 포트가 열릴 때까지 잠깐 폴링해
                # 콜드 스타트 레이스를 흡수한다('already'/기타는 대기 불필요).
                if _sess_status == "started":
                    if not chrome_session.wait_until_running(_attach_port, timeout=10.0):
                        log.warning("  ⚠️  발행 세션 크롬 디버깅 포트가 제때 열리지 않음 — "
                                    "attach 를 시도합니다(실패 시 fail-closed)")
            except Exception as e:  # noqa: BLE001 — 자동 기동 실패는 무해(attach 시도로 판정)
                log.warning(f"  ⚠️  발행 세션 자동 기동 실패(무시): {e}")
        from core.naver_uploader import attach_to_session
        driver = attach_to_session(_attach_port)
        if driver is None:
            if gov24_driver:
                try:
                    gov24_driver.quit()
                except Exception:
                    pass
            log.error("  ❌ 발행 세션 attach 실패 — fail-closed")
            # 무인 실행 중 세션이 없거나 로그인이 안 돼 있으면 사용자가 알 방법이
            # 없으므로 카카오로 처방까지 한 줄. (알림 실패는 흐름에 무영향)
            try:
                from core.notifier import send_kakao
                send_kakao("긴급 알림",
                           "🖥 발행 세션 크롬을 열고 네이버 로그인을 확인하세요.\n"
                           "👉 GUI 🏠 홈에서 [발행 세션 열기]를 누른 뒤 그 창에서 "
                           "네이버에 로그인하고 다시 실행하세요.")
            except Exception:  # noqa: BLE001
                pass
            status_path = LOG_DIR / "pipeline_status.json"
            atomic_write_json(str(status_path), {
                "date": (target_date or date.today().strftime("%Y-%m-%d")),
                "steps": pipeline_status,
                "total": total, "success": 0,
            })
            return {"total": total, "success": 0, "error": "attach_failed",
                    "pipeline_status": pipeline_status}

    if (not dry_run and not s.USE_NAVER_API and not save_only
            and _publish_mode != "attach"):
        # attach 세션 크롬이 열려 있으면 같은 프로필 user-data-dir 을 두 프로세스가
        # 잡아 Chrome SingletonLock 충돌이 날 수 있다 — spawn 기동 전 감지+경고만.
        # (동작 불변: 경고 로그뿐, 발행 흐름은 그대로 진행)
        try:
            from core import chrome_session
            if chrome_session.is_running(int(getattr(s, "CHROME_ATTACH_PORT", 9222))):
                log.warning("  ⚠️  attach 세션 크롬이 열려 있어 spawn 발행이 프로필 충돌할 수 "
                            "있습니다 — 세션 창을 닫은 뒤 실행하세요(무인 발행은 spawn 전용).")
        except Exception:  # noqa: BLE001 — 감지 실패는 무해
            pass
        from core.naver_uploader import init_driver, naver_login
        driver = init_driver()
        if not naver_login(driver, auth_event=auth_event, auth_cb=auth_cb):
            driver.quit()
            driver = None
            if gov24_driver:
                try:
                    gov24_driver.quit()
                except Exception:
                    pass
            log.error("  ❌ 로그인 실패")
            # UX P2 — 자동 실행(무인) 중 로그인이 풀리면 사용자가 알 방법이
            # 없었음 → 카카오로 처방까지 한 줄. (알림 실패는 흐름에 무영향)
            try:
                from core.notifier import send_kakao
                send_kakao("긴급 알림",
                           "🔐 네이버 로그인이 풀려 오늘 발행을 못 했어요.\n"
                           "👉 GUI 🏠 홈에서 [지금 로그인]을 누른 뒤 다시 실행하세요.")
            except Exception:  # noqa: BLE001
                pass
            # 파이프라인 상태 저장 (대시보드가 stale 데이터를 표시하지 않도록)
            status_path = LOG_DIR / "pipeline_status.json"
            atomic_write_json(str(status_path), {
                "date": (target_date or date.today().strftime("%Y-%m-%d")),
                "steps": pipeline_status,
                "total": total, "success": 0,
            })
            return {"total": total, "success": 0, "error": "login_failed",
                    "pipeline_status": pipeline_status}

    # ★ 30일 자동 리프레시(화·목) — 메인과 같은 드라이버·로그인 세션 공유(재인증 방지).
    # driver 는 위에서 GUI 인증(auth_event/auth_cb)으로 로그인된 공유 세션이며, 여기서
    # 리프레시에 그대로 주입해 로그인·인증을 다시 하지 않는다('오늘 실행'당 인증 1회).
    #  · 실 Selenium 발행: driver 존재 → run_refresh_batch 가 그 세션을 재사용(재로그인·
    #    재인증 없음). 로그인 실패 시엔 위에서 이미 return 되어 이 지점에 도달하지 않는다.
    #  · dry_run / USE_NAVER_API: driver=None 이지만 run_refresh_batch 내부가 각각
    #    드라이 프리뷰 / HTTP API 경로로 처리한다(드라이버 불필요).
    #  · save_only: 발행하지 않는 모드이므로 리프레시도 skip.
    # fast-path(첫 대상이 보관본)일 땐 리프레시(편당 LLM 재생성+재발행, 최대 3편)가
    # 첫 발행을 막지 않도록 루프 종료 후로 미룬다(드라이버가 아직 살아 있는 지점,
    # `_run_refresh_batch_deferred` 참조 — 공유 세션 그대로 재사용).
    if _fast_path:
        if date.today().weekday() in (1, 3) and not save_only:
            log.info("  ⏭  30일 글 리프레시 — 첫 발행(보관본) 뒤로 지연")
        pipeline_status["post_refresher"] = {
            "success": True, "message": "보관본 우선 발행 — 첫 발행 뒤로 지연"}
    elif date.today().weekday() in (1, 3) and not save_only:   # 화=1, 목=3
        try:
            from core.post_refresher import run_refresh_batch
            log.info("\n♻️  30일 글 리프레시 확인 중...(공유 세션)")
            # reserve_times/target_date 는 이 함수 스코프에 이미 존재
            # (reserve_times: L1284/재로드 시 L1395). 유효 슬롯을 못 만든 글은 skip.
            run_refresh_batch(max_count=3, dry_run=dry_run,
                              target_date=target_date, after_slots=reserve_times,
                              driver=driver)
            pipeline_status["post_refresher"] = {"success": True, "message": "리프레시 완료"}
        except (ImportError, FileNotFoundError) as e:
            log.warning(f"  ⚠️  리프레시 실패: {e}")
            pipeline_status["post_refresher"] = {"success": False, "message": str(e)}

    # ════════════════════════════════════════════════════
    # 인터리브 루프: 생성 → 즉시 업로드 → 최소 간격 보정 대기
    # ════════════════════════════════════════════════════
    log.info(f"\n{'━'*54}")
    log.info(f"📝 글 생성 + 업로드 ({total}개)")
    log.info(f"{'━'*54}")

    success = 0
    # Per-item outcomes keep failures, manual checks, and skips distinct.
    _publication_results = {}
    generated_count = 0
    draft_used_count = 0   # 2026-05-09: 글 생성 직전 보관본 활용 카운트 (drafts_used)

    try:
        for i, info in enumerate(posts):
            loop_start = time.monotonic()

            if stop_flag and stop_flag.is_set():
                log.info("\n⏹ 사용자에 의해 중지됨")
                break

            title        = info["title"]
            reserve_time = reserve_times[i]
            pid          = f"{date.today().strftime('%m%d')}_{i+1:02d}"

            # ── 정지 스위치 · 중복 발행 차단 (글 경계, 2026-09-18) ─────────
            # 파일 기반이라 GUI 가 떠 있지 않아도 듣는다. 브리지가 없거나 깨지면
            # 아무 것도 막지 않는다(기존 동작 보존).
            try:
                from core import publish_state_bridge as _psb
            except Exception:  # noqa: BLE001 — 모듈 부재 = 기존 경로
                _psb = None
            if _psb is not None:
                _hold = _psb.guard_new_item()
                if _hold:
                    log.warning(
                        "⏹ 운영 정지 신호 감지 — 새 글을 시작하지 않습니다 "
                        f"({'전체 비상정지' if _hold == 'estop' else '이 계정 일시중지'})")
                    break
                # 결말이 미상인 글(예약 등록 신호를 못 읽은 글)은 다시 올리지 않는다.
                _blocked = _psb.blocked_reason(title=title)
                if _blocked:
                    log.warning(
                        f"  ⏭  확인 필요로 보류된 글 — 다시 올리지 않습니다 ({_blocked}): "
                        f"{title[:40]}")
                    _publication_results[i] = "unconfirmed"
                    write_log(title, False, reserve_time, f"awaiting_manual_check:{_blocked}")
                    if progress_cb:
                        progress_cb(title, "verify", i + 1, total)
                    continue

            log.info(f"\n[{i+1}/{total}] {title[:46]}")

            if progress_cb:
                progress_cb(title, "start", i+1, total)

            # ── 생성 단계 ──
            # draft_hit 은 try 블록 안에서 보관본 lookup 결과로 채워진다.
            # try 외부의 업로드 단계에서 mark_status 호출에 사용하므로
            # break/continue 경로에서도 변수가 정의돼 있도록 미리 초기화.
            draft_hit = None
            try:
                # ② 중복 검사
                is_dup, matched, score = check_duplicate(title)
                # 하네스 slug 산출 (이번 글 동안 hook B/C/D 가 공유)
                _slug = None
                if _he is not None:
                    try:
                        _slug = _he.make_slug(title, target_date or date.today().strftime("%Y-%m-%d"))
                        post_slugs[i] = _slug
                    except Exception:
                        _slug = None
                _clean_kw = re.sub(r'^\d{4}-\d{2}-\d{2} — ', '', title)
                _cat = detect_category(_clean_kw)
                if is_dup:
                    log.info(f"  ⏭  중복 감지 ({score*100:.0f}%) → Skip")
                    # 다음 사이클에서 후보 단계에서 사전 제외되도록 등록.
                    # 실패해도 메인 흐름을 막지 않는다 (로그만 남기고 진행).
                    try:
                        add_excluded(title, reason="duplicate")
                    except Exception as _ex_err:
                        log.warning(f"  excluded_titles 추가 실패(무시): {_ex_err}")
                    log.info(f"     유사 기존 글: {matched[:42]}")
                    write_log(title, False, reserve_time, f"중복({score*100:.0f}%)")
                    skipped_dup_count += 1
                    if progress_cb:
                        progress_cb(title, "duplicate", i+1, total)
                    # [HOOK B-skip] 하네스 — 중복 스킵 brief
                    if _he is not None and _slug:
                        try:
                            _he.emit_post_brief(
                                _slug, title, _cat,
                                dup_status={"is_dup": True, "matched": matched, "score": score})
                            published_records.append({
                                "slug": _slug, "title": _clean_kw, "category": _cat,
                                "reserve_time": reserve_time, "success": False,
                                "url": "", "skipped": "duplicate"})
                        except Exception:
                            pass
                    continue

                # ②-inst 제도-단위 중복 게이트 (title-only 다음) — 최근 30일 같은
                # 정규화 제도명(에너지바우처·긴급생계지원 등) 발행 편수가 임계(기본 3)를
                # 초과하면 신규 발행을 차단한다. 폐기가 아니라 "기존 최약 사촌 1편을
                # 리프레시" 신호로 브리프를 남긴다(build_merge_brief). 검사 오류는
                # fail-open(발행 막지 않음). settings.INSTITUTION_DEDUP_ENABLED 로 on/off.
                if getattr(s, "INSTITUTION_DEDUP_ENABLED", True):
                    try:
                        _inst_gate = institution_dedup.evaluate_gate(
                            title,
                            window_days=getattr(s, "INSTITUTION_DEDUP_WINDOW_DAYS", 30),
                            max_per_window=getattr(s, "INSTITUTION_DEDUP_MAX_PER_WINDOW", 3),
                        )
                    except (OSError, ValueError, TypeError) as _ig_err:
                        _inst_gate = {"blocked": False}
                        log.warning(f"  제도 중복 게이트 오류(무시하고 진행): {_ig_err}")
                    if _inst_gate.get("blocked"):
                        _inst_name = _inst_gate.get("institution", "")
                        _inst_cnt = _inst_gate.get("count", 0)
                        _inst_lim = _inst_gate.get("limit", 0)
                        log.info(
                            f"  ⏭  제도 포화 차단 — '{_inst_name}' 최근 "
                            f"{getattr(s, 'INSTITUTION_DEDUP_WINDOW_DAYS', 30)}일 "
                            f"{_inst_cnt}편 > {_inst_lim} → 신규 대신 최약 사촌 리프레시 권장"
                        )
                        # 사촌 병합/리프레시 브리프 산출(발행/재작성은 기존 트리거).
                        try:
                            _merge_brief = institution_dedup.build_merge_brief(
                                title, category=_cat)
                            if _merge_brief and _he is not None and _slug:
                                _he.emit_post_brief(
                                    _slug, title, _cat,
                                    dup_status={
                                        "is_dup": True,
                                        "institution_gate": _merge_brief,
                                        "score": score,
                                    })
                        except Exception as _mb_err:  # noqa: BLE001 — 브리프 실패가 흐름을 막지 않음
                            log.debug(f"  제도 병합 브리프 생략: {_mb_err}")
                        write_log(title, False, reserve_time,
                                  f"제도포화({_inst_name}:{_inst_cnt})")
                        skipped_institution_count += 1
                        if progress_cb:
                            progress_cb(title, "institution_saturated", i + 1, total)
                        continue

                # ②-c 검색량 게이트 (2차 방어선) — 제목이 실제 검색 수요 있는 키워드에
                # 앵커돼 있지 않으면(무수요 초-롱테일) 발행 스킵. keyword_expander 가
                # 못 거른 무수요 제목을 발행 직전에 최종 차단한다. 검사 오류는 fail-open.
                if getattr(s, "SEARCH_VOLUME_GATE_ENABLED", False):
                    try:
                        from core.search_volume import meets_threshold, best_keyword
                        _min_vol = getattr(s, "MIN_SEARCH_VOLUME", 500)
                        _vol_ok = meets_threshold(title, _min_vol)
                        _anchor, _anchor_vol = best_keyword(title)
                    except (OSError, ValueError, TypeError, ImportError) as _sv_err:
                        _vol_ok, _anchor, _anchor_vol = True, "", 0  # 오류 시 발행 막지 않음
                        log.warning(f"  검색량 게이트 오류(무시하고 진행): {_sv_err}")
                    # 엑셀 미달이어도 네이버 자동완성에 뜨는 '실검색어'면 통과(키워드 목록 자동 확장).
                    # 좁은 엑셀 화이트리스트가 실검색 온-브랜드 주제(주거급여·보금자리론 등)를
                    # 오차단하던 문제 해소(2026-07-15). _demand: True=실검색 / False=수요약함 /
                    # None=조회실패(불확정). None 이면 영구 제외하지 않고 이번만 스킵(다음 재시도).
                    _demand = None
                    _demand_consulted = False
                    if not _vol_ok and getattr(s, "LIVE_DEMAND_GATE_ENABLED", True):
                        _demand_consulted = True
                        try:
                            from core.live_demand import live_demand_ok
                            _demand = live_demand_ok(title)
                        except Exception as _ld_err:  # noqa: BLE001 — 판정 실패는 불확정 취급
                            _demand = None
                            log.debug(f"  자동완성 수요 판정 오류(무시): {_ld_err}")
                        if _demand is True:
                            _vol_ok = True
                            log.info(f"  ✅ 자동완성 실검색 확인 → 통과(엑셀 미등록): {title[:40]}")
                    if not _vol_ok:
                        log.info(
                            f"  ⏭  검색량 하한 미달 (앵커 볼륨 {_anchor_vol} < {_min_vol}) "
                            f"→ Skip: {title[:40]}")
                        # 자동완성 조회 실패(None)면 영구 제외 보류 — 네트워크 블립으로 실검색
                        # 주제를 영구 차단하는 사고 방지. 그 외(게이트 미사용·확정 수요약함)는 기존대로 제외.
                        if not (_demand_consulted and _demand is None):
                            try:
                                add_excluded(title, reason="low_search_volume")
                            except Exception as _ex_err:  # noqa: BLE001 — 등록 실패가 흐름을 막지 않음
                                log.warning(f"  excluded_titles 추가 실패(무시): {_ex_err}")
                        else:
                            log.info("  (자동완성 조회 실패 — 영구 제외 보류, 다음 재시도)")
                        write_log(title, False, reserve_time, f"검색량미달({_anchor_vol})")
                        skipped_low_volume_count += 1
                        if progress_cb:
                            progress_cb(title, "low_volume", i + 1, total)
                        continue

                # ── 보관함 lookup: 같은 키워드 보관본이 있으면 글 생성 스킵하고 업로드만 ──
                # save_only=True 모드에서는 보관함을 채우는 흐름이므로 lookup 안 함.
                # dry_run 일 때는 lookup 하더라도 어차피 업로드를 안 하니 무해 — 그러나
                # 이미지·HTML 변환 결과까지 보관본 그대로 쓰면 dry_run 진단 가치가 사라지므로 스킵.
                # (draft_hit 은 try 블록 진입 직전에 None 으로 미리 초기화돼 있다.)
                if not save_only and not dry_run:
                    try:
                        from core.draft_store import find_draft_for_title
                        draft_hit = find_draft_for_title(title)
                    except Exception as _de:
                        log.debug(f"  보관함 lookup 예외 — 무시하고 신규 생성: {_de}")
                        draft_hit = None

                if draft_hit is not None:
                    if not _prepare_archive_images(
                            draft_hit,
                            repair_legacy_images=bool(
                                repair_legacy_images and not dry_run and not save_only)):
                        image_result = draft_hit.get("_image_prepare_result") or {}
                        image_reason = _archive_image_reason(image_result.get("reason"))
                        if image_reason == "unverified_assets":
                            count = image_result.get("unverified", 0)
                            log.warning(
                                "보관본 이미지 %d장이 출처 미검증 상태라 발행을 보류했습니다. "
                                "저장 글 → 해당 글 선택 → 미검증 이미지 다시 만들기를 실행하세요.",
                                count if isinstance(count, int) and count >= 0 else 0)
                        elif image_reason in {"unverifiable_assets", "reference_assets",
                                              "placeholder_assets", "tampered_assets",
                                              "undecodable_assets"}:
                            log.warning(
                                "보관본 이미지 출처를 안전하게 확인할 수 없어 발행을 보류했습니다 (%s)",
                                image_reason)
                        elif image_reason == "login_required":
                            log.warning("이미지 생성기 로그인이 필요해 발행을 보류했습니다")
                        elif image_reason == "rate_limited":
                            log.warning("이미지 생성기 한도에 도달해 발행을 보류했습니다")
                        elif image_reason == "already_running":
                            log.warning("이미지 채움이 진행 중이라 발행을 보류했습니다")
                        else:
                            log.warning("보관본 이미지 준비가 완료되지 않아 발행을 보류했습니다 (%s)",
                                        image_reason)
                        write_log(title, False, reserve_time, image_reason)
                        skipped_quality_count += 1
                        if progress_cb:
                            progress_cb(title, "fail", i + 1, total)
                        continue
                    # 보관본은 이미 QA 통과한 글 — _failed 게이트도 거치지 않는다.
                    # 이미지·HTML 변환·preview 저장·HOOK C 모두 스킵하고 곧바로 업로드.
                    _drft_post = draft_hit.get("post_data") or {}
                    post_data = {
                        "title":        _drft_post.get("title", title),
                        "content":      "",                                  # plain 본문은 안 쓰지만 호환을 위해 빈 문자열
                        "content_html": _drft_post.get("content", ""),       # 보관본은 HTML 본문이 "content" 키
                        "tags":         list(_drft_post.get("tags", []) or []),
                        "category":     _drft_post.get("category") or _cat,
                    }
                    # 젠스파크 이미지 자동 채움(Phase 3) — 채운 이미지의 per-post 삽입 신호를
                    # 승계한다. **참일 때만** 키를 추가해 "새 키 0" 계약을 지킨다(부재/거짓이면
                    # 키 미추가 = 기존 dict 와 동일). 없으면 _should_insert_images 가
                    # SKIP_IMAGE_INSERT 를 따라 삽입 안 함(기존 동작 유지).
                    if _drft_post.get("_force_image_insert"):
                        post_data["_force_image_insert"] = True
                    # ★버그 수정: image_paths 는 draft_hit 최상위가 아니라 post_data 안에 저장된다
                    #   (save_draft payload → load_drafts 가 post_data 로 반환). 기존 코드는 최상위를
                    #   읽어 항상 []였다. post_data 우선 + 최상위 fallback(방어).
                    all_images = list(_drft_post.get("image_paths")
                                      or draft_hit.get("image_paths") or [])
                    generated_count += 1
                    draft_used_count += 1
                    log.info(
                        f"  📤 보관본 활용: {draft_hit.get('slug')} "
                        f"({len(post_data['content_html'])}자 / 태그 {len(post_data['tags'])}개, "
                        f"created_at: {draft_hit.get('created_at')})"
                    )
                    log.info(
                        f"  ✅ 보관본 사용: {len(post_data['content_html'])}자 / "
                        f"이미지 {len(all_images)}장"
                    )
                else:
                    # [HOOK B-start] 하네스 — 본격 생성 직전 brief
                    if _he is not None and _slug:
                        try:
                            _he.emit_post_brief(
                                _slug, title, _cat,
                                dup_status={"is_dup": False, "score": score})
                        except Exception:
                            pass

                    # 중지 확인 (글 생성 전)
                    if stop_flag and stop_flag.is_set():
                        log.info("\n⏹ 사용자에 의해 중지됨 (글 생성 직전)")
                        break

                    # 글 생성 (①품질검증 내장)
                    if info.get("_auto_topic") and isinstance(info.get("topic_context"), dict):
                        post_data = generate_blog_post(
                            title, stop_flag=stop_flag,
                            topic_context=info["topic_context"])
                    else:
                        post_data = generate_blog_post(title, stop_flag=stop_flag)
                    # 중지 요청으로 조기 반환된 경우
                    if post_data.get("_stopped") or (stop_flag and stop_flag.is_set()):
                        log.info("\n⏹ 사용자에 의해 중지됨 (글 생성 후)")
                        break

                    # ②-b 본문 중복 검사 — 제목은 통과했지만 생성된 본문이 기존 글과
                    # 거의 동일한 경우(다른 제목·같은 내용) 발행 전에 스킵한다.
                    # 검사 오류는 발행을 막지 않도록 fail-open(중복 아님)으로 처리한다.
                    # (본문 비교는 record_post 가 body_kw 를 저장한 이후 글부터 유효.)
                    try:
                        _body_dup, _body_match, _body_score = check_duplicate(
                            title, content=post_data.get("content", "") or "")
                    except (OSError, ValueError, TypeError) as _bd_err:
                        _body_dup, _body_match, _body_score = False, "", 0.0
                        log.warning(f"  본문 중복 검사 오류(무시하고 진행): {_bd_err}")
                    if _body_dup:
                        log.info(
                            f"  ⏭  본문 중복 감지 ({_body_score*100:.0f}%) → Skip "
                            f"(유사 기존 글: {_body_match[:42]})")
                        try:
                            add_excluded(title, reason="body_duplicate")
                        except Exception as _ex_err:  # noqa: BLE001 — 등록 실패가 흐름을 막지 않음
                            log.warning(f"  excluded_titles 추가 실패(무시): {_ex_err}")
                        write_log(title, False, reserve_time,
                                  f"본문중복({_body_score*100:.0f}%)")
                        skipped_dup_count += 1
                        if progress_cb:
                            progress_cb(title, "duplicate", i + 1, total)
                        continue

                    # ⑤ 내부 링크 포함 HTML 변환
                    post_data["content_html"] = markdown_to_naver_html(
                        post_data["content"], title, post_data.get("category",""))

                    # 미리보기 저장
                    save_preview(post_data, reserve_time, pid)

                    # 중지 확인 (이미지 생성 전)
                    if stop_flag and stop_flag.is_set():
                        log.info("\n⏹ 사용자에 의해 중지됨 (이미지 생성 직전)")
                        break

                    # 이미지 생성 (SKIP_IMAGE_INSERT 또는 save_only 시 건너뜀)
                    # save_only 모드는 gov24_driver=None 이라 공공사이트 캡처 불가 + 사용자가 발행 시점에
                    # 이미지를 직접 다루므로 보관함에는 본문/태그만 저장한다.
                    all_images = []
                    _planned_new_post_fill = (
                        not dry_run and not save_only and
                        _plan_fill_covers_new_post_publish())
                    if (not s.SKIP_IMAGE_INSERT and not save_only and
                            not _planned_new_post_fill):
                        gov_shots = []
                        if not dry_run:
                            # 카테고리 기반 공공기관 사이트 디스패치
                            # (정부지원 7종 → gov24, 절세 → 홈택스, 부동산 → 청약홈, 노후 → 국민연금공단)
                            try:
                                from core.public_site_screenshot import (
                                    capture_for_category, site_for_category,
                                )
                                target_site = site_for_category(_cat)
                                if target_site == "gov24" or target_site is None:
                                    # gov24 또는 매핑 없음 → 기존 함수 (호환)
                                    gov_shots = capture_gov24_page(
                                        title, pid, driver=_ensure_gov24_driver())
                                else:
                                    gov_shots = capture_for_category(
                                        title, pid, _cat, driver=_ensure_gov24_driver())
                            except Exception as _cap_err:  # noqa: BLE001 — 캡처 실패가 발행을 깨뜨리면 안 됨
                                log.warning(f"  ⚠️  공공사이트 캡처 dispatch 실패: {_cap_err} — gov24 fallback")
                                gov_shots = capture_gov24_page(title, pid, driver=_ensure_gov24_driver())
                        all_images = generate_all_images(title, post_data, gov_shots, pid)
                    elif _planned_new_post_fill:
                        log.info(
                            "  ⏭️  기존 이미지 생성기 건너뜀 — 새 글 계획 기반 채움이 "
                            "이미지 생성과 완전성 확인을 담당합니다.")
                    else:
                        log.info("  ⏭️  이미지 생성 건너뜀 (SKIP_IMAGE_INSERT=True)")

                    # ── ★ Codex CLI 이미지 자동 생성 (2026-05-16, Day 1~3 통합) ──
                    # settings.IMAGE_INJECT_ENABLED 가 True 이고 dry_run 이 아니면
                    # ImageDispatcher 가 본문 figure placeholder 6장(디지털 8장)을 채운다.
                    # save_only / 인터리브 모드 양쪽에서 동작 — 보관본은 곧바로 발행 가능한
                    # 상태로 저장된다. Tier 1 실패 시 dispatcher 가 자동으로 Tier 2 폴백.
                    # SKIP_IMAGE_INSERT(='이미지 건너뛰기' 체크박스)가 켜지면 본문 '삽입'뿐
                    # 아니라 이미지 '생성'(image_dispatch)도 건너뛴다. 예전엔 삽입만 막히고
                    # 생성은 계속 돌아 이미지가 안 쓰이는데도 편당 수 분씩 낭비됐다
                    # (2026-07-14 사용자 지적). 이제 체크박스 하나로 생성·삽입이 함께 OFF.
                    # 2026-09-10: 보관함 채우기(save_only)에서 계획 기반 채움이 켜져
                    # 있으면 여기서는 **만들지 않는다** — 아래 [HOOK G] 가 계획대로
                    # 한 번만 생성한다(구 dispatcher 와의 이중 생성 차단).
                    _dispatch_on = (not dry_run
                                    and not s.SKIP_IMAGE_INSERT
                                    and getattr(s, "IMAGE_INJECT_ENABLED", False)
                                    and bool(_slug))
                    if _dispatch_on and _plan_fill_is_sole_image_producer(save_only):
                        _dispatch_on = False
                        log.info(
                            "  ⏭  구 이미지 dispatcher 건너뜀 — 보관함 채우기는 "
                            "계획 기반 이미지 채움이 한 번만 생성합니다(이중 생성 방지).")
                    # 2026-09-18: 새 글을 곧장 발행하는 경로도 아래 "업로드 단계" 직전에
                    # 계획 기반 채움(_fill_new_post_images)이 한 번만 생성한다 — 여기서
                    # 구 dispatcher 까지 돌리면 계획·슬롯을 모르는 이미지가 먼저 섞여
                    # 같은 글에 이미지를 두 번 만드는 셈이 된다(이중 생성 방지).
                    elif (_dispatch_on and not save_only
                          and _plan_fill_covers_new_post_publish()):
                        _dispatch_on = False
                        log.info(
                            "  ⏭  구 이미지 dispatcher 건너뜀 — 새 글 발행은 "
                            "계획 기반 이미지 채움이 한 번만 생성합니다(이중 생성 방지).")
                    if _dispatch_on:
                        try:
                            _codex_images = _run_image_dispatcher(
                                slug=_slug,
                                title=post_data.get("title", title),
                                category=_cat,
                                post_data=post_data,
                            )
                        except Exception as _img_err:  # noqa: BLE001 — 이미지 실패가 발행을 안 막음
                            log.warning(
                                f"  ⚠️  Codex 이미지 dispatch 예외 — "
                                f"{type(_img_err).__name__}: {_img_err}"
                            )
                            _codex_images = []
                        if _codex_images:
                            all_images = list(all_images) + list(_codex_images)

                    log.info(f"  ✅ 생성 완료: {len(post_data['content'])}자 / 이미지 {len(all_images)}장")
                    generated_count += 1

                    # [HOOK C] 하네스 — 생성 결과 메타 기록 (본문은 복사하지 않음)
                    if _he is not None and _slug:
                        try:
                            _preview_path = str((_DATA_DIR / "previews" / f"{pid}.html"))
                            _he.emit_run_record(
                                _slug, post_data,
                                qa_meta={}, image_count=len(all_images),
                                status="generated", preview_path=_preview_path)
                        except Exception:
                            pass

            except Exception as e:
                _publication_results[i] = "failed"
                log.error(f"  ❌ 생성 오류: {e}\n{traceback.format_exc()}")
                write_log(title, False, reserve_time, str(e))
                if progress_cb: progress_cb(title, "fail", i+1, total)
                failed = safe_load_json(str(FAILED_POSTS_PATH), default=list)
                failed.append({
                    "title": title,
                    "reason": "생성 실패",
                    "date": datetime.now().strftime("%Y-%m-%d %H:%M"),
                })
                failed = failed[-50:]
                atomic_write_json(str(FAILED_POSTS_PATH), failed)
                # [HOOK D-genfail] 생성 단계에서 실패 → publish_brief 도 실패로 기록
                if _he is not None and _slug:
                    try:
                        _he.emit_publish_brief(
                            _slug, _clean_kw, _cat,
                            tags=[], schedule_at=str(reserve_time),
                            qa_status="FAIL", success=False, url="",
                            error=f"생성 실패: {e}")
                        published_records.append({
                            "slug": _slug, "title": _clean_kw, "category": _cat,
                            "reserve_time": str(reserve_time), "success": False,
                            "url": "", "error": "생성 실패"})
                    except Exception:
                        pass
                continue

            # ── dry_run: 업로드 건너뜀 ──
            if dry_run:
                log.info(f"  🔍 [테스트] {len(post_data['content'])}자 / 이미지 {len(all_images)}장")
                write_log(title, True, reserve_time, "dry_run")
                if progress_cb: progress_cb(title, "ok", i+1, total)
                continue

            # ── 품질 게이트: content_generator가 _failed 플래그를 세웠으면 업로드 스킵 ──
            # (2026-05-08 사고 — Claude CLI 타임아웃으로 "생성 오류" placeholder가 그대로
            # 발행된 사례 재발 방지. 발행 차단은 content_generator가 이미 판정했음.)
            if post_data.get("_failed"):
                _fail_reason = post_data.get("_failed_reason", "unknown")
                log.warning(
                    f"  ⛔ 품질 미달 — 업로드 스킵: {title[:46]} "
                    f"(사유: {_fail_reason}, 본문 {len(post_data.get('content', ''))}자)"
                )
                # 품질/사실 실패(retries_exceeded·qa_failed·fact_gate)로 막힌 글은 본문이
                # 살아 있으므로 버리지 말고 needs_fix 로 보관한다(사람/후속 수정 대기).
                # 깡통 사유(placeholder/too_short)는 헬퍼 내부에서 저장하지 않는다.
                # save_only/dry_run 은 이 게이트에 도달하지 않으므로 발행 경로 전용.
                _saved_for_fix = _save_failed_draft_for_fix(
                    title, post_data, all_images, _cat, _fail_reason,
                    slug=_slug, target_date=target_date, schedule_at=str(reserve_time))
                write_log(title, False, reserve_time, f"quality_gate_blocked:{_fail_reason}")
                skipped_quality_count += 1
                if progress_cb: progress_cb(title, "fail", i+1, total)
                # [HOOK D-quality] 하네스 — 품질 게이트 차단 brief
                if _he is not None and _slug:
                    try:
                        _he.emit_publish_brief(
                            _slug, post_data.get("title", _clean_kw), _cat,
                            tags=post_data.get("tags", []),
                            schedule_at=str(reserve_time),
                            qa_status="FAIL", success=False, url="",
                            error=f"quality_gate_blocked:{_fail_reason}",
                            dia_score=post_data.get("_dia_score"),
                            aeo_score=post_data.get("_aeo_score"))
                        published_records.append({
                            "slug": _slug, "title": post_data.get("title", _clean_kw),
                            "category": _cat, "reserve_time": str(reserve_time),
                            "success": False, "url": "",
                            "skipped": f"quality:{_fail_reason}"})
                    except Exception:
                        pass
                continue

            # ── save_only 모드: drafts 폴더에 보관하고 업로드는 스킵 (Stage 1 분리) ──
            # (2026-05-09: 생성과 발행 분리 1단계. _failed 게이트를 먼저 통과한
            # 글만 여기 도달 — 품질 미달 placeholder 가 drafts 에 쌓이지 않는다.
            # 하네스 publish_brief 는 emit 하지 않음 — 발행 직전 단계가 아니므로
            # 차후 실제 업로드 시 emit 하면 된다.)
            if save_only:
                try:
                    from core.draft_store import save_draft
                    _draft_slug = _slug
                    if not _draft_slug:
                        if _he is not None:
                            try:
                                _draft_slug = _he.make_slug(
                                    title,
                                    target_date or date.today().strftime("%Y-%m-%d"),
                                )
                            except Exception:
                                _draft_slug = None
                        if not _draft_slug:
                            _draft_slug = _make_fallback_slug(title, target_date)
                    draft_post_data = {
                        "title":    post_data["title"],
                        "content":  post_data["content_html"],
                        "tags":     post_data.get("tags", []),
                        "category": _cat,
                    }
                    if info.get("_auto_topic") and isinstance(info.get("topic_context"), dict):
                        _editorial = post_data.get("editorial_brief")
                        _editorial = dict(_editorial) if isinstance(_editorial, dict) else {}
                        _editorial["topic_context"] = info["topic_context"]
                        draft_post_data["editorial_brief"] = _editorial
                    # 집필 호출에서 같이 받은 이미지 기획을 보관본으로 넘긴다(2026-09-10).
                    # 참일 때만 키를 더하므로 기획이 없는 경로의 payload 는 바이트 그대로다
                    # (그 경우 draft_store 저장 경계가 로컬 결정론 계획으로 채운다).
                    if isinstance(post_data.get("image_plan"), dict):
                        draft_post_data["image_plan"] = post_data["image_plan"]
                    _draft_images = list(all_images or [])
                    # [HOOK G] 이미지 자동 채움(Phase 2b) — IMAGE_FILL_ENABLED 게이트+
                    # fail-open. OFF 면 이 블록 전체가 no-op → draft_images·draft_post_data 는
                    # 기존과 바이트 동일(새 키 0). 어떤 실패도 초안 저장을 막지 않는다(텍스트만
                    # 저장으로 폴백). ★fill 에는 마크다운 본문(post_data["content"])을 넘긴다 —
                    # Stage1 품질이 여기서 산다(draft 로 저장하는 content_html 아님).
                    # 2026-09-09: 기본 엔진이 젠스파크(브라우저) → Codex CLI 로 바뀌었다.
                    # 게이트도 엔진 중립 키로 옮기되, 새 키가 없는 기존 설치는 예전
                    # GENSPARK_ENABLED 값으로 폴백해 **켜짐/꺼짐은** 그대로 유지된다.
                    # (엔진 선택 자체는 폴백하지 않고 codex 가 기본 — draft_image_fill 헤더 참조.)
                    _fill_on = getattr(s, "IMAGE_FILL_ENABLED", None)
                    if _fill_on is None:
                        _fill_on = getattr(s, "GENSPARK_ENABLED", False)
                    if _fill_on:
                        try:
                            from core import draft_image_fill
                            _fill_src = {
                                "title":    post_data["title"],
                                "content":  post_data.get("content", ""),
                                "category": _cat,
                            }
                            # 집필 때 만든 계획을 그대로 넘긴다 — 채움이 기획을 다시
                            # 만들지 않게(계획 없는 글만 채움 쪽에서 로컬로 1회 생성).
                            if isinstance(post_data.get("image_plan"), dict):
                                _fill_src["image_plan"] = post_data["image_plan"]
                            # 채움이 단일 생성 주체인 경로에서는 이미 있는 자산을 채움에
                            # 넘겨 **재사용**시킨다(reuse=True 기본 — 실재하는 파일만 살고
                            # 사라진 경로는 meta.dropped_missing 으로 계수돼 그 슬롯이 다시
                            # 대상이 된다). 레거시 경로(구 dispatcher 가 이미 만든 경우)는
                            # 예전 그대로 넘기지 않고 아래에서 뒤에 병합만 한다.
                            _fill_sole = _plan_fill_is_sole_image_producer(save_only)
                            _prior_imgs = list(all_images or [])
                            if _fill_sole and _prior_imgs:
                                _fill_src["image_paths"] = list(_prior_imgs)
                            _fill_res = draft_image_fill.fill_images_for_draft(
                                _fill_src, slug=_draft_slug, category=_cat)
                            _fill_imgs = _fill_res.get("image_paths") or []
                            _fill_meta = _fill_res.get("image_meta") or {}
                            if _fill_imgs:
                                # ★교체가 아니라 병합 — save_only 진입 전 _run_image_dispatcher 가
                                # 이미 만든 기존 이미지(all_images)를 무경고 폐기하지 않는다.
                                # 새로 채운 것을 앞에 둬 썸네일이 [0] 유지 + 중복 제거.
                                # 단일 생성 주체 경로에서는 기존 자산을 이미 채움에
                                # 넘겼으므로(위) 여기서 다시 붙이지 않는다 — 파일이
                                # 사라진 경로를 되살리지 않기 위함.
                                _prior = [] if _fill_sole else list(_prior_imgs)
                                _draft_images = list(_fill_imgs) + [
                                    x for x in _prior if x not in _fill_imgs]
                                if _prior:
                                    log.info(
                                        f"  🖼  기존 이미지 {len(_prior)}장을 새로 채운 "
                                        f"{len(_fill_imgs)}장 뒤에 병합(보존)")
                                if _fill_src.get("_force_image_insert"):
                                    draft_post_data["_force_image_insert"] = True
                                log.info(
                                    f"  🖼  이미지 자동 생성 {len(_fill_imgs)}장 "
                                    f"(새로 만든 {_fill_meta.get('generated', 0)}장 · "
                                    f"실패 슬롯 {_fill_meta.get('slots_failed', 0)})")
                            else:
                                # ★글이 먼저다 — 이미지가 한 장도 안 나와도 본문은 그대로
                                #   저장하고, 왜 못 만들었는지만 남긴다. 다음 보충이
                                #   빠진 슬롯만 다시 만든다(아래 계획·슬롯 승계).
                                log.warning(
                                    "  ⚠ 이미지 자동 생성 0장 — 본문만 보관합니다"
                                    f"(사유 {_fill_meta.get('skipped') or '알 수 없음'}). "
                                    "다음 '이미지 채움'이 빠진 자리만 다시 만듭니다.")
                            # 계획·슬롯 상태는 **성공/실패와 무관하게** 승계한다 —
                            # 계획이 저장돼야 다음 보충이 기획을 다시 만들지 않고,
                            # 슬롯 map 이 저장돼야 이미 만든 자리를 중복 생성하지 않는다.
                            if isinstance(_fill_src.get("image_plan"), dict):
                                draft_post_data["image_plan"] = _fill_src["image_plan"]
                            if isinstance(_fill_src.get("image_slots"), dict):
                                draft_post_data["image_slots"] = _fill_src["image_slots"]
                        except Exception as _ge:
                            log.warning(
                                f"  ⚠ 이미지 자동 채움 실패(무시): {type(_ge).__name__}")
                    else:
                        # 사용자가 명시적으로 꺼 둔 상태 — 존중하되 조용히 넘기지 않는다.
                        log.info(
                            "  ⏭  이미지 자동 채움 OFF — 본문만 보관합니다"
                            "(설정 → 이미지 채움에서 켜면 보관함 채우기가 이미지까지 함께 만듭니다).")
                    _saved_path = save_draft(
                        slug=_draft_slug,
                        post_data=draft_post_data,
                        image_paths=_draft_images,
                        schedule_at=str(reserve_time),
                    )
                    if _saved_path and info.get("_auto_topic"):
                        from core import auto_topics
                        auto_topics.record_success(info)
                    saved_draft_count += 1
                    log.info(f"  💾 초안 저장 (save_only): {title[:46]}")
                    if progress_cb: progress_cb(title, "ok", i+1, total)
                except Exception as _se:
                    log.error(f"  ❌ 초안 저장 실패: {type(_se).__name__}: {_se}")
                    _report_draft_save_failure(progress_cb, title, i, total)
                continue

            # ── 새 글 발행 경로: 계획 기반 이미지 완전성 게이트 (2026-09-18) ──
            # 보관본 재사용(draft_hit)은 위에서 이미 _prepare_archive_images 가 계획·
            # 완전성을 확인했다. 여기는 그 자리에서 막 집필해 곧장 업로드로 넘어가는
            # 새 글 전용 — 지금까지는 이 지점에서 어떤 완전성 검사도 없이 all_images 를
            # 그대로 업로드에 넘겨, 이미지 0장도 "의도한 텍스트 전용"과 구분 없이
            # 발행됐다. 게이트가 꺼져 있으면(둘 중 하나라도 OFF) 예전 그대로 통과한다.
            if draft_hit is None and _plan_fill_covers_new_post_publish():
                _img_slug = _slug or _make_fallback_slug(title, target_date)
                _images_complete, all_images, _new_img_meta = _fill_new_post_images(
                    post_data, slug=_img_slug, category=_cat, existing_images=all_images)
                if not _images_complete:
                    log.warning(
                        "  ⛔ 이미지 미완성 — 새 글 발행을 보류하고 보관합니다 "
                        f"(생성 {_new_img_meta.get('generated', 0)} / "
                        f"실패 슬롯 {_new_img_meta.get('slots_failed', 0)})")
                    _saved_for_fix = _save_failed_draft_for_fix(
                        title, post_data, all_images, _cat, "new_post_images_not_ready",
                        slug=_img_slug, target_date=target_date,
                        schedule_at=str(reserve_time))
                    write_log(title, False, reserve_time, "new_post_images_not_ready")
                    skipped_quality_count += 1
                    if progress_cb: progress_cb(title, "fail", i + 1, total)
                    continue

            # ── 업로드 단계 ──
            log.info(f"  ⏰ 예약: {reserve_time}")
            _state_slug = (draft_hit or {}).get("slug") or _slug or ""
            _upload_result = None
            _publish_state_marked = ""
            try:
                # ★ 외부 확정 행동 직전 — 비상정지 확인. 여기서 멈추면 아직 아무것도
                #   제출하지 않았으므로 산출물(집필·이미지·QA)만 남고 중복 위험이 없다.
                if _psb is not None and _psb.guard_external_action() == "estop":
                    log.warning(
                        "⏹ 전체 비상정지 — 발행 제출로 넘어가지 않습니다"
                        "(여기까지의 글·이미지는 보관됩니다)")
                    write_log(title, False, reserve_time, "estop_before_submit")
                    if progress_cb: progress_cb(title, "fail", i + 1, total)
                    break
                if info.get("_auto_topic") and not dry_run:
                    from core import auto_topics
                    try:
                        _pending_recorded = auto_topics.record_pending(info)
                    except (OSError, ValueError, TypeError, RuntimeError) as _history_err:
                        _pending_recorded = False
                        log.error(
                            "자동 주제 원문 중복 이력 저장 중 예외가 발생했습니다: %s",
                            type(_history_err).__name__, exc_info=True)
                    if not _pending_recorded:
                        log.error(
                            "⛔ 자동 주제 원문 중복 이력을 저장하지 못해 제출을 중단합니다. "
                            "초안과 이미지 자산은 유지하고 저장소를 확인해 주세요.")
                        if draft_hit is None:
                            try:
                                from core.draft_store import save_draft
                                _blocked_slug = _slug or _make_fallback_slug(title, target_date)
                                save_draft(
                                    slug=_blocked_slug, post_data=post_data,
                                    image_paths=list(all_images or []),
                                    schedule_at=str(reserve_time))
                            except (OSError, ValueError, TypeError, RuntimeError) as _draft_err:
                                log.error(
                                    "자동 주제 이력 실패 후 초안 보관도 실패했습니다. "
                                    "현재 생성 결과와 이미지 파일을 확인해 주세요.",
                                    exc_info=True)
                        _publication_results[i] = "failed"
                        write_log(title, False, reserve_time,
                                  "auto_topic_history_write_failed_before_submit")
                        if progress_cb:
                            progress_cb(title, "fail", i + 1, total)
                        continue
                # ★ write-ahead — 에디터에 쓰기 시작하기 **전에** 디스크에 남긴다.
                #   이 값이 남아 있는 채로 크래시하면 결말 미상 = 재시작 후 대조 대상.
                if _psb is not None:
                    _psb.mark("reserving", slug=_state_slug, title=title,
                              publish_at=str(reserve_time),
                              attempt_at=datetime.now().astimezone()
                              .isoformat(timespec="seconds"))
                if s.USE_NAVER_API and s.NAVER_API_CLIENT_ID:
                    from core.naver_api_uploader import upload_post_api
                    result = upload_post_api(post_data, all_images, reserve_time)
                    ok     = result.get("success", False)
                    url    = result.get("url", "")
                else:
                    from core.naver_uploader import upload_post
                    _upload_pd = {
                        "title":    post_data["title"],
                        "content":  post_data["content_html"],
                        "tags":     post_data.get("tags", []),
                        "category": post_data.get("category"),
                    }
                    # 젠스파크 채운 이미지의 per-post 삽입 신호를 upload_post 로 전달한다.
                    # 참일 때만 키 추가 → 없으면 dict 는 기존과 동일(_should_insert_images 가
                    # SKIP_IMAGE_INSERT 기본값을 따라 삽입 안 함, 옛 동작 유지).
                    if post_data.get("_force_image_insert"):
                        _upload_pd["_force_image_insert"] = True
                    _reserve_arg = reserve_time if s.USE_RESERVATION else None
                    # 무인 운영 컨텍스트이거나 미확인 재시도 금지가 켜졌을 때만
                    # 확장 계약(upload_post_ex)을 쓴다. 그 외에는 예전 그대로
                    # upload_post(bool) — 기존 설치의 동작이 한 줄도 바뀌지 않는다.
                    _want_ex = bool(
                        _psb is not None
                        and (_psb.recording()
                             or getattr(s, "UPLOAD_RECONCILE_ON_UNCONFIRMED", False)))
                    _up_ex = None
                    if _want_ex:
                        try:
                            from core.naver_uploader import upload_post_ex as _up_ex
                        except ImportError:
                            _up_ex = None
                    if _up_ex is not None:
                        _upload_result = _up_ex(
                            driver, _upload_pd, all_images,
                            reserve_datetime=_reserve_arg)
                        ok = bool((_upload_result or {}).get("ok"))
                    else:
                        ok = upload_post(driver, _upload_pd, all_images,
                                         reserve_datetime=_reserve_arg)
                    url = ""

                # ★ 발행 결과를 상태로 번역해 기록한다. 증거 없는 성공은 reserved 가
                #   아니라 unconfirmed 이고, 제출 뒤 신호 불명도 unconfirmed 다
                #   (= 자동 재발행 금지, 사람 확인 큐).
                if _psb is not None:
                    _publish_state_marked, _ps_reason = _psb.classify_upload_result(
                        _upload_result if _upload_result is not None else ok)
                    _psb.mark(_publish_state_marked, slug=_state_slug, title=title,
                              reason=_ps_reason, publish_at=str(reserve_time))

                # An optimistic uploader response is not a confirmed reservation.
                if _publish_state_marked == "unconfirmed":
                    ok = False

                if ok:
                    # 드레인 항목(보관함 우선 소진)은 대응하는 엑셀 셀이 없으므로
                    # mark_done 을 스킵한다. 재드레인 방지는 아래 draft_hit
                    # mark_status(uploaded) 가 담당한다(draft_ready → uploaded).
                    if (not info.get("_drain") and not info.get("_auto_topic")
                            and info.get("excel_path") and info.get("row") and info.get("col")):
                        mark_done(info["excel_path"], info["row"], info["col"], cfg=s)
                    if info.get("_auto_topic"):
                        from core import auto_topics
                        auto_topics.record_success(info)
                    # 본문 키워드까지 히스토리에 저장 → 이후 본문 중복 검사가 유효.
                    record_post(title, url,
                                content=post_data.get("content")
                                or post_data.get("content_html", ""))
                    write_log(title, True, reserve_time, url)
                    # 원격 URL 을 실제로 받은 경로(API 발행)만 live 로 올린다.
                    # Selenium 예약 경로는 URL 이 없으므로 reserved/unconfirmed 그대로 두고,
                    # 공개 이후 원격 대조가 log_no 를 확보했을 때만 live 가 된다.
                    if url and _psb is not None:
                        _log_no = ""
                        _m = re.search(r"/(\d{6,})(?:[/?#]|$)", str(url))
                        if _m:
                            _log_no = _m.group(1)
                        _psb.mark("live", slug=_state_slug, title=title,
                                  url=url, log_no=_log_no or None,
                                  reason="api_url_returned")

                    # 보관본을 사용한 발행이면 draft_store 상태를 'uploaded' 로 갱신.
                    # mark_status 실패는 발행 결과에 영향 없음 — warning 만 남기고 계속.
                    if draft_hit is not None:
                        try:
                            from core.draft_store import mark_status
                            mark_status(draft_hit["slug"], "uploaded", url=url)
                        except Exception as _ms_err:
                            log.warning(
                                f"  ⚠️  draft_store mark_status(uploaded) 실패 — "
                                f"{draft_hit.get('slug')}: {_ms_err}"
                            )

                    try:
                        from core.series_planner import find_episode_in_series, mark_episode_published
                        s_obj, ep_obj = find_episode_in_series(title)
                        if s_obj and ep_obj:
                            mark_episode_published(
                                s_obj["id"], ep_obj["ep"],
                                post_data.get("title", title), url)
                    except (ImportError, KeyError):
                        pass

                    try:
                        from core.channel_dispatcher import (
                            send_kakao_notification, enqueue_cafe_posts)
                        dispatch_data = {**post_data, "url": url}
                        send_kakao_notification(dispatch_data, url)
                        enqueue_cafe_posts(dispatch_data, url)
                    except (ImportError, urllib.error.URLError) as _ce:
                        log.warning(f"  ⚠️  채널 디스패치 오류: {_ce}")

                    try:
                        from core.notifier import notify_success
                        notify_success(post_data["title"], reserve_time, success+1, total)
                    except ImportError:
                        pass
                    success += 1
                    _publication_results[i] = "registered"
                    log.info(f"  ✅ 예약발행 등록 완료 ({success}/{total}) → {reserve_time}")
                    # 예약 대장 기록 — brief 유무와 무관하게 '가장 먼 예약 다음날' 이월 보장.
                    try:
                        from core import reservation_ledger as _rl
                        _rl.record(str(reserve_time), slug=_slug or "", category=_cat)
                    except Exception as _rl_exc:  # noqa: BLE001
                        log.debug(f"  예약 대장 기록 실패(무시): {_rl_exc}")
                    if progress_cb: progress_cb(title, "ok", i+1, total)
                    # [HOOK D-success] 하네스 — 발행 성공 brief
                    if _he is not None and _slug:
                        try:
                            _he.emit_publish_brief(
                                _slug, post_data.get("title", _clean_kw), _cat,
                                tags=post_data.get("tags", []),
                                schedule_at=str(reserve_time),
                                qa_status="PASS", success=True, url=url,
                                publish_method="api" if (s.USE_NAVER_API and s.NAVER_API_CLIENT_ID) else "selenium",
                                dia_score=post_data.get("_dia_score"),
                                aeo_score=post_data.get("_aeo_score"))
                            published_records.append({
                                "slug": _slug, "title": post_data.get("title", _clean_kw),
                                "category": _cat, "reserve_time": str(reserve_time),
                                "success": True, "url": url})
                        except Exception:
                            pass
                else:
                    # 제출 뒤 신호를 못 읽은 경우(unconfirmed)는 **실패가 아니다** —
                    # 실제로는 올라갔을 수 있으므로 실패로 확정하지도, 자동 재발행
                    # 하지도 않는다. 보관본은 needs_fix 로 내려 재드레인 대상에서
                    # 빼고(중복 발행 입구 차단) 사람 확인 큐에 남긴다.
                    _unconf = (_publish_state_marked == "unconfirmed")
                    _publication_results[i] = "unconfirmed" if _unconf else "failed"
                    write_log(title, False, reserve_time,
                              "publish_unconfirmed" if _unconf else "upload_failed")
                    if progress_cb:
                        progress_cb(title, _publish_failure_progress_state(_unconf), i+1, total)
                    if _unconf:
                        log.warning(
                            "  ⛔ 발행 성공 신호를 읽지 못했습니다 — 같은 글을 다시 "
                            "올리지 않습니다. 네이버에서 실제 등록 여부를 확인하세요.")
                    # 보관본 발행이 실패하면 draft_store 상태를 'failed' 로 표시 —
                    # 다음 회차에 같은 키워드로 다시 활용되지 않게 한다.
                    if draft_hit is not None:
                        try:
                            from core.draft_store import mark_status
                            if _unconf:
                                mark_status(
                                    draft_hit["slug"], "needs_fix",
                                    error="publish_unconfirmed: 네이버에서 실제 "
                                          "등록 여부 확인 필요(자동 재발행 안 함)")
                            else:
                                mark_status(
                                    draft_hit["slug"], "failed", error="upload_failed"
                                )
                        except Exception as _ms_err:
                            log.warning(
                                f"  ⚠️  draft_store mark_status(failed) 실패 — "
                                f"{draft_hit.get('slug')}: {_ms_err}"
                            )
                    # [HOOK D-fail] 하네스 — 업로드 실패 brief
                    if _he is not None and _slug:
                        try:
                            _he.emit_publish_brief(
                                _slug, post_data.get("title", _clean_kw), _cat,
                                tags=post_data.get("tags", []),
                                schedule_at=str(reserve_time),
                                qa_status="PASS", success=False, url="",
                                error="upload_failed",
                                publish_method="api" if (s.USE_NAVER_API and s.NAVER_API_CLIENT_ID) else "selenium",
                                dia_score=post_data.get("_dia_score"),
                                aeo_score=post_data.get("_aeo_score"))
                            published_records.append({
                                "slug": _slug, "title": post_data.get("title", _clean_kw),
                                "category": _cat, "reserve_time": str(reserve_time),
                                "success": False, "url": "", "error": "upload_failed"})
                        except Exception:
                            pass

            except Exception as e:
                if _publication_results.get(i) == "registered":
                    success -= 1
                _publication_results[i] = (
                    "failed" if _publish_state_marked == "failed" else "unconfirmed")
                log.error(f"  ❌ 업로드 오류: {e}\n{traceback.format_exc()}")
                write_log(title, False, reserve_time, str(e))
                # 업로드 예외는 remote 결과를 판정하지 못한 상태다 — 실패 표시를
                # 통해 재발행 대상으로 오인시키지 않도록 확인 필요로만 전달한다.
                if progress_cb: progress_cb(title, "verify", i+1, total)
                # ★ 업로드 도중 예외 = **결말 미상**. 응답이 끊겼을 뿐 실제로는
                #   등록됐을 수 있으므로 실패로 확정하지 않고 unconfirmed 로 남긴다
                #   (재시작 후 원격 대조가 결말을 시도한다). 이미 write-ahead 로
                #   reserving 이 적혀 있으면 그 위에 얹는다.
                if _psb is not None and _publish_state_marked != "failed":
                    try:
                        _psb.mark("unconfirmed", slug=_state_slug, title=title,
                                  reason=f"upload_exception:{type(e).__name__}",
                                  publish_at=str(reserve_time))
                        _publish_state_marked = "unconfirmed"
                    except Exception:  # noqa: BLE001
                        pass
                failed = safe_load_json(str(FAILED_POSTS_PATH), default=list)
                failed.append({"title": title, "reason": "업로드 실패",
                               "date": datetime.now().strftime("%Y-%m-%d %H:%M")})
                atomic_write_json(str(FAILED_POSTS_PATH), failed[-50:])
                # 보관본 발행 중 예외 — 'failed' 로 마킹하여 재시도 루프를 끊는다.
                if draft_hit is not None:
                    try:
                        from core.draft_store import mark_status
                        if _publish_state_marked == "unconfirmed":
                            mark_status(
                                draft_hit["slug"], "needs_fix",
                                error="publish_unconfirmed(exception): 네이버에서 "
                                      "실제 등록 여부 확인 필요(자동 재발행 안 함)")
                        else:
                            mark_status(
                                draft_hit["slug"], "failed",
                                error=f"upload_exception: {e}")
                    except Exception as _ms_err:
                        log.warning(
                            f"  ⚠️  draft_store mark_status(failed) 실패 — "
                            f"{draft_hit.get('slug')}: {_ms_err}"
                        )
                # [HOOK D-exc] 하네스 — 업로드 예외 brief
                if _he is not None and _slug:
                    try:
                        _he.emit_publish_brief(
                            _slug, post_data.get("title", _clean_kw),
                            _cat, tags=post_data.get("tags", []),
                            schedule_at=str(reserve_time),
                            qa_status="UNKNOWN", success=False, url="",
                            error=f"upload_exception: {e}",
                            dia_score=post_data.get("_dia_score"),
                            aeo_score=post_data.get("_aeo_score"))
                        published_records.append({
                            "slug": _slug, "title": post_data.get("title", _clean_kw),
                            "category": _cat, "reserve_time": str(reserve_time),
                            "success": False, "url": "", "error": "업로드 예외"})
                    except Exception:
                        pass

            # ── 최소 간격 보정 대기 + 랜덤 지터 (마지막 글 제외) ──
            # T-02: 탐지 회피를 위한 랜덤 지터 (-30 ~ +60초)
            if i < len(posts) - 1:
                jitter = random.randint(-30, 60)
                interval_with_jitter = max(120, s.POST_INTERVAL_SECONDS + jitter)
                elapsed = time.monotonic() - loop_start
                remaining = interval_with_jitter - elapsed
                if remaining > 0:
                    log.info(f"  ⏳ {remaining:.0f}초 대기 (간격 {interval_with_jitter}초 = "
                             f"기본 {s.POST_INTERVAL_SECONDS}초 + 지터 {jitter:+d}초)")
                    # 중지 신호가 최장 ~3분 지연되지 않도록 1초 단위로 분할 대기하며
                    # stop_flag 를 폴링한다. (2026-07-03)
                    while remaining > 0:
                        if stop_flag and stop_flag.is_set():
                            log.warning("  ⏹ 간격 대기 중 중지 신호 감지 — 대기 중단")
                            break
                        _s = min(1.0, remaining)
                        time.sleep(_s)
                        remaining -= _s
                else:
                    log.info(f"  ⚡ 대기 불필요 (생성+업로드 {elapsed:.0f}초 ≥ {interval_with_jitter}초)")

        # ── S2: fast-path 로 미뤄둔 준비 작업을 이제(첫 발행 뒤) 실행 ──────
        # driver 가 아직 살아 있는 지점(finally 의 _release_driver 이전)에서
        # 실행해야 리프레시가 공유 세션을 재사용할 수 있다. 이 블록 자체가
        # 예외를 던져 업로드 결과(success/pipeline_status)를 지우지 않도록
        # 각 단계는 이미 자체 try/except 로 fail-open 되어 있고, 바깥에도
        # 한 겹 더 감싼다.
        if _limited_run:
            log.info("  ⏭  지정 발행 — 준비 작업(재시도·마감·트렌드·통계·리프레시) 생략")
        elif _fast_path:
            try:
                log.info(f"\n{'━'*54}")
                log.info("🧰 지연했던 준비 작업 실행 (첫 발행 완료 후)")
                log.info(f"{'━'*54}")
                if progress_cb:
                    try:
                        progress_cb("지연했던 준비 작업 실행 중", "prep", total, total)
                    except Exception:    # noqa: BLE001
                        log.debug("  progress_cb(prep) 호출 실패 — 로그만 유지")
                if not save_only:
                    pipeline_status["auto_retry_failed"] = (
                        _run_retry_failed_deferred(stop_flag, progress_cb))
                pipeline_status["deadline_tracker"] = _run_deadline_detection_deferred(s)
                pipeline_status["trend_updater"] = _run_trend_update_deferred(s)
                pipeline_status["stats_fetcher"] = _run_stats_and_rank_deferred(dry_run)
                if not save_only:
                    pipeline_status["post_refresher"] = _run_refresh_batch_deferred(
                        dry_run, target_date, reserve_times, driver)
            except Exception as _defer_err:  # noqa: BLE001 — 지연 작업 실패가 이미 끝난 발행 결과를 해치면 안 됨
                log.warning(f"  ⚠️  지연 준비 작업 중 예외(무시): {_defer_err}")

    finally:
        if gov24_driver:
            try:
                gov24_driver.quit()
            except Exception:
                log.debug("gov24 드라이버 종료 실패")
        # attach 세션 브라우저는 닫지 않는다(다음 발행 재사용) — spawn 은 기존대로 quit.
        _release_driver(driver, _publish_mode)

        _failed_count = sum(state == "failed" for state in _publication_results.values())
        _unconfirmed_count = sum(state == "unconfirmed" for state in _publication_results.values())
        publication_outcomes = {
            "registered": success,
            "failed": _failed_count,
            "unconfirmed": _unconfirmed_count,
            "duplicate": skipped_dup_count,
            # Includes quality gates, saved drafts, and candidates not attempted.
            "other_skipped": max(total - success - _failed_count
                                 - _unconfirmed_count - skipped_dup_count, 0),
        }

        # 파이프라인 상태 저장 (대시보드 표시용) — 루프 내 미포착 예외로
        # 함수가 비정상 종료해도 상태파일이 stale 로 남지 않도록 finally 에서
        # 항상 기록한다. 저장 자체 실패는 삼켜 드라이버 정리를 방해하지 않는다.
        # (2026-07-03: pipeline_status stale 방지)
        try:
            _status_path = LOG_DIR / "pipeline_status.json"
            atomic_write_json(str(_status_path), {
                "date": (target_date or date.today().strftime("%Y-%m-%d")),
                "steps": pipeline_status,
                "total": total,
                "success": success,
                "publication_outcomes": publication_outcomes,
            })
        except (OSError, ValueError, TypeError) as _sp_err:
            log.warning(f"  ⚠️  pipeline_status.json 저장 실패: {_sp_err}")

    try:
        from core.notifier import notify_daily_summary
        # 스킵 사유를 사람 말 한 줄로 — 알림만 보고도 무슨 일이 있었는지 알게.
        _skip_parts = [
            (f"품질 {skipped_quality_count}", skipped_quality_count),
            (f"검색량 {skipped_low_volume_count}", skipped_low_volume_count),
            (f"제도중복 {skipped_institution_count}", skipped_institution_count),
        ]
        _skip_detail = " · ".join(p for p, n in _skip_parts if n > 0)
        notify_daily_summary(
            success, total, target_date or date.today().strftime("%Y-%m-%d"),
            drafts_used=draft_used_count, skipped_detail=_skip_detail,
            publication_outcomes=publication_outcomes)
    except ImportError:
        pass

    # Definite failures and unknown remote outcomes need different actions.
    # Duplicate/quality skips alone do not indicate a failed publication.
    if ((_failed_count or _unconfirmed_count) and not dry_run and not save_only
            and not (stop_flag and stop_flag.is_set())):
        _counts = (f"예약 등록 {success}편 · 실제 실패 {_failed_count}편 · "
                   f"확인 필요 {_unconfirmed_count}편 · 중복 건너뜀 {skipped_dup_count}편 · "
                   f"기타 보류/미시도 {publication_outcomes['other_skipped']}편")
        if _unconfirmed_count:
            alert_title = "발행 확인 필요"
            alert_msg = (f"⚠️ {_counts}\n"
                         "👉 확인 필요 글은 네이버에서 실제 등록 여부를 확인하세요. "
                         "확인 전 같은 글을 다시 올리지 않습니다.")
            log.warning(f"블로그 발행 확인 필요: {_counts}")
        else:
            alert_title = ("긴급 알림" if _failed_count == total else "일부 발행 실패")
            alert_msg = (f"❌ {_counts}\n"
                         "👉 GUI에서 실패 사유를 확인하세요.")
            log.error(f"블로그 발행 실패: {_counts}")
        try:
            from core.notifier import send_kakao
            send_kakao(alert_title, alert_msg)
        except ImportError:
            log.warning("카카오 알림 모듈 로드 실패 — 로그만 기록")

    log.info(f"\n{'━'*54}")
    if dry_run:
        log.info(f"🔍 [테스트 모드] {generated_count}개 글 생성 완료")
    else:
        log.info(
            f"완료: {success}/{total}개 예약발행 등록 "
            f"(실제 실패 {_failed_count}건, 확인 필요 {_unconfirmed_count}건) "
            f"(중복 스킵 {skipped_dup_count}건, 품질 게이트 스킵 {skipped_quality_count}건, "
            f"검색량 스킵 {skipped_low_volume_count}건, 제도포화 스킵 {skipped_institution_count}건)"
        )
        if skipped_quality_count > 0:
            log.warning(
                f"⚠️  품질 미달로 업로드를 스킵한 글이 {skipped_quality_count}건 있습니다. "
                f"logs/app.log에서 'quality_gate_blocked' 로 검색해 사유를 확인하세요."
            )
        if saved_draft_count > 0:
            log.info(f"💾 초안 보관: {saved_draft_count}건 (runs/<slug>/post_data.json)")
        if draft_used_count > 0:
            log.info(f"📤 보관본 활용: {draft_used_count}건 (글 생성 스킵)")

    # 파이프라인 상태 저장은 드라이버 정리 finally 블록에서 항상 수행된다
    # (루프 비정상 종료 시에도 stale 방지). 여기서는 중복 기록하지 않는다.

    # [HOOK E] 하네스 — 일일 다이제스트 리포트 (관측 전용)
    if _he is not None:
        try:
            _failed_records = [r for r in published_records if not r.get("success")]
            _he.emit_daily_report(
                target_date or date.today().strftime("%Y-%m-%d"),
                pipeline_status, total=total, success=success,
                generated=generated_count, skipped_dup=skipped_dup_count,
                failed=_failed_records, published=published_records)
        except Exception:
            pass

    # [HOOK F] 발행 후 쇼츠 영상 자동 생성(Phase②, 옵트인) — 실발행·성공 시에만.
    # detached 백그라운드라 발행 흐름을 막지 않는다. 어떤 예외도 발행 결과에 무영향.
    if (not dry_run and not save_only and success > 0
            and getattr(s, "AUTO_VIDEO_AFTER_PUBLISH", False)):
        try:
            from core.video_batch import spawn_video_batch
            r = spawn_video_batch(limit=getattr(s, "AUTO_VIDEO_LIMIT", 5))
            if r.get("launched"):
                log.info(f"🎬 발행 후 영상 배치 기동 — 대상 {r.get('count', 0)}편(백그라운드)")
        except Exception as _vb_exc:  # noqa: BLE001
            log.warning(f"영상 배치 기동 실패(무시): {_vb_exc}")

    return {"total": total, "success": success, "pipeline_status": pipeline_status,
            "publication_outcomes": publication_outcomes,
            "saved_drafts": saved_draft_count, "drafts_used": draft_used_count,
            "skipped_quality": skipped_quality_count,
            "skipped_low_volume": skipped_low_volume_count,
            "skipped_institution": skipped_institution_count}


# ─────────────────────────────────────────────────────
# 발행 전 보관함 매칭 분석 (pre-flight) — 2026-05-16
# ─────────────────────────────────────────────────────


def preflight_check(target_date: str = None,
                    max_count: int = None) -> dict:
    """오늘 실행 직전, 발행 후보와 보관함 매칭을 분석한다.

    부수 효과 0 — Selenium / 네이버 로그인 / gov24 캡처 / Claude CLI / 엑셀 쓰기
    모두 호출하지 않는다. 엑셀은 읽기만 한다(get_pending_posts 가 openpyxl 로 open
    하지만 _backup_excel 은 run_pipeline 본체에서만 호출되므로 무관).

    Args:
        target_date: 'YYYY-MM-DD'. None 이면 today.
        max_count: 분석할 후보 상한. None 이면 settings.DAILY_POST_LIMIT.

    Returns:
        draft_store.check_coverage 결과 + 후보 메타:
        {
            "by_category": {...},
            "missing_categories": [...],
            "total_candidates": N,
            "fully_covered": bool,
            "candidates": [{"title", "category", "reserve_time"}, ...],
            "target_date": "YYYY-MM-DD",
            "queue": {                       # (2026-05-26) 예약 이월 가시화
                "queued_count": int,         # 현재 미래 예약 건수
                "last_reserved_date": "YYYY-MM-DD" | None,
                "next_start_date": "YYYY-MM-DD",  # assign_reservation_times 시작점
                "carryover_days": int,       # 옛 D+1 대비 며칠 미뤄졌는지
            },
        }
        예외 발생 시 안전 dict 반환 (fully_covered=True, total_candidates=0).
    """
    _t = target_date or date.today().strftime("%Y-%m-%d")
    base_d_plus_1 = (
        datetime.strptime(_t, "%Y-%m-%d").date() + timedelta(days=1)
    )

    # 큐 메타는 후보 유무·예외 발생과 무관하게 항상 계산해 두면
    # 빈 후보일 때도 GUI 가 "큐: 6/3까지 등록됨" 한 줄을 띄울 수 있다.
    try:
        _existing_for_meta = _load_existing_reservations()
    except Exception as _meta_exc:   # noqa: BLE001
        log.debug(f"[preflight] 예약 큐 로드 실패(메타): {_meta_exc}")
        _existing_for_meta = []
    _last_reserved = (
        max(r["date"] for r in _existing_for_meta)
        if _existing_for_meta else None
    )
    _next_start = (
        max(base_d_plus_1, _last_reserved + timedelta(days=1))
        if _last_reserved else base_d_plus_1
    )
    _queue_meta = {
        "queued_count": len(_existing_for_meta),
        "last_reserved_date": (
            _last_reserved.strftime("%Y-%m-%d") if _last_reserved else None
        ),
        "next_start_date": _next_start.strftime("%Y-%m-%d"),
        "carryover_days": max((_next_start - base_d_plus_1).days, 0),
    }

    safe_default = {
        "by_category": {},
        "missing_categories": [],
        "total_candidates": 0,
        "fully_covered": True,
        "candidates": [],
        "target_date": _t,
        "queue": _queue_meta,
    }

    try:
        s = _cfg()
    except Exception as exc:
        log.warning(f"[preflight] settings 로드 실패: {exc}")
        return safe_default

    try:
        automatic = _automatic_topic_posts(
            s, refresh=False, limit=max_count or getattr(s, "DAILY_POST_LIMIT", 1))
        if automatic is not None:
            posts = automatic
        else:
            posts = get_pending_posts(target_date, cfg=s, limit=max_count)
    except (ImportError, FileNotFoundError, OSError, KeyError, TypeError, ValueError) as exc:
        log.warning(f"[preflight] 자동 주제 분석 실패: {exc}")
        posts = []

    if not posts:
        return safe_default

    # 보관함에 이미 있는 키워드는 사전 차감 — 실제 run_pipeline(save_only=False) 흐름에서는
    # 보관본을 그대로 활용하므로 매칭 카테고리에 영향 없음. 따라서 후보 카테고리 분석은
    # 차감 전 원본을 기준으로 수행한다. (사용자가 "발행 후보 N편 중 부족 카테고리" 를
    # 정확히 알고 싶어 함.)

    # 사전 중복 필터 — post_history 와 100%/threshold 매칭 후보는 실제 실행 시
    # Skip 되므로 preflight 단계에서도 미리 차감해 GUI 의 "필요 N편" 계산을 정확히 한다.
    # preflight 는 부수 효과 0 을 유지해야 하므로 mutate_excluded=False 로 호출.
    if posts:
        try:
            kept, dropped = _prefilter_known_duplicates(
                posts, reason="preflight_prefilter",
                mutate_excluded=False)
            if dropped:
                log.info(
                    f"[preflight] 사전 중복 필터: {len(dropped)}건 차감 "
                    f"({len(posts)}건 → {len(kept)}건)"
                )
            posts = kept
        except Exception as exc:    # noqa: BLE001
            log.warning(f"[preflight] 사전 중복 필터 예외 — 필터 없이 진행: {exc}")

    if not posts:
        return safe_default

    if max_count is not None and max_count > 0 and len(posts) > max_count:
        posts = posts[:max_count]

    try:
        reserve_times = assign_reservation_times(len(posts), target_date, posts, cfg=s)
    except (TypeError, ValueError, AttributeError) as exc:
        log.debug(f"[preflight] assign_reservation_times 예외 — reserve_time 생략: {exc}")
        reserve_times = [""] * len(posts)

    # 후보별 카테고리 정규화
    candidates = []
    for idx, info in enumerate(posts):
        raw_title = info.get("title", "")
        clean = re.sub(r'^\d{4}-\d{2}-\d{2}\s*[—\-]\s*', '', raw_title)
        try:
            cat = detect_category(clean)
        except (TypeError, ValueError):
            cat = "default"
        candidates.append({
            "title": raw_title,
            "category": cat,
            "reserve_time": reserve_times[idx] if idx < len(reserve_times) else "",
        })

    try:
        from core.draft_store import check_coverage
        summary = check_coverage(candidates)
    except Exception as exc:
        log.warning(f"[preflight] check_coverage 실패: {exc}")
        return {**safe_default, "candidates": candidates}

    summary["candidates"] = candidates
    summary["target_date"] = _t

    # 큐 메타: 함수 시작에서 계산한 _queue_meta 를 기본으로 쓰되, 정상 경로에서
    # reserve_times[0] 이 더 정확한 next_start_date 를 알려주면 그쪽을 신뢰한다.
    queue_meta = dict(_queue_meta)
    if reserve_times:
        first = reserve_times[0]
        if first and " " in first:
            try:
                ns = datetime.strptime(
                    first.split(" ")[0], "%Y-%m-%d"
                ).date()
                queue_meta["next_start_date"] = ns.strftime("%Y-%m-%d")
                queue_meta["carryover_days"] = max(
                    (ns - base_d_plus_1).days, 0
                )
            except ValueError:
                pass
    summary["queue"] = queue_meta
    return summary
